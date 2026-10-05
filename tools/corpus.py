"""
Differential run of every oracle over a corpus of workbooks.

    python tools/corpus.py <dir or file>... [--jobs N] [--timeout S]
                           [--checks name,...] [--out results.jsonl]
                           [--files-from list.txt] [--heavy-mb 1.0]
                           [--mem-limit-mb 3000]

Every file is checked in its own process, so a crash, a hang or a memory
blow-up caused by one file is reported for that file and does not stop the
run. Checks (all by default):

* ``full``, ``data_only``, ``read_only``, ``read_only_data_only``,
  ``rich_text``, ``keep_vba``, ``no_links``: tests/oracle/compare.py - every
  cell, style, worksheet and workbook attribute, warnings, load errors and
  the saved package (byte for byte);
* ``ext``, ``ext_data_only``: openrsxl.extended streaming with every flag
  against openpyxl's full mode (tests/oracle/extended.py);
* ``ext_subset``: a random (per file, reproducible) subset of the flags.

The XML backend follows ``OPENPYXL_LXML`` as usual; ``OPENRSXL_PURE_PYTHON=1``
checks the Python fallbacks. Exit status 1 if any file differs.
"""

import json
import os
import random
import subprocess
import sys
import time
import traceback
import zlib
from concurrent.futures import ThreadPoolExecutor, as_completed

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
EXTS = (".xlsx", ".xlsm", ".xltx", ".xltm")
ALL_CHECKS = (
    "full",
    "data_only",
    "read_only",
    "read_only_data_only",
    "rich_text",
    "keep_vba",
    "no_links",
    "ext",
    "ext_data_only",
    "ext_subset",
)


def _subset(path):
    from oracle.extended import ALL_FLAGS

    rnd = random.Random(zlib.crc32(os.path.basename(path).encode()))
    names = [k for k in ALL_FLAGS if rnd.random() < 0.5] or [rnd.choice(list(ALL_FLAGS))]
    return {k: True for k in names}, rnd.random() < 0.5


#: read-only iteration pads every row to the declared dimension: a sheet
#: claiming A1:XFD1048576 yields 17 billion cells (in both libraries); the
#: streaming checks of such files are skipped
MAX_STREAMED_CELLS = 5_000_000
STREAMED = {"read_only", "read_only_data_only", "ext", "ext_data_only", "ext_subset"}


def streamed_area(path):
    import openpyxl

    try:
        wb = openpyxl.load_workbook(path, read_only=True)
    except Exception:
        return 0
    area = 0
    for ws in wb.worksheets:
        try:
            area = max(area, (ws.max_row or 0) * (ws.max_column or 0))
        except Exception:
            pass
    wb.close()
    return area


def check_one(path, checks):
    sys.path.insert(0, os.path.join(ROOT, "tests"))
    from oracle.compare import compare_workbooks
    from oracle.extended import compare_extended

    area = streamed_area(path) if STREAMED & set(checks) else 0

    runs = {
        "full": lambda: compare_workbooks(path),
        "data_only": lambda: compare_workbooks(path, data_only=True),
        "read_only": lambda: compare_workbooks(path, read_only=True),
        "read_only_data_only": lambda: compare_workbooks(path, read_only=True, data_only=True),
        "rich_text": lambda: compare_workbooks(path, rich_text=True),
        "keep_vba": lambda: compare_workbooks(path, keep_vba=True),
        "no_links": lambda: compare_workbooks(path, keep_links=False),
        "ext": lambda: compare_extended(path),
        "ext_data_only": lambda: compare_extended(path, data_only=True),
    }

    def subset():
        flags, data_only = _subset(path)
        return compare_extended(path, data_only=data_only, flags=flags)

    runs["ext_subset"] = subset
    out = {}
    for name in checks:
        if name in STREAMED and area > MAX_STREAMED_CELLS:
            out[name] = {"diffs": [], "skipped": f"read-only dimension of {area} cells", "time": 0}
            continue
        t = time.perf_counter()
        try:
            diffs = runs[name]()
        except BaseException as e:  # harness failure: report, keep going
            diffs = [f"HARNESS {type(e).__name__}: {e}\n{traceback.format_exc()[-1500:]}"]
        out[name] = {"diffs": [d[:1500] for d in diffs[:8]], "time": round(time.perf_counter() - t, 2)}
    return out


def collect(paths):
    files = []
    for p in paths:
        if os.path.isdir(p):
            for d, dirs, names in os.walk(p):
                dirs[:] = sorted(x for x in dirs if x != ".git")
                files.extend(os.path.join(d, n) for n in sorted(names) if n.lower().endswith(EXTS))
        else:
            files.append(p)
    return files


#: committed memory (MB) a worker may use; above it the file is reported as
#: "memory" (the oracle - openpyxl's full mode - needs GBs on some files)
MEM_LIMIT_MB = 3000


def _tree_memory(pid):
    import psutil

    try:
        procs = [psutil.Process(pid)]
        procs += procs[0].children(recursive=True)
    except psutil.Error:
        return 0
    total = 0
    for proc in procs:
        try:
            mi = proc.memory_info()
            total += getattr(mi, "private", 0) or mi.rss
        except psutil.Error:
            pass
    return total / 2**20


def _kill_tree(proc):
    import psutil

    try:
        for child in psutil.Process(proc.pid).children(recursive=True):
            child.kill()
    except psutil.Error:
        pass
    proc.kill()
    proc.communicate()


def run_file(path, checks, timeout):
    cmd = [sys.executable, os.path.abspath(__file__), "--one", path, "--checks", ",".join(checks)]
    t = time.perf_counter()
    proc = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, cwd=ROOT)
    peak = 0.0
    while True:
        try:
            out, err = proc.communicate(timeout=1)
            break
        except subprocess.TimeoutExpired:
            peak = max(peak, _tree_memory(proc.pid))
            if peak > MEM_LIMIT_MB:
                _kill_tree(proc)
                return {
                    "file": path,
                    "status": "memory",
                    "time": round(time.perf_counter() - t, 2),
                    "peak_mb": round(peak),
                }
            if time.perf_counter() - t > timeout:
                _kill_tree(proc)
                return {"file": path, "status": "timeout", "time": timeout}
    p = subprocess.CompletedProcess(cmd, proc.returncode, out, err)
    res = {"file": path, "time": round(time.perf_counter() - t, 2)}
    lines = p.stdout.decode("utf-8", "replace").strip().splitlines()
    if p.returncode != 0 or not lines:
        res.update(status="crash", returncode=p.returncode, stderr=p.stderr.decode("utf-8", "replace")[-3000:])
        return res
    res["checks"] = json.loads(lines[-1])
    bad = {k: v for k, v in res["checks"].items() if v["diffs"]}
    res["status"] = "different" if bad else "identical"
    return res


def lower_priority():
    """Run at idle priority (inherited by the worker processes): a corpus run
    takes hours and must not make the machine unresponsive."""
    try:
        import psutil

        psutil.Process().nice(psutil.IDLE_PRIORITY_CLASS if os.name == "nt" else 19)
    except Exception:
        pass


def main(argv):
    # diffs contain arbitrary text: never fail on a console encoding
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    lower_priority()
    if argv[:1] == ["--one"]:
        path = argv[1]
        checks = argv[argv.index("--checks") + 1].split(",") if "--checks" in argv else list(ALL_CHECKS)
        print(json.dumps(check_one(path, checks)))
        return 0
    jobs, timeout, checks, out = max(1, min(4, (os.cpu_count() or 2) // 4)), 900, list(ALL_CHECKS), None
    heavy_mb = 1.0
    paths = []
    it = iter(argv)
    for a in it:
        if a == "--jobs":
            jobs = int(next(it))
        elif a == "--timeout":
            timeout = int(next(it))
        elif a == "--heavy-mb":
            heavy_mb = float(next(it))
        elif a == "--mem-limit-mb":
            global MEM_LIMIT_MB
            MEM_LIMIT_MB = int(next(it))
        elif a == "--checks":
            checks = next(it).split(",")
        elif a == "--out":
            out = next(it)
        elif a == "--files-from":
            with open(next(it), encoding="utf-8") as f:
                paths.extend(line.strip() for line in f if line.strip())
        else:
            paths.append(a)
    files = collect(paths)
    print(
        f"{len(files)} files, checks: {','.join(checks)}, "
        f"OPENPYXL_LXML={os.environ.get('OPENPYXL_LXML', '')} "
        f"OPENRSXL_PURE_PYTHON={os.environ.get('OPENRSXL_PURE_PYTHON', '')}",
        flush=True,
    )
    failed = skipped = 0
    sink = open(out, "w", encoding="utf-8") if out else None
    # big files need GBs in openpyxl's full mode (the oracle): they run one at
    # a time, after the others
    small = [f for f in files if os.path.getsize(f) < heavy_mb * 2**20]
    heavy = [f for f in files if os.path.getsize(f) >= heavy_mb * 2**20]

    def results():
        for group, n in ((small, jobs), (heavy, 1)):
            with ThreadPoolExecutor(n) as ex:
                futures = [ex.submit(run_file, f, checks, timeout) for f in group]
                for fut in as_completed(futures):
                    yield fut.result()

    for res in results():
        if sink:
            sink.write(json.dumps(res) + "\n")
            sink.flush()
        skips = sorted({v["skipped"] for v in res.get("checks", {}).values() if "skipped" in v})
        if skips:
            skipped += 1
            print(f"SKIPPED streaming checks ({', '.join(skips)}): {res['file']}", flush=True)
        if res["status"] != "identical":
            failed += 1
            print(f"{res['status'].upper()}: {res['file']}", flush=True)
            for k, v in res.get("checks", {}).items():
                for d in v["diffs"][:3]:
                    print(f"  [{k}] {d[:600]}", flush=True)
            if "stderr" in res:
                print("  " + res["stderr"][-800:], flush=True)
    if sink:
        sink.close()
    print(f"{len(files) - failed}/{len(files)} identical ({skipped} with skipped streaming checks)", flush=True)
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
