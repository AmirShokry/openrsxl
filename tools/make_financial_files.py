"""
Generate randomized financial workbooks: the kinds of sheets finance teams
keep (general ledger, trial balance, monthly P&L, AR ageing, budget vs
actual, cash flow statement, FX rates, payroll, fixed asset register), with
formulas *and* their cached results (as Excel saves them), accounting number
formats, merged headers, subtotal rows and the sparse layout of reports.

Nothing is tuned to a particular file: every sheet is drawn from a seed
and a size, so the compression's analysis is measured on a population of
financial sheets rather than on one workbook.

    python tools/make_financial_files.py <outdir> [--seeds 1,2,3] [--scale 1.0]
"""

import argparse
import datetime
import os
import random
import re
import zipfile

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ACCOUNTING = '_(* #,##0.00_);_(* (#,##0.00);_(* "-"??_);_(@_)'
AMOUNT = "#,##0.00"
PCT = "0.0%"
BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin")
TOP = Border(top=THIN)
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


class Book:
    """A workbook plus the cached result of every formula cell"""

    def __init__(self):
        self.wb = openpyxl.Workbook()
        self.wb.remove(self.wb.active)
        self.cached = {}  # sheet title -> {coordinate: value}

    def sheet(self, title):
        ws = self.wb.create_sheet(title)
        self.cached[title] = {}
        return ws

    def formula(self, ws, row, col, f, value, fmt=None):
        c = ws.cell(row, col, f)
        self.cached[ws.title][c.coordinate] = value
        if fmt:
            c.number_format = fmt
        return c

    def save(self, path):
        self.wb.save(path)
        _inject_cached(path, self.wb, self.cached)


def _inject_cached(path, wb, cached):
    """write the formulas' cached results into the saved file (openpyxl
    writes formulas without them; Excel always saves both)"""
    names = {ws.title: f"xl/worksheets/sheet{i}.xml" for i, ws in enumerate(wb.worksheets, 1)}
    src = zipfile.ZipFile(path)
    items = {i.filename: src.read(i.filename) for i in src.infolist()}
    src.close()
    for title, values in cached.items():
        part = names[title]
        xml = items[part].decode("utf-8")

        def repl(m, values=values):
            ref, attrs, f = m.group(1), m.group(2), m.group(3)
            v = values.get(ref)
            if v is None or v == "":
                return m.group(0)
            if isinstance(v, bool):
                return f'<c r="{ref}"{attrs} t="b"><f>{f}</f><v>{int(v)}</v></c>'
            if isinstance(v, str):
                return f'<c r="{ref}"{attrs} t="str"><f>{f}</f><v>{v}</v></c>'
            return f'<c r="{ref}"{attrs}><f>{f}</f><v>{repr(float(v)) if isinstance(v, float) else v}</v></c>'

        xml = re.sub(r'<c r="([A-Z]+[0-9]+)"((?: s="[0-9]+")?)><f>(.*?)</f><v\s*/?>(?:</v>)?</c>', repl, xml)
        items[part] = xml.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in items.items():
            z.writestr(name, data)


def money(rnd, lo, hi):
    return round(rnd.uniform(lo, hi), 2)


def chart_of_accounts(rnd, n):
    kinds = [
        (1000, "Cash"),
        (1100, "Receivables"),
        (1500, "Fixed Assets"),
        (2000, "Payables"),
        (2100, "Accruals"),
        (3000, "Equity"),
        (4000, "Revenue"),
        (5000, "Cost of Sales"),
        (6000, "Operating Expenses"),
        (7000, "Finance Costs"),
    ]
    accts = []
    for i in range(n):
        base, name = kinds[i % len(kinds)]
        accts.append((base + 10 * (i // len(kinds)) + rnd.randint(0, 9), f"{name} - {rnd.choice('ABCDEFGH')}{i}"))
    return accts


def header(ws, row, labels, col=1):
    for j, h in enumerate(labels):
        c = ws.cell(row, col + j, h)
        c.font = BOLD
        c.fill = HEAD_FILL


def general_ledger(b, rnd, rows, title="General Ledger"):
    ws = b.sheet(title)
    ws["A1"] = f"{title} - FY{rnd.randint(2019, 2025)}"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:J1")
    header(
        ws,
        3,
        [
            "Date",
            "Journal",
            "Account",
            "Account Name",
            "Cost Centre",
            "Ccy",
            "Description",
            "Debit",
            "Credit",
            "Balance",
        ],
    )
    accts = chart_of_accounts(rnd, rnd.randint(30, 90))
    centres = [f"CC{rnd.randint(100, 999)}" for _ in range(rnd.randint(4, 15))]
    ccys = rnd.sample(["GBP", "EUR", "USD", "SGD", "JPY", "CHF"], rnd.randint(1, 3))
    day = datetime.date(2024, 1, 1)
    bal = 0.0
    j = rnd.randint(1000, 5000)
    for r in range(4, 4 + rows):
        if rnd.random() < 0.3:
            day += datetime.timedelta(days=1)
        if rnd.random() < 0.4:
            j += 1
        code, name = rnd.choice(accts)
        amt = money(rnd, 5, 50000) if rnd.random() < 0.9 else money(rnd, 50000, 2e6)
        debit = rnd.random() < 0.5
        ws.cell(r, 1, day).number_format = "dd/mm/yyyy"
        ws.cell(r, 2, f"JE{j}")
        ws.cell(r, 3, code)
        ws.cell(r, 4, name)
        ws.cell(r, 5, rnd.choice(centres))
        ws.cell(r, 6, rnd.choice(ccys))
        ws.cell(
            r, 7, rnd.choice(["Invoice", "Payment", "Accrual", "Reversal", "Journal"]) + f" {rnd.randint(10000, 99999)}"
        )
        if debit:
            ws.cell(r, 8, amt).number_format = ACCOUNTING
            bal = round(bal + amt, 2)
        else:
            ws.cell(r, 9, amt).number_format = ACCOUNTING
            bal = round(bal - amt, 2)
        f = f"=H{r}-I{r}" if r == 4 else f"=J{r - 1}+H{r}-I{r}"
        b.formula(ws, r, 10, f, bal, ACCOUNTING)
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:J{3 + rows}"
    for col, w in zip("ABCDEFGHIJ", (11, 9, 9, 28, 10, 5, 22, 14, 14, 14)):
        ws.column_dimensions[col].width = w


def trial_balance(b, rnd, n):
    ws = b.sheet("Trial Balance")
    ws["A1"] = "Trial Balance"
    ws["A1"].font = BOLD
    header(ws, 3, ["Account", "Name", "Opening", "Debits", "Credits", "Closing"])
    accts = sorted(chart_of_accounts(rnd, n))
    r = 4
    group_start = r
    group = accts[0][0] // 1000
    totals = []

    def subtotal(r, start, end):
        ws.cell(r, 2, "Subtotal").font = BOLD
        for col in range(3, 7):
            letter = get_column_letter(col)
            vals = [b.cached[ws.title].get(f"{letter}{i}", ws.cell(i, col).value) for i in range(start, end + 1)]
            c = b.formula(
                ws, r, col, f"=SUM({letter}{start}:{letter}{end})", round(sum(v or 0 for v in vals), 2), ACCOUNTING
            )
            c.font = BOLD
            c.border = TOP
        totals.append(r)

    for code, name in accts:
        if code // 1000 != group:
            subtotal(r, group_start, r - 1)
            r += 2
            group_start = r
            group = code // 1000
        opening = money(rnd, -1e6, 1e6) if rnd.random() < 0.8 else 0
        dr = money(rnd, 0, 5e5) if rnd.random() < 0.7 else 0
        cr = money(rnd, 0, 5e5) if rnd.random() < 0.7 else 0
        ws.cell(r, 1, code)
        ws.cell(r, 2, name)
        for col, v in ((3, opening), (4, dr), (5, cr)):
            ws.cell(r, col, v).number_format = ACCOUNTING
        b.formula(ws, r, 6, f"=C{r}+D{r}-E{r}", round(opening + dr - cr, 2), ACCOUNTING)
        r += 1
    subtotal(r, group_start, r - 1)


def pnl_monthly(b, rnd, lines):
    ws = b.sheet("P&L Monthly")
    ws["A1"] = "Profit and Loss"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:N1")
    header(ws, 3, ["Line"] + MONTHS + ["FY"])
    sections = [("Revenue", lines), ("Cost of Sales", max(2, lines // 2)), ("Operating Expenses", lines)]
    r = 4
    totals = {}
    for sec, n in sections:
        ws.cell(r, 1, sec).font = BOLD
        r += 1
        start = r
        for i in range(n):
            ws.cell(r, 1, f"  {sec[:3]} item {i + 1}").alignment = Alignment(indent=1)
            base = money(rnd, 1000, 90000)
            row_vals = []
            for m in range(12):
                v = round(base * rnd.uniform(0.8, 1.25), 2) if rnd.random() < 0.95 else 0
                ws.cell(r, 2 + m, v).number_format = ACCOUNTING
                row_vals.append(v)
            b.formula(ws, r, 14, f"=SUM(B{r}:M{r})", round(sum(row_vals), 2), ACCOUNTING)
            r += 1
        ws.cell(r, 1, f"Total {sec}").font = BOLD
        for col in range(2, 15):
            letter = get_column_letter(col)
            s = sum(b.cached[ws.title].get(f"{letter}{i}", ws.cell(i, col).value) or 0 for i in range(start, r))
            c = b.formula(ws, r, col, f"=SUM({letter}{start}:{letter}{r - 1})", round(s, 2), ACCOUNTING)
            c.font = BOLD
            c.border = TOP
        totals[sec] = r
        r += 2
    ws.cell(r, 1, "Gross Profit").font = BOLD
    gp = r
    for col in range(2, 15):
        L = get_column_letter(col)
        v = b.cached[ws.title][f"{L}{totals['Revenue']}"] - b.cached[ws.title][f"{L}{totals['Cost of Sales']}"]
        b.formula(ws, r, col, f"={L}{totals['Revenue']}-{L}{totals['Cost of Sales']}", round(v, 2), ACCOUNTING)
    r += 1
    ws.cell(r, 1, "Gross Margin %")
    for col in range(2, 15):
        L = get_column_letter(col)
        rev = b.cached[ws.title][f"{L}{totals['Revenue']}"]
        b.formula(
            ws,
            r,
            col,
            f"=IF({L}{totals['Revenue']}=0,0,{L}{gp}/{L}{totals['Revenue']})",
            b.cached[ws.title][f"{L}{gp}"] / rev if rev else 0,
            PCT,
        )
    ws.freeze_panes = "B4"


def ar_ageing(b, rnd, n):
    ws = b.sheet("AR Ageing")
    header(ws, 1, ["Customer", "Credit Limit", "Current", "1-30", "31-60", "61-90", "90+", "Total", "Over Limit"])
    for r in range(2, 2 + n):
        ws.cell(r, 1, f"Customer {rnd.randint(1000, 9999)} Ltd")
        limit = rnd.choice([5000, 10000, 25000, 50000, 100000])
        ws.cell(r, 2, limit).number_format = AMOUNT
        vals = []
        for col in range(3, 8):
            v = money(rnd, 0, limit / 2) if rnd.random() < (0.8 if col == 3 else 0.25) else 0
            ws.cell(r, col, v).number_format = ACCOUNTING
            vals.append(v)
        total = round(sum(vals), 2)
        b.formula(ws, r, 8, f"=SUM(C{r}:G{r})", total, ACCOUNTING)
        b.formula(ws, r, 9, f'=IF(H{r}>B{r},"YES","")', "YES" if total > limit else "")


def budget_vs_actual(b, rnd, n):
    ws = b.sheet("Budget v Actual")
    header(ws, 2, ["Department", "Line", "Budget", "Actual", "Variance", "Var %"])
    depts = [f"Dept {c}" for c in "ABCDEFG"[: rnd.randint(3, 7)]]
    r = 3
    for d in depts:
        for i in range(n // len(depts)):
            ws.cell(r, 1, d)
            ws.cell(r, 2, rnd.choice(["Salaries", "Travel", "Software", "Rent", "Marketing", "Training"]))
            bud = rnd.choice([0, money(rnd, 1000, 200000)])
            act = round(bud * rnd.uniform(0.7, 1.3), 2) if bud else money(rnd, 0, 5000)
            ws.cell(r, 3, bud).number_format = ACCOUNTING
            ws.cell(r, 4, act).number_format = ACCOUNTING
            b.formula(ws, r, 5, f"=D{r}-C{r}", round(act - bud, 2), ACCOUNTING)
            b.formula(ws, r, 6, f'=IF(C{r}=0,"",E{r}/C{r})', (act - bud) / bud if bud else "", PCT)
            r += 1


def cash_flow(b, rnd, years):
    ws = b.sheet("Cash Flow")
    ws["B2"] = "Statement of Cash Flows"
    ws["B2"].font = Font(bold=True, size=13)
    for j in range(years):
        c = ws.cell(4, 4 + j, f"FY{2025 - years + 1 + j}")
        c.font = BOLD
        c.alignment = Alignment(horizontal="right")
    r = 6
    for block in ("Operating activities", "Investing activities", "Financing activities"):
        ws.cell(r, 2, block).font = BOLD
        r += 1
        start = r
        for i in range(rnd.randint(4, 9)):
            ws.cell(r, 2, f"{block.split()[0]} line {i + 1}").alignment = Alignment(indent=1)
            for j in range(years):
                if rnd.random() < 0.85:
                    ws.cell(r, 4 + j, money(rnd, -5e5, 5e5)).number_format = ACCOUNTING
            r += 1
        ws.cell(r, 2, f"Net cash from {block.split()[0].lower()}").font = BOLD
        for j in range(years):
            L = get_column_letter(4 + j)
            s = sum(ws.cell(i, 4 + j).value or 0 for i in range(start, r))
            b.formula(ws, r, 4 + j, f"=SUM({L}{start}:{L}{r - 1})", round(s, 2), ACCOUNTING).border = TOP
        r += 3


def fx_rates(b, rnd, days):
    ws = b.sheet("FX Rates")
    ccys = ["EUR", "USD", "SGD", "JPY", "CHF", "AUD"]
    header(ws, 1, ["Date"] + ccys)
    rates = [rnd.uniform(0.5, 180) for _ in ccys]
    d = datetime.date(2024, 1, 1)
    for r in range(2, 2 + days):
        ws.cell(r, 1, d).number_format = "yyyy-mm-dd"
        for j in range(len(ccys)):
            rates[j] *= rnd.uniform(0.995, 1.005)
            ws.cell(r, 2 + j, round(rates[j], 4)).number_format = "0.0000"
        d += datetime.timedelta(days=1)


def payroll(b, rnd, n):
    ws = b.sheet("Payroll")
    header(ws, 1, ["Emp ID", "Name", "Dept", "Grade", "Salary", "Tax", "NI", "Pension", "Net"])
    for r in range(2, 2 + n):
        ws.cell(r, 1, f"E{10000 + r}")
        ws.cell(
            r,
            2,
            f"{rnd.choice(['A', 'B', 'C', 'D', 'J', 'M', 'S'])}. {rnd.choice(['Smith', 'Jones', 'Patel', 'Khan', 'Brown', 'Lee', 'Garcia'])}",
        )
        ws.cell(r, 3, rnd.choice(["Finance", "Sales", "Ops", "IT", "HR"]))
        ws.cell(r, 4, rnd.choice(["G1", "G2", "G3", "G4"]))
        sal = round(rnd.choice([2500, 3000, 3500, 4200, 5000, 6500]) * rnd.uniform(0.95, 1.05), 2)
        ws.cell(r, 5, sal).number_format = AMOUNT
        tax = round(sal * 0.2, 2)
        ni = round(sal * 0.08, 2)
        pen = round(sal * 0.05, 2)
        b.formula(ws, r, 6, f"=ROUND(E{r}*0.2,2)", tax, AMOUNT)
        b.formula(ws, r, 7, f"=ROUND(E{r}*0.08,2)", ni, AMOUNT)
        b.formula(ws, r, 8, f"=ROUND(E{r}*0.05,2)", pen, AMOUNT)
        b.formula(ws, r, 9, f"=E{r}-F{r}-G{r}-H{r}", round(sal - tax - ni - pen, 2), AMOUNT)


def fixed_assets(b, rnd, n, years=10):
    ws = b.sheet("Fixed Assets")
    header(ws, 1, ["Asset", "Class", "Cost", "Life"] + [f"Dep Y{y + 1}" for y in range(years)] + ["NBV"])
    for r in range(2, 2 + n):
        ws.cell(r, 1, f"FA-{rnd.randint(1000, 9999)}")
        ws.cell(r, 2, rnd.choice(["IT", "Vehicles", "Furniture", "Plant", "Buildings"]))
        cost = money(rnd, 500, 2e5)
        life = rnd.choice([3, 5, 7, 10])
        ws.cell(r, 3, cost).number_format = AMOUNT
        ws.cell(r, 4, life)
        dep = round(cost / life, 2)
        total = 0
        for y in range(years):
            v = dep if y < life else 0
            total += v
            b.formula(ws, r, 5 + y, f"=IF({y + 1}<=$D{r},ROUND($C{r}/$D{r},2),0)", v, AMOUNT)
        last = get_column_letter(4 + years)
        b.formula(ws, r, 5 + years, f"=C{r}-SUM(E{r}:{last}{r})", round(cost - total, 2), AMOUNT)


def make(path, seed, scale=1.0):
    rnd = random.Random(seed)
    b = Book()
    s = lambda n: max(5, int(n * scale * rnd.uniform(0.5, 1.5)))  # noqa: E731
    general_ledger(b, rnd, s(3000))
    trial_balance(b, rnd, s(120))
    pnl_monthly(b, rnd, max(3, s(12)))
    ar_ageing(b, rnd, s(400))
    budget_vs_actual(b, rnd, s(150))
    cash_flow(b, rnd, rnd.randint(2, 5))
    fx_rates(b, rnd, s(500))
    payroll(b, rnd, s(300))
    fixed_assets(b, rnd, s(250))
    b.save(path)
    return path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("outdir")
    ap.add_argument("--seeds", default="1,2,3,4,5")
    ap.add_argument("--scale", type=float, default=1.0)
    a = ap.parse_args()
    os.makedirs(a.outdir, exist_ok=True)
    for seed in a.seeds.split(","):
        print(make(os.path.join(a.outdir, f"finance_{seed}.xlsx"), int(seed), a.scale))


if __name__ == "__main__":
    main()
