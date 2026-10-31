"""
Token benchmark of openrsxl.extended.llm: how many tokens an LLM reads with
and without the compression, for the *same* information (exact counts with
tiktoken, o200k_base).

    python tools/llm_bench.py book.xlsx [more.xlsx ...] [--types] [--json out.json]

For every selection (everything, values, formulas, styles, metadata) the
**baseline** is that information loaded as is, losslessly and without any
compression - what a program that dumps the workbook for a model produces:

* values      {"Sheet": {"A1": {"v": value, "t": data_type}, ...}} (JSON,
              every cell with a value; t keeps the type: 1 vs "1", TRUE vs
              "TRUE", #N/A vs "#N/A")
* formulas    {"Sheet": {"D2": "=B2*C2", ...}} (every formula cell)
* styles      the workbook's styles.xml + {"Sheet": {"A1": style id}} (every
              styled cell)
* metadata    the XML of the parts that hold it (the worksheets outside
              their cell data, comments, tables, workbook.xml, docProps)
* everything  all of it: one {"v", "t", "f", "s"} object per cell, styles.xml
              and the metadata XML

**ours** is ``compress_workbook(include=<selection>)`` - the complete
document a model gets, legend and headers included. ``saved`` = 1 - ours /
baseline. ``--types`` adds the same measurement per value type (a copy of
the workbook holding only the cells of one type, in their places).
"""

import collections
import json
import os
import re
import sys
import tempfile
import zipfile

import openpyxl
import tiktoken
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter

from openrsxl.extended.llm import compress_workbook

ENC = tiktoken.get_encoding("o200k_base")
CHUNK = 1 << 20


def count(text):
    """exact tokens; long texts in ~1 MB pieces cut at line ends"""
    if len(text) <= CHUNK:
        return len(ENC.encode_ordinary(text))
    parts, i = [], 0
    while i < len(text):
        j = min(len(text), i + CHUNK)
        if j < len(text):
            k = max(text.rfind("\n", i, j), text.rfind(">", i, j))
            j = k + 1 if k > i else j
        parts.append(text[i:j])
        i = j
    return sum(len(t) for t in ENC.encode_ordinary_batch(parts, num_threads=3))


class Counter:
    """counts the tokens of a text written piece by piece (bounded memory)"""

    def __init__(self):
        self.parts, self.size, self.tokens = [], 0, 0

    def write(self, s):
        self.parts.append(s)
        self.size += len(s)
        if self.size > 4 * CHUNK:
            self.flush()

    def flush(self):
        if self.parts:
            self.tokens += count("".join(self.parts))
            self.parts, self.size = [], 0
        return self.tokens


# -- baselines ---------------------------------------------------------------------
def _json_value(v):
    if isinstance(v, (bool, int, float, str)) or v is None:
        return v
    return v.isoformat() if hasattr(v, "isoformat") else str(v)


def _formula_text(v):
    if isinstance(v, str):
        return v
    text = getattr(v, "text", None)
    ref = getattr(v, "ref", None)
    return {"array": ref, "text": text} if text is not None else str(v)


SHEET_PART = re.compile(r"xl/worksheets/sheet[0-9]+\.xml$")
SHEET_DATA = re.compile(r"<(?:\w+:)?sheetData\b.*?</(?:\w+:)?sheetData>|<(?:\w+:)?sheetData\s*/>", re.S)


def metadata_xml(path):
    """the XML holding the metadata sections, as in the file"""
    out = Counter()
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            keep = (
                SHEET_PART.match(name)
                or re.match(r"xl/comments[0-9]*\.xml$", name)
                or name.startswith("xl/tables/")
                or name == "xl/workbook.xml"
                or (name.startswith("docProps/") and name.endswith(".xml"))
            )
            if not keep:
                continue
            data = z.read(name).decode("utf-8", "replace")
            if SHEET_PART.match(name):
                data = SHEET_DATA.sub("", data)
            out.write(data + "\n")
    return out.flush()


def styles_xml(path):
    with zipfile.ZipFile(path) as z:
        try:
            return count(z.read("xl/styles.xml").decode("utf-8", "replace"))
        except KeyError:
            return 0


def cell_baselines(path):
    """tokens of the per-cell JSON baselines: values, formulas, styles, everything"""
    out = {k: Counter() for k in ("values", "formulas", "styles", "everything")}
    vb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    fb = openpyxl.load_workbook(path, read_only=True)
    cells = collections.Counter()
    for k in out:
        out[k].write("{")
    for n, (vs, fs) in enumerate(zip(vb.worksheets, fb.worksheets)):
        vs.reset_dimensions()
        fs.reset_dimensions()
        head = ("" if n == 0 else ", ") + json.dumps(vs.title, ensure_ascii=False) + ": {"
        for k in out:
            out[k].write(head)
        first = dict.fromkeys(out, True)

        def put(kind, key, obj):
            out[kind].write(
                ("" if first[kind] else ", ") + json.dumps(key) + ": " + json.dumps(obj, ensure_ascii=False)
            )
            first[kind] = False

        for vrow, frow in zip(vs.iter_rows(), fs.iter_rows()):
            for vc, fc in zip(vrow, frow):
                if not hasattr(vc, "_style_id"):
                    continue  # padding (EmptyCell)
                ref = f"{get_column_letter(vc.column)}{vc.row}"
                every = {}
                if vc.value is not None:
                    v = {"v": _json_value(vc.value), "t": vc.data_type}
                    put("values", ref, v)
                    every.update(v)
                    cells["values"] += 1
                if fc.data_type == "f":
                    f = _formula_text(fc.value)
                    put("formulas", ref, f)
                    every["f"] = f
                    cells["formulas"] += 1
                if vc._style_id:
                    put("styles", ref, vc._style_id)
                    every["s"] = vc._style_id
                    cells["styles"] += 1
                if every:
                    put("everything", ref, every)
        for k in out:
            out[k].write("}")
    for k in out:
        out[k].write("}")
    vb.close()
    fb.close()
    return {k: c.flush() for k, c in out.items()}, cells


def baselines(path):
    cells, counts = cell_baselines(path)
    sx = styles_xml(path)
    meta = metadata_xml(path)
    base = {
        "everything": cells["everything"] + sx + meta,
        "values": cells["values"],
        "formulas": cells["formulas"],
        "styles": cells["styles"] + sx,
        "metadata": meta,
    }
    return base, counts


def ours(path, include):
    doc = compress_workbook(path, include=include, output="text", path=os.path.join(TMP, "bench"))
    n = Counter()
    with open(doc.files["text"], encoding="utf-8") as f:
        for line in f:
            n.write(line)
    return n.flush()


SELECTIONS = [
    ("everything", "all"),
    ("values", "values"),
    ("formulas", "formulas"),
    ("styles", "styles"),
    ("metadata", "metadata"),
]


def measure(path):
    base, counts = baselines(path)
    rows = []
    for name, include in SELECTIONS:
        b = base[name]
        o = ours(path, include) if b else 0
        rows.append({"selection": name, "cells": counts.get(name), "baseline": b, "ours": o})
    return rows


# -- per value type -------------------------------------------------------------
TYPES = ("integer", "decimal", "text", "boolean", "date/time", "error")


def _kind(cell):
    v = cell.value
    if v is None:
        return None
    if cell.data_type == "e":
        return "error"
    if isinstance(v, bool):
        return "boolean"
    if isinstance(v, int):
        return "integer"
    if isinstance(v, float):
        return "decimal"
    if isinstance(v, str):
        return "text"
    return "date/time"


def split_by_type(path, tmp):
    """one workbook per value type, holding only those cells in their places"""
    src = openpyxl.load_workbook(path, read_only=True, data_only=True)
    outs = {t: openpyxl.Workbook(write_only=True) for t in TYPES}
    counts = collections.Counter()
    for ws in src.worksheets:
        ws.reset_dimensions()
        dst = {t: outs[t].create_sheet(ws.title) for t in TYPES}
        last = dict.fromkeys(TYPES, 0)
        for row in ws.iter_rows():
            per = collections.defaultdict(list)
            for c in row:
                k = _kind(c)
                if k is not None:
                    per[k].append(c)
            for k, cs in per.items():
                r = cs[0].row
                while last[k] < r - 1:
                    dst[k].append([])
                    last[k] += 1
                line = [None] * cs[-1].column
                for c in cs:
                    oc = WriteOnlyCell(dst[k], c.value)
                    if k == "text":
                        oc.data_type = "s"
                    elif k == "date/time":
                        oc.number_format = c.number_format
                    line[c.column - 1] = oc
                    counts[k] += 1
                dst[k].append(line)
                last[k] = r
    src.close()
    files = {}
    for t, wb in outs.items():
        if counts[t]:
            files[t] = os.path.join(tmp, t.replace("/", "_") + ".xlsx")
            wb.save(files[t])
    return files, counts


def measure_types(path):
    with tempfile.TemporaryDirectory() as tmp:
        files, counts = split_by_type(path, tmp)
        rows = []
        for t, p in files.items():
            cells, _ = cell_baselines(p)
            rows.append({"type": t, "cells": counts[t], "baseline": cells["values"], "ours": ours(p, "values")})
        return rows


def pct(b, o):
    return f"{100 * (1 - o / b):5.1f}%" if b else "    -"


def main(argv):
    global TMP
    want_types = "--types" in argv
    out_json = argv[argv.index("--json") + 1] if "--json" in argv else None
    files = [a for a in argv if not a.startswith("--") and a != out_json]
    report = {}
    with tempfile.TemporaryDirectory() as TMP:
        for path in files:
            rows = measure(path)
            print(f"\n== {os.path.basename(path)}   (o200k_base tokens; baseline = the same information uncompressed)")
            print(f"{'loaded':12s} {'cells':>11s} {'baseline':>14s} {'ours':>12s} {'saved':>7s}")
            for r in rows:
                cells = f"{r['cells']:,}" if r["cells"] is not None else ""
                print(
                    f"{r['selection']:12s} {cells:>11s} {r['baseline']:>14,} {r['ours']:>12,} {pct(r['baseline'], r['ours']):>7s}"
                )
            report[path] = {"selections": rows}
            if want_types:
                trows = measure_types(path)
                print(f"\n{'value type':12s} {'cells':>11s} {'baseline':>14s} {'ours':>12s} {'saved':>7s}")
                for r in trows:
                    print(
                        f"{r['type']:12s} {r['cells']:>11,} {r['baseline']:>14,} {r['ours']:>12,} {pct(r['baseline'], r['ours']):>7s}"
                    )
                report[path]["types"] = trows
    if out_json:
        with open(out_json, "w") as f:
            json.dump(report, f, indent=1)


TMP = None

if __name__ == "__main__":
    main(sys.argv[1:])
