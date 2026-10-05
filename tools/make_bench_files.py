"""
Generate benchmark / verification workbooks with different profiles, so that
performance and fidelity are not measured on a single file only.

    python tools/make_bench_files.py <outdir> [--rows N]

Files are written with openpyxl (the oracle) in write-only mode.
"""

import argparse
import datetime
import os
import random

import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def numbers(ws, rows, rnd):
    for r in range(rows):
        ws.append(
            [
                rnd.random() * 1e6,
                rnd.randint(-(10**9), 10**9),
                r,
                rnd.random(),
                rnd.randint(0, 100),
                rnd.uniform(-1, 1),
                r * 0.5,
                rnd.random() < 0.5,
                rnd.randint(0, 10**15),
                rnd.gauss(0, 1),
            ]
        )


def strings(ws, rows, rnd):
    words = ["alpha", "beta", "gamma", "delta", " padded ", "naïve", "日本語", "x&y<z>", "line\nbreak", ""]
    for r in range(rows):
        ws.append(
            [
                f"row {r}",
                rnd.choice(words),
                rnd.choice(words) * 3,
                f"id-{rnd.randint(0, 10**6)}",
                rnd.choice(words),
                "=" if r % 997 == 0 else "text",
                "#N/A" if r % 991 == 0 else "ok",
                str(r),
                rnd.choice(words),
                f"{rnd.random():.6f}",
            ]
        )


def dates(ws, rows, rnd):
    base = datetime.datetime(2000, 1, 1)
    for r in range(rows):
        d = base + datetime.timedelta(days=rnd.randint(0, 20000), seconds=rnd.randint(0, 86399))
        ws.append(
            [
                d,
                d.date(),
                d.time(),
                datetime.timedelta(hours=rnd.randint(0, 1000), seconds=rnd.randint(0, 59)),
                d + datetime.timedelta(microseconds=rnd.randint(0, 999999)),
                r,
                d.replace(hour=0, minute=0, second=0),
            ]
        )


def sparse(ws, rows, rnd):
    for r in range(rows):
        row = [None] * 200
        for _ in range(3):
            row[rnd.randrange(200)] = rnd.random()
        ws.append(row)


def styled(ws, rows, rnd):
    fonts = [Font(bold=True), Font(italic=True, color="FF0000FF"), Font(name="Arial", size=9)]
    fills = [PatternFill("solid", fgColor=c) for c in ("FFFFFF00", "FF00FF00", "FFDDDDDD")]
    border = Border(left=Side("thin"), bottom=Side("double"))
    for r in range(rows):
        cells = []
        for c in range(8):
            cell = WriteOnlyCell(ws, value=rnd.random() if c % 2 else f"s{r}")
            if (r + c) % 3 == 0:
                cell.font = fonts[(r + c) % 3]
            if (r + c) % 4 == 0:
                cell.fill = fills[c % 3]
            if c == 7:
                cell.border = border
                cell.alignment = Alignment(horizontal="center", wrap_text=True)
                cell.number_format = "0.00%"
            cells.append(cell)
        ws.append(cells)


def formulas(ws, rows, rnd):
    for r in range(1, rows + 1):
        ws.append(
            [
                rnd.random(),
                rnd.random(),
                f"=A{r}+B{r}",
                f"=SUM($A$1:A{r})",
                f'=IF(A{r}>0.5,"hi","lo")',
                f"=A{r}*{rnd.randint(1, 10**6)}",  # unique constants: no template sharing
                f"=VLOOKUP(A{r},$A$1:$B${rows},2,FALSE)",
                f"=Sheet1!A{r}",
            ]
        )


PROFILES = {
    "numbers": numbers,
    "strings": strings,
    "dates": dates,
    "sparse": sparse,
    "styled": styled,
    "formulas": formulas,
}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("outdir")
    ap.add_argument("--rows", type=int, default=100_000)
    a = ap.parse_args()
    os.makedirs(a.outdir, exist_ok=True)
    for name, fill in PROFILES.items():
        rnd = random.Random(42)
        wb = openpyxl.Workbook(write_only=True)
        ws = wb.create_sheet("Sheet1")
        fill(ws, a.rows, rnd)
        path = os.path.join(a.outdir, f"bench_{name}.xlsx")
        wb.save(path)
        print(path, os.path.getsize(path))


if __name__ == "__main__":
    main()
