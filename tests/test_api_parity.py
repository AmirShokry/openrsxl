"""
API parity of full mode (and read-only mode) beyond cell values: dict views
and errors of ``ws._cells``, unusual value types written into loaded and new
workbooks, cell introspection, worksheet / workbook attributes, error paths,
resources released after a failed save, and pandas running on openrsxl
(``install_as_openpyxl``). Every check runs the same code with openpyxl and
with openrsxl and compares the results.
"""

import collections.abc
import datetime
import decimal
import enum
import fractions
import gc
import glob
import io
import os
import subprocess
import sys
import tempfile
import warnings

import openpyxl
import pytest
from oracle.compare import ADDR_RE, norm, norm_text, saved_parts

import openrsxl

try:
    import numpy as np
except ImportError:  # pragma: no cover
    np = None


def make_source():
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.worksheet.formula import ArrayFormula

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["a", 1, 2.5, True, None, "=A1", datetime.datetime(2020, 1, 1)])
    ws.append(["  x  ", -3, 1e20, False, "", "#N/A", datetime.time(1, 2)])
    ws["A4"] = CellRichText(["plain ", TextBlock(InlineFont(b=True), "bold"), " end"])
    ws["B4"] = ArrayFormula("B4:B5", "=A1:A2")
    ws["C5"] = "z"
    ws["J9"] = 7
    ws["D7"] = datetime.timedelta(hours=5)
    ws["E7"] = datetime.date(2001, 2, 3)
    ws["A1"].font = openpyxl.styles.Font(bold=True)
    ws.merge_cells("F8:G9")
    wb.create_sheet("Empty")
    gap = wb.create_sheet("Gap")
    gap["C3"] = 1
    gap["A10"] = 2
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


SOURCE = make_source()


def outcome(f):
    """Result (normalised) or exception, plus the warnings emitted."""
    with warnings.catch_warnings(record=True) as w:
        warnings.simplefilter("always")
        try:
            r = ("ok", norm(f()))
        except Exception as e:
            r = ("error", type(e).__name__, ADDR_RE.sub("", norm_text(str(e))))
    return r, [(x.category.__name__, norm_text(str(x.message))) for x in w if x.category is not ResourceWarning]


def both(fn):
    return [fn(lib) for lib in (openpyxl, openrsxl)]


def load(lib, **kw):
    return lib.load_workbook(io.BytesIO(SOURCE), **kw)


# -- ws._cells: dict views and key errors ------------------------------------


def views(cells):
    k, v, i = cells.keys(), cells.values(), cells.items()
    first = next(iter(cells))
    return [
        isinstance(k, collections.abc.KeysView),
        isinstance(v, collections.abc.ValuesView),
        isinstance(i, collections.abc.ItemsView),
        list(reversed(k)),
        [c.coordinate for c in reversed(v)],
        [(kk, c.coordinate) for kk, c in reversed(i)],
        sorted(k & {first, (99, 99)}),
        sorted(k | {(99, 99)}),
        sorted(k - {first}),
        sorted(k ^ {first, (99, 99)}),
        k.isdisjoint({first}),
        first in k,
        (first, cells[first]) in i,
        (first, None) in i,
        cells[first] in v,
        len(k) == len(v) == len(i),
        k == set(cells),
        dict(k.mapping) == dict(cells),
        type(k.mapping).__name__,
        repr(k),
        norm_text(repr(v)).count("<Cell"),
        norm_text(repr(i)).count("<Cell"),
        repr(k).split("(")[0],
        repr(v).split("(")[0],
        repr(i).split("(")[0],
    ]


@pytest.mark.parametrize("sheet", ["Sheet", "Gap"])
def test_cells_views_behave_like_dict_views(sheet):
    a, b = both(lambda lib: views(load(lib)[sheet]._cells))
    assert a == b


def test_cells_views_of_empty_and_modified_store():
    def run(lib):
        ws = load(lib)["Gap"]
        cells = ws._cells
        out = []
        ws["B1"] = 5
        del ws._cells[(3, 3)]
        out.append(views(cells))
        cells.clear()
        out.append([list(reversed(cells.keys())), repr(cells.keys()), repr(cells.items()), len(cells.values())])
        return out

    a, b = both(run)
    assert a == b


def test_cells_keys_after_cell_access():
    """
    Keys iterated after cells were accessed in different ways (openrsxl
    shares the int objects of the key tuples it builds and keeps the tuples
    used to look up cells) are equal to openpyxl's and behave as dict keys
    """

    def run(lib):
        ws = load(lib)["Sheet"]
        cells = ws._cells
        out = []
        for row in ws.iter_rows():
            for c in row:
                out.append(c.coordinate)
        ws.cell(20, 3).value = 1
        out.append(ws["B5"].coordinate)
        first, second = list(cells), list(cells)
        out.append(first == second)
        out.append(all(type(r) is int and type(c) is int for r, c in first))
        out.append(list(reversed(cells)) == first[::-1])
        out.append(sorted(cells.items()) == sorted(zip(first, cells.values())))
        key, cell = cells.popitem()
        out.append((key, cell.coordinate, key in cells, len(cells)))
        cells[key] = cell
        ws.insert_rows(2)
        out.append(sorted(cells) == sorted((c.row, c.column) for c in cells.values()))
        out.append(sorted(cells))
        return out

    a, b = both(run)
    assert a == b


def test_cells_keep_the_key_used_to_look_them_up():
    ws = load(openrsxl)["Sheet"]
    cells = ws._cells
    keys = [(r, c) for r, c in list(cells)]
    for k in keys:
        cells[k]
    merged = {k for r in ws.merged_cells.ranges for k in r.cells}
    assert all(a is b for a, b in zip(list(cells), keys) if a not in merged)
    # replacing a cell object keeps the key
    cells[keys[0]] = cells[keys[0]]
    assert next(iter(cells)) is keys[0]


KEY_OPS = {
    "in": lambda d, k: k in d,
    "get": lambda d, k: d.get(k),
    "get_default": lambda d, k: d.get(k, 1),
    "getitem": lambda d, k: d[k],
    "pop": lambda d, k: d.pop(k),
    "pop_default": lambda d, k: d.pop(k, "dflt"),
    "setdefault": lambda d, k: d.setdefault(k),
    "setitem": lambda d, k: d.__setitem__(k, 1),
    "delitem": lambda d, k: d.__delitem__(k),
    "keys_in": lambda d, k: k in d.keys(),
    "items_in": lambda d, k: (k, None) in d.items(),
    "keys_and": lambda d, k: d.keys() & [k],
    "update": lambda d, k: d.update([(k, 1)]),
}


class Unhashable:
    __hash__ = None


class BadHash:
    def __hash__(self):
        raise ValueError("no hash for you")


@pytest.mark.parametrize("op", sorted(KEY_OPS))
@pytest.mark.parametrize("empty", [False, True])
@pytest.mark.parametrize(
    "key", [[1, 2], (1, [2]), {1}, Unhashable(), BadHash(), (1, 1), (1.0, True), (99, 99), "A1", None]
)
def test_cells_key_errors_like_dict(op, empty, key):
    def run(lib):
        cells = load(lib).active._cells
        if empty:
            cells.clear()
        return outcome(lambda: KEY_OPS[op](cells, key))

    a, b = both(run)
    assert a == b
    # errors are those of a plain dict (the message depends on the Python
    # version)
    ref = outcome(lambda: KEY_OPS[op]({} if empty else {(1, 1): None}, key))[0]
    if ref[0] == "error":
        assert b[0] == ref


# -- unusual value types -----------------------------------------------------


class MyInt(int):
    pass


class MyFloat(float):
    def __repr__(self):
        return "MYFLOAT"

    def __str__(self):
        return "MYSTR"


class MyStr(str):
    pass


class Color(enum.IntEnum):
    RED = 3


VALUES = {
    "myint": lambda: MyInt(4),
    "myfloat": lambda: MyFloat(1.5),
    "mystr_formula": lambda: MyStr("=1+1"),
    "intenum": lambda: Color.RED,
    "fraction": lambda: fractions.Fraction(1, 3),
    "decimal_exp": lambda: decimal.Decimal("1E+5"),
    "decimal_nan": lambda: decimal.Decimal("NaN"),
    "bigint": lambda: 10**30,
    "negzero": lambda: -0.0,
    "subnormal": lambda: 5e-324,
    "datetime_tz": lambda: datetime.datetime(2020, 1, 1, tzinfo=datetime.timezone.utc),
    "date_min": lambda: datetime.date(1, 1, 1),
    "datetime_max": lambda: datetime.datetime.max,
    "time_us": lambda: datetime.time(23, 59, 59, 999999),
    "timedelta_big": lambda: datetime.timedelta(days=99999),
    "bytes": lambda: b"abc",
    "bytearray": lambda: bytearray(b"abc"),
    "list": lambda: [1],
    "str_ctrl": lambda: "a\x00b",
    "str_crlf": lambda: "a\r\nb",
    "str_escape": lambda: "_x000D_",
    "str_long": lambda: "y" * 40000,
    "str_eq": lambda: "=",
    "str_error": lambda: "#REF!",
}
if np is not None:
    VALUES.update(
        {
            "np_int64": lambda: np.int64(5),
            "np_int8": lambda: np.int8(-5),
            "np_uint64": lambda: np.uint64(2**63 + 1),
            "np_float64": lambda: np.float64(0.1),
            "np_float32": lambda: np.float32(0.1),
            "np_float16": lambda: np.float16(0.1),
            "np_bool": lambda: np.bool_(True),
            "np_nan": lambda: np.float64("nan"),
            "np_datetime64": lambda: np.datetime64("2020-01-01"),
            "np_str": lambda: np.str_("hey"),
        }
    )


@pytest.mark.parametrize("name", sorted(VALUES))
def test_unusual_value_in_loaded_sheet(name):
    def run(lib):
        def go():
            ws = load(lib).active
            c = ws["B2"]
            c.value = VALUES[name]()
            r = [type(c.value).__name__, repr(c.value)[:60], c.data_type, c.number_format]
            c2 = ws.cell(20, 2, VALUES[name]())
            r += [type(c2.value).__name__, c2.data_type]
            ws.append([VALUES[name]()])
            return r, saved_parts(ws.parent)

        return outcome(go)

    a, b = both(run)
    assert a == b


@pytest.mark.parametrize("name", sorted(VALUES))
def test_unusual_value_in_new_workbook(name):
    def run(lib):
        def go():
            wb = lib.Workbook()
            wb.active["A1"] = VALUES[name]()
            return saved_parts(wb)

        return outcome(go)

    a, b = both(run)
    assert a == b


# -- a failed save -----------------------------------------------------------


def test_failed_save_releases_the_temporary_file():
    """
    A value that cannot be written (tz-aware datetime) makes save() raise;
    openpyxl closes the temporary worksheet file when the writer is
    discarded, so does openrsxl (no ResourceWarning, the file can be removed).
    """

    def run(lib):
        tmp = tempfile.gettempdir()
        before = set(glob.glob(os.path.join(tmp, "openpyxl.*")))
        with warnings.catch_warnings(record=True) as w:
            warnings.simplefilter("always")
            wb = load(lib)
            wb.active["A1"] = datetime.datetime(2020, 1, 1, tzinfo=datetime.timezone.utc)
            with pytest.raises(TypeError) as e:
                wb.save(io.BytesIO())
            del e
            gc.collect()
        left = set(glob.glob(os.path.join(tmp, "openpyxl.*"))) - before
        for f in left:
            os.remove(f)  # fails on Windows if the file is still open
        return len(left), [str(x.message) for x in w if x.category is ResourceWarning]

    a, b = both(run)
    assert b[1] == []
    if getattr(sys, "_is_gil_enabled", lambda: True)():
        assert a == b
    else:
        # free-threaded Python: openpyxl's own writer is sometimes finalized
        # with a ResourceWarning (timing dependent); openrsxl's must not be
        assert a[0] == b[0]


@pytest.mark.parametrize("bad", [None, "a\x1cb", datetime.datetime(2020, 1, 1, tzinfo=datetime.timezone.utc)])
def test_discarded_write_only_workbook_releases_the_temporary_file(bad):
    """
    A write-only worksheet streams into a temporary file from its first
    append; when the workbook is discarded without being saved (eg. after
    an append that raised), the garbage collector closes it in both libraries
    (openrsxl's native row writer takes part in the worksheet <-> writer
    reference cycle).
    """

    def run(lib):
        tmp = tempfile.gettempdir()
        before = set(glob.glob(os.path.join(tmp, "openpyxl.*")))
        with warnings.catch_warnings(record=True) as w:
            warnings.simplefilter("always")
            wb = lib.Workbook(write_only=True)
            ws = wb.create_sheet("x")
            ws.append([1, "a", 2.5])
            err = None
            if bad is not None:
                try:
                    ws.append([bad])
                except Exception as e:
                    err = type(e).__name__
            del wb, ws
            gc.collect()
        left = set(glob.glob(os.path.join(tmp, "openpyxl.*"))) - before
        for f in left:
            os.remove(f)  # fails on Windows if the file is still open
        return err, len(left), [str(x.message) for x in w if x.category is ResourceWarning]

    a, b = both(run)
    assert a == b
    assert b[2] == []


# -- introspection -----------------------------------------------------------


def test_cell_introspection():
    def run(lib):
        ws = load(lib).active
        c = ws["A1"]
        return [
            type(c).__qualname__,
            len(type(c).__mro__),
            hasattr(c, "__dict__"),
            c is ws["A1"],
            c is ws.cell(1, 1),
            c.parent is ws,
            c.internal_value,
            c.is_date,
            c.col_idx,
            c.column_letter,
            c.encoding,
            c.base_date,
            c.offset(2, 3).coordinate,
            norm_text(repr(c)).split(" ")[0],
            [ws[k].is_date for k in ("G1", "G2", "D7", "E7")],
            [ws[k].number_format for k in ("G1", "G2", "D7", "E7")],
            [(ws[k].data_type, norm(ws[k].value)) for k in ("F1", "F2", "E2", "B4", "A4", "F8", "G9")],
            [type(ws[k]).__name__ for k in ("F8", "G9")],
            sorted(a for a in dir(c) if not a.startswith("__")),
            sorted(a for a in dir(ws["G9"]) if not a.startswith("__")),
        ]

    a, b = both(run)
    assert a == b


def test_worksheet_and_workbook_attributes():
    def run(lib):
        wb = load(lib)
        ws = wb.active
        return [
            ws.max_row,
            ws.max_column,
            ws.min_row,
            ws.min_column,
            ws.dimensions,
            ws.calculate_dimension(),
            type(ws.iter_rows()).__name__,
            type(ws.rows).__name__,
            type(ws.values).__name__,
            type(ws.columns).__name__,
            ws._current_row,
            len(ws._cells),
            [type(x).__name__ for x in ws["A1:B2"]],
            [type(ws[k]).__name__ for k in ("A", 1, "A:B", "1:2")],
            len(ws["A:B"]),
            len(ws["1:2"][0]),
            norm(ws.array_formulae),
            str(ws.merged_cells),
            sorted(ws.row_dimensions),
            sorted(ws.column_dimensions),
            [wb.read_only, wb.data_only, wb.write_only, wb.template, wb.is_template, wb.iso_dates],
            [wb.epoch, wb.excel_base_date, wb.encoding, wb.mime_type, wb.path, wb.code_name],
            [wb.vba_archive, len(wb._cell_styles), len(wb._fonts), len(wb._fills)],
            [type(x).__name__ for x in wb._sheets],
            [(s.title, s.max_row, s.max_column, s.dimensions) for s in wb],
        ]

    a, b = both(run)
    assert a == b


def test_iteration_edge_cases():
    def run(lib):
        ws = load(lib).active
        steps = [
            lambda: list(ws.iter_rows(min_row=0, max_row=1, values_only=True)),
            lambda: list(ws.iter_rows(min_row=3, max_row=2, values_only=True)),
            lambda: list(ws.iter_rows(min_col=3, max_col=1, values_only=True)),
            lambda: list(ws.iter_rows(min_row=8, values_only=True)),
            lambda: list(ws.iter_cols(min_col=9, values_only=True)),
            lambda: [c.coordinate for row in ws.iter_rows(min_row=4, max_row=5, max_col=2) for c in row],
            lambda: len(ws._cells),
            lambda: [[c.coordinate for c in row] for row in ws["B3:A1"]],
            lambda: ws["Z100"].value,
            lambda: (len(ws._cells), ws.max_row, ws.max_column, ws._current_row),
            lambda: list(ws.iter_rows(min_row=True, max_row=2.0, max_col=2, values_only=True)),
            lambda: list(ws.iter_rows(min_row=-1, max_row=1, values_only=True)),
        ]
        return [outcome(f) for f in steps]

    a, b = both(run)
    assert a == b


def test_iteration_interleaved_with_changes():
    def run(lib):
        ws = load(lib).active
        g = ws.iter_rows(values_only=True)
        first = next(g)
        ws["A2"] = "changed"
        ws.insert_rows(1)
        rest = list(g)
        out = [first, rest]
        for k in ws._cells:
            ws._cells[k] = ws._cells[k]
        it = iter(ws._cells)
        next(it)
        del ws._cells[(2, 2)]
        out.append(outcome(lambda: next(it)))
        return norm(out)

    a, b = both(run)
    assert a == b


ERROR_STEPS = [
    lambda wb, ws: ws.cell(0, 1),
    lambda wb, ws: ws.cell(1, 0),
    lambda wb, ws: ws.cell(-1, 1),
    lambda wb, ws: ws.cell("1", 1),
    lambda wb, ws: ws[""],
    lambda wb, ws: ws["A0"],
    lambda wb, ws: ws[0],
    lambda wb, ws: ws["XFE1"],
    lambda wb, ws: ws["A1048577"],
    lambda wb, ws: ws[1.5],
    lambda wb, ws: ws["1:0"],
    lambda wb, ws: ws["A:"],
    lambda wb, ws: ws[None],
    lambda wb, ws: ws.cell(1, 1, value=None).value,
    lambda wb, ws: ws.cell(1, 1, value=0).value,
    lambda wb, ws: ws.cell(1048577, 1),
    lambda wb, ws: ws.cell(row=2.0, column=1).coordinate,
    lambda wb, ws: ws.cell(True, True).coordinate,
    lambda wb, ws: next(ws.iter_rows(min_row=-1)),
    lambda wb, ws: wb["nope"],
    lambda wb, ws: ws.merge_cells("A1"),
    lambda wb, ws: ws.unmerge_cells("A1:B2"),
    lambda wb, ws: ws.move_range("A1:B2", rows=-5),
    lambda wb, ws: ws.insert_rows(0),
    lambda wb, ws: ws.delete_rows(-2, 3),
    lambda wb, ws: ws.delete_cols(100),
    lambda wb, ws: ws.append("abc"),
    lambda wb, ws: ws.append(5),
    lambda wb, ws: ws.append({"A": 1, 3: 2, "zz": 3}),
    lambda wb, ws: ws.append({1.5: 1}),
    lambda wb, ws: ws.append(iter([1, 2])),
    lambda wb, ws: setattr(ws["A1"], "value", {}),
    lambda wb, ws: setattr(ws["A1"], "data_type", "q"),
    lambda wb, ws: setattr(ws["A1"], "hyperlink", None),
    lambda wb, ws: setattr(ws["G9"], "value", 1),
    lambda wb, ws: ws["F8"].offset(-10, 0),
    lambda wb, ws: ws["A1"].check_error("#BAD"),
    lambda wb, ws: ws["A1"].check_string(b"\xff"),
    lambda wb, ws: ws["A1"].check_string("x" * 33000)[-3:],
    lambda wb, ws: setattr(ws["A1"], "style", "nostyle"),
    lambda wb, ws: wb.remove(ws),
    lambda wb, ws: wb.active.title,
    lambda wb, ws: setattr(wb, "active", 5),
    lambda wb, ws: wb.create_sheet("Gap").title,
    lambda wb, ws: wb.create_sheet("a/b"),
    lambda wb, ws: wb.copy_worksheet(wb["Gap"]).title,
    lambda wb, ws: saved_parts(wb),
]


def test_error_paths():
    def run(lib):
        wb = load(lib)
        ws = wb.active
        return [outcome(lambda f=f: f(wb, ws)) for f in ERROR_STEPS]

    a, b = both(run)
    for i, (x, y) in enumerate(zip(a, b)):
        assert x == y, i


# -- read-only mode ----------------------------------------------------------


def cellrow(c):
    return (
        type(c).__name__,
        getattr(c, "coordinate", None) if type(c).__name__ != "EmptyCell" else None,
        norm(c.value),
        c.data_type,
        getattr(c, "number_format", None),
        getattr(c, "is_date", None),
        getattr(c, "row", None),
        getattr(c, "column", None),
        getattr(c, "style_id", None),
        getattr(c, "has_style", None),
        norm(getattr(c, "font", None)),
    )


@pytest.mark.parametrize("read_only", [False, True])
@pytest.mark.parametrize("data_only", [False, True])
@pytest.mark.parametrize("rich_text", [False, True])
def test_cells_every_mode(read_only, data_only, rich_text):
    def run(lib):
        wb = load(lib, read_only=read_only, data_only=data_only, rich_text=rich_text)
        r = []
        for ws in wb.worksheets:
            r.append((type(ws).__name__, ws.title, ws.max_row, ws.max_column, ws.min_row, ws.min_column))
            r.append(ws.calculate_dimension())
            r.append([[cellrow(c) for c in row] for row in ws.iter_rows()])
            r.append(norm(list(ws.values)))
            r.append(norm(list(ws.iter_rows(min_row=2, max_row=12, min_col=2, max_col=11, values_only=True))))
            r.append([[cellrow(c) for c in row] for row in ws.iter_rows(min_row=8, max_row=12, max_col=12)])
        return r

    a, b = both(run)
    assert a == b


@pytest.mark.parametrize("data_only", [False, True])
def test_read_only_worksheet_operations(data_only):
    def run(lib):
        wb = load(lib, read_only=True, data_only=data_only)
        ws = wb.active
        steps = [
            lambda: ws.max_row,
            lambda: ws.reset_dimensions(),
            lambda: (ws.max_row, ws.max_column),
            lambda: ws.calculate_dimension(force=True),
            lambda: (ws.max_row, ws.max_column),
            lambda: len(list(ws.rows)),
            lambda: len(list(ws.rows)),
            lambda: ws["C5"].value,
            lambda: [c.value for c in ws["A1:B2"][0]],
            lambda: ws.cell(2, 2).value,
            lambda: ws.calculate_dimension(),
            lambda: (lambda e: (e.max_row, e.max_column, list(e.rows), e.calculate_dimension()))(wb["Empty"]),
            lambda: wb.close(),
            lambda: len(list(ws.rows)),
        ]
        return [outcome(f) for f in steps]

    a, b = both(run)
    assert a == b


def test_rich_text_values():
    def run(lib):
        out = []
        for ro in (False, True):
            ws = load(lib, read_only=ro, rich_text=True).active
            row = next(ws.iter_rows(min_row=4, max_row=4))
            v, f = row[0].value, row[1].value
            out.append([type(v).__name__, str(v), norm_text(repr(v)), [(type(x).__name__, str(x)) for x in v]])
            out.append([type(f).__name__, f.text, f.ref])
        return out

    a, b = both(run)
    assert a == b


# -- pandas on openrsxl ------------------------------------------------------

PANDAS_SCRIPT = r"""
import hashlib, io, json, sys, warnings
lib, path = sys.argv[1:3]
if lib == "openrsxl":
    import openrsxl.extended
    openrsxl.extended.install_as_openpyxl()
import openpyxl
import pandas as pd
assert openpyxl.__spec__.name == lib
warnings.simplefilter("ignore")
out = {}
for header in (None, 0):
    sheets = pd.read_excel(path, sheet_name=None, header=header, engine="openpyxl")
    for name, df in sheets.items():
        cells = [[(type(v).__name__, repr(v)) for v in row] for row in df.itertuples(index=False)]
        b = io.BytesIO()
        df.to_excel(b, engine="openpyxl", index=False)
        back = pd.read_excel(io.BytesIO(b.getvalue()), engine="openpyxl")
        out[f"{header}:{name}"] = [list(map(str, df.dtypes)), list(df.shape), cells,
                                   repr(back.to_dict("split"))]
print(json.dumps(out))
"""


def test_pandas_on_openrsxl(tmp_path):
    pytest.importorskip("pandas")
    path = tmp_path / "f.xlsx"
    path.write_bytes(SOURCE)
    script = tmp_path / "run_pandas.py"
    script.write_text(PANDAS_SCRIPT)
    res = [
        subprocess.run([sys.executable, str(script), lib, str(path)], capture_output=True, text=True, check=True).stdout
        for lib in ("openpyxl", "openrsxl")
    ]
    assert res[0] == res[1]
    assert res[0].strip()
