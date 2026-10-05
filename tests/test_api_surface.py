"""
API surface conformance: every module, class, function, method, property,
constant and signature of openpyxl must exist identically in openrsxl (and
in openrsxl.extended).
"""

import importlib
import inspect
import pkgutil
import re

import openpyxl
import pytest

import openrsxl

ADDR = re.compile(r" at 0x[0-9A-Fa-f]+")


def norm(s):
    return ADDR.sub("", re.sub(r"\bopenrsxl\b", "openpyxl", str(s)))


def openpyxl_modules():
    names = ["openpyxl"]
    for m in pkgutil.walk_packages(openpyxl.__path__, "openpyxl."):
        if ".tests" in m.name:
            continue
        names.append(m.name)
    return sorted(names)


MODULES = openpyxl_modules()


def describe(obj):
    if inspect.isclass(obj):
        return "class"
    if inspect.isfunction(obj) or inspect.isbuiltin(obj):
        return "function"
    if inspect.ismodule(obj):
        return "module"
    return "value:" + type(obj).__name__


def signature(obj):
    try:
        return norm(inspect.signature(obj))
    except (TypeError, ValueError):
        return None


def extends(x, y):
    """y's signature is x's plus keyword-only parameters with defaults"""
    try:
        sx, sy = inspect.signature(x), inspect.signature(y)
    except (TypeError, ValueError):
        return False
    px, py = list(sx.parameters.values()), list(sy.parameters.values())
    if norm(str(px)) != norm(str(py[: len(px)])):
        return False
    return all(p.kind is p.KEYWORD_ONLY and p.default is not p.empty for p in py[len(px) :])


def counterpart(name, prefix):
    return prefix + name[len("openpyxl") :]


@pytest.mark.parametrize("prefix", ["openrsxl", "openrsxl.extended"])
@pytest.mark.parametrize("modname", MODULES)
def test_module(modname, prefix):
    try:
        a = importlib.import_module(modname)
    except ImportError as err:
        # an optional dependency of openpyxl (eg. numpy for utils.dataframe):
        # openrsxl's module must fail the same way
        with pytest.raises(type(err)):
            importlib.import_module(counterpart(modname, prefix))
        pytest.skip(f"{modname} needs {err.name}")
    b = importlib.import_module(counterpart(modname, prefix))

    def public(mod):
        # submodules appear as attributes once imported anywhere (import order
        # dependent); they are checked as modules of their own
        return {
            n
            for n in dir(mod)
            if not n.startswith("__")
            and not (
                inspect.ismodule(getattr(mod, n))
                and getattr(getattr(mod, n), "__name__", "").startswith(mod.__name__ + ".")
            )
        }

    names_a = public(a)
    names_b = public(b)
    missing = names_a - names_b
    assert not missing, f"{modname}: missing {sorted(missing)}"
    if prefix == "openrsxl":
        extra = {n for n in names_b - names_a if not n.startswith("_")}
        assert not extra, f"{modname}: unexpected public names {sorted(extra)}"
    for n in sorted(names_a):
        x, y = getattr(a, n), getattr(b, n)
        assert describe(x) == describe(y), f"{modname}.{n}"
        if inspect.ismodule(x):
            assert norm(x.__name__) == norm(y.__name__)
            continue
        if callable(x):
            # openrsxl.extended may add keyword-only options
            assert signature(x) == signature(y) or (
                prefix == "openrsxl.extended" and extends(x, y)
            ), f"{modname}.{n} signature"
        if inspect.isclass(x):
            check_class(f"{modname}.{n}", x, y)
        elif not callable(x) and not inspect.ismodule(x):
            if isinstance(x, (int, float, str, bytes, tuple, frozenset, bool, type(None))):
                assert norm(repr(x)) == norm(repr(y)), f"{modname}.{n} value"


def check_class(path, x, y):
    assert norm(x.__qualname__) == norm(y.__qualname__), path
    assert [norm(f"{c.__module__}.{c.__qualname__}") for c in x.__mro__] == [
        norm(f"{c.__module__}.{c.__qualname__}") for c in y.__mro__
    ], f"{path} mro"
    for attr in dir(x):
        if attr in (
            "__dict__",
            "__weakref__",
            "__doc__",
            "__module__",
            "__init_subclass__",
            "__subclasshook__",
            "__annotations__",
            "__firstlineno__",
            "__static_attributes__",
        ):
            continue
        missing = object()
        ax = inspect.getattr_static(x, attr, missing)
        ay = inspect.getattr_static(y, attr, missing)
        assert ay is not missing, f"{path}.{attr} missing"
        assert type(ax).__name__ == type(ay).__name__, f"{path}.{attr} kind"
        if inspect.isfunction(ax) or isinstance(ax, (classmethod, staticmethod)):
            assert signature(getattr(x, attr)) == signature(getattr(y, attr)), f"{path}.{attr} signature"
        if isinstance(ax, (int, float, str, bytes, tuple, frozenset, bool, type(None))):
            assert norm(repr(ax)) == norm(repr(ay)), f"{path}.{attr} value"


def test_version_and_constants():
    for k in ("__version__", "__license__", "__author__"):
        assert getattr(openpyxl, k) == getattr(openrsxl, k)


def test_extended_is_alias():
    import openrsxl.extended as ex
    import openrsxl.styles
    from openrsxl.extended.styles import Font

    assert Font is openrsxl.styles.Font
    # load_workbook is extended (streaming features), the rest is aliased
    assert ex.load_workbook is not openrsxl.load_workbook
    assert ex.Workbook is openrsxl.Workbook
