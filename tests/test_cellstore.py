"""
CellStore (the compact ``Worksheet._cells`` of loaded worksheets) must behave
exactly like the dict openpyxl uses.  Random sequences of worksheet / dict
operations are applied to workbooks loaded by openpyxl and by openrsxl, every
result is compared, as well as the final state and the saved file.
"""

import datetime
import io
import os

import openpyxl
import pytest
from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st
from oracle.compare import ADDR_RE, norm, norm_text, saved_parts, sheet_state

import openrsxl

EXAMPLES = int(os.environ.get("OPENRSXL_FUZZ_EXAMPLES", "300"))
SETTINGS = settings(max_examples=EXAMPLES, deadline=None, suppress_health_check=list(HealthCheck))


def make_source():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for r in range(1, 9):
        for c in range(1, 6):
            if (r * c) % 7 == 3:
                continue  # holes
            v = [r * c, r / 3, f"s{r}{c}", None, True, datetime.datetime(2020, r, c)][(r + c) % 6]
            ws.cell(row=r, column=c, value=v)
    ws["F2"] = "=SUM(A2:E2)"
    ws["F3"] = "=SUM(A3:E3)"
    ws["G1"] = "=A1*2"
    ws["G2"] = "=A2*2"
    ws["B2"].font = openpyxl.styles.Font(bold=True)
    ws["C3"].number_format = "0.00"
    ws.merge_cells("H1:I2")
    ws["A12"] = "far"
    other = wb.create_sheet("Other")
    other["B2"] = 1
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


SOURCE = make_source()

coords = st.tuples(st.integers(1, 14), st.integers(1, 10))
values = st.one_of(
    st.none(),
    st.integers(-5, 5),
    st.floats(allow_nan=False, width=32),
    st.text(alphabet="ab =#", max_size=4),
    st.booleans(),
)

ops = st.one_of(
    st.tuples(st.just("read"), coords),
    st.tuples(st.just("write"), coords, values),
    st.tuples(st.just("cell"), coords, values),
    st.tuples(st.just("contains"), coords),
    st.tuples(st.just("contains_float"), coords),
    st.tuples(st.just("get"), coords),
    st.tuples(st.just("del"), coords),
    st.tuples(st.just("pop"), coords),
    st.tuples(st.just("popitem")),
    st.tuples(st.just("len")),
    st.tuples(st.just("keys")),
    st.tuples(st.just("items")),
    st.tuples(st.just("dims")),
    st.tuples(
        st.just("iter"), st.integers(1, 6), st.integers(1, 6), st.integers(0, 3), st.integers(0, 3), st.booleans()
    ),
    st.tuples(st.just("values")),
    st.tuples(st.just("insert_rows"), st.integers(1, 10), st.integers(1, 3)),
    st.tuples(st.just("delete_rows"), st.integers(1, 10), st.integers(1, 3)),
    st.tuples(st.just("insert_cols"), st.integers(1, 6), st.integers(1, 2)),
    st.tuples(st.just("delete_cols"), st.integers(1, 6), st.integers(1, 2)),
    st.tuples(st.just("move"), st.integers(0, 3), st.integers(0, 3), st.booleans()),
    st.tuples(st.just("merge"), coords),
    st.tuples(st.just("unmerge")),
    st.tuples(st.just("append"), st.lists(values, max_size=4)),
    st.tuples(st.just("style"), coords),
    st.tuples(st.just("mutate_during_iter")),
    st.tuples(st.just("odd_key"), st.sampled_from([("x", 1), (1, 1, 1), 5, (1.0, 2), (True, 1)])),
    st.tuples(st.just("setitem_existing"), coords),
    st.tuples(st.just("reversed")),
    st.tuples(st.just("copy_sheet")),
    st.tuples(st.just("array_formulae")),
    st.tuples(st.just("epoch_mac")),
    st.tuples(st.just("iso_dates")),
)


def describe(x):
    """comparable description of an operation result"""
    if hasattr(x, "coordinate") and hasattr(x, "parent"):
        return (
            "cell",
            type(x).__name__,
            x.coordinate,
            norm(x.value),
            x.data_type,
            list(x._style) if getattr(x, "_style", None) is not None else None,
        )
    if isinstance(x, (list, tuple)):
        return [describe(y) for y in x]
    return norm(x)


def apply(lib, ws, op):
    kind = op[0]
    cells = ws._cells
    if kind == "read":
        return ws.cell(row=op[1][0], column=op[1][1]).value
    if kind == "write":
        ws.cell(row=op[1][0], column=op[1][1]).value = op[2]
        return None
    if kind == "cell":
        return ws.cell(row=op[1][0], column=op[1][1], value=op[2])
    if kind == "contains":
        return op[1] in cells
    if kind == "contains_float":
        return (float(op[1][0]), op[1][1]) in cells
    if kind == "get":
        return cells.get(op[1], "missing")
    if kind == "del":
        del cells[op[1]]
        return None
    if kind == "pop":
        return cells.pop(op[1], "missing")
    if kind == "popitem":
        return cells.popitem()
    if kind == "len":
        return len(cells)
    if kind == "keys":
        return list(cells)
    if kind == "items":
        return [(k, v) for k, v in cells.items()]
    if kind == "dims":
        return (ws.min_row, ws.max_row, ws.min_column, ws.max_column, ws.dimensions, ws._current_row)
    if kind == "iter":
        min_row, min_col, extra_r, extra_c, vo = op[1:]
        return list(
            ws.iter_rows(
                min_row=min_row, max_row=min_row + extra_r, min_col=min_col, max_col=min_col + extra_c, values_only=vo
            )
        )
    if kind == "values":
        return list(ws.values)
    if kind == "insert_rows":
        ws.insert_rows(op[1], op[2])
        return None
    if kind == "delete_rows":
        ws.delete_rows(op[1], op[2])
        return None
    if kind == "insert_cols":
        ws.insert_cols(op[1], op[2])
        return None
    if kind == "delete_cols":
        ws.delete_cols(op[1], op[2])
        return None
    if kind == "move":
        ws.move_range("A1:C4", rows=op[1], cols=op[2], translate=op[3])
        return None
    if kind == "merge":
        r, c = op[1]
        ws.merge_cells(start_row=r, start_column=c, end_row=r + 1, end_column=c + 1)
        return None
    if kind == "unmerge":
        if ws.merged_cells.ranges:
            ws.unmerge_cells(str(sorted(ws.merged_cells.ranges, key=str)[0]))
        return None
    if kind == "append":
        ws.append(op[1])
        return None
    if kind == "style":
        c = ws.cell(row=op[1][0], column=op[1][1])
        c.font = lib.styles.Font(italic=True)
        return c.has_style
    if kind == "mutate_during_iter":
        for k in cells:
            ws.cell(row=k[0] + 50, column=1)
        return None
    if kind == "odd_key":
        key = op[1]
        try:
            hash(key)
        except TypeError:
            return "unhashable"
        res = [key in cells, cells.get(key, "missing")]
        return res
    if kind == "setitem_existing":
        if op[1] in cells:
            cells[op[1]] = cells[op[1]]
        return list(cells)[:5]
    if kind == "reversed":
        return list(reversed(cells))
    if kind == "copy_sheet":
        cp = ws.parent.copy_worksheet(ws)
        return list(cp._cells)
    if kind == "array_formulae":
        return ws.array_formulae
    if kind == "epoch_mac":
        from importlib import import_module

        ws.parent.epoch = import_module(lib.__name__ + ".utils.datetime").MAC_EPOCH
        return None
    if kind == "iso_dates":
        ws.parent.iso_dates = True
        return None
    raise AssertionError(kind)


def run(lib, seq):
    wb = lib.load_workbook(io.BytesIO(SOURCE))
    ws = wb["Data"]
    out = []
    for op in seq:
        try:
            out.append(("ok", describe(apply(lib, ws, op))))
        except Exception as e:
            out.append(("error", type(e).__name__, ADDR_RE.sub("", norm_text(str(e)))))
    final = sheet_state(ws)
    try:
        names, parts = saved_parts(wb)
        saved = (names, parts)
    except Exception as e:
        saved = ("error", type(e).__name__, norm_text(str(e)))
    return out, final, saved


@SETTINGS
@given(st.lists(ops, min_size=1, max_size=12))
def test_operation_sequences(seq):
    a = run(openpyxl, seq)
    b = run(openrsxl, seq)
    for i, (x, y) in enumerate(zip(a[0], b[0])):
        assert x == y, (i, seq[i], x, y)
    assert a[1] == b[1]
    assert a[2] == b[2]


def test_loaded_cells_are_store():
    wb = openrsxl.load_workbook(io.BytesIO(SOURCE))
    from openrsxl.worksheet._cellstore import CellStore

    assert isinstance(wb["Data"]._cells, CellStore)
    # identity semantics of a dict
    ws = wb["Data"]
    assert ws["A1"] is ws["A1"]
    assert ws.cell(1, 1) is ws._cells[(1, 1)]


def test_store_dict_protocol():
    wb = openrsxl.load_workbook(io.BytesIO(SOURCE))
    pwb = openpyxl.load_workbook(io.BytesIO(SOURCE))
    c, d = wb["Data"]._cells, pwb["Data"]._cells
    assert list(c) == list(d)
    assert len(c) == len(d)
    assert list(c.keys()) == list(d.keys())
    assert [x.coordinate for x in c.values()] == [x.coordinate for x in d.values()]
    with pytest.raises(KeyError) as e1:
        c[(99, 99)]
    with pytest.raises(KeyError) as e2:
        d[(99, 99)]
    assert str(e1.value) == str(e2.value)
    with pytest.raises(TypeError):
        c[[1, 2]]
    assert repr(c).count("<Cell") == repr(d).count("<Cell")
    assert dict(c).keys() == d.keys()
    assert c.copy().keys() == d.copy().keys()
    assert (c == c) and not (c != c)


def test_deepcopy_and_pickle_roundtrip():
    import copy
    import pickle

    def run(lib):
        wb = lib.load_workbook(io.BytesIO(SOURCE))
        out = []
        for clone in (copy.deepcopy(wb), pickle.loads(pickle.dumps(wb))):
            clone["Data"]["A1"] = "changed"
            out.append(saved_parts(clone))
        # the original is untouched
        out.append(saved_parts(wb))
        return out

    assert run(openpyxl) == run(openrsxl)
