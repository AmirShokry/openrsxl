"""
openrsxl.extended.llm: lossless LLM readable compression.

Every workbook is checked by the oracle (tests/oracle/llm.py): the decoded
document must equal what openpyxl reads - values and types, formulas, the
style of every cell, merged cells, hyperlinks, comments, dimensions, data
validation, conditional formatting, tables, sheet and workbook settings -
in every encoding and every output form.
"""

import datetime
import glob
import io
import json
import math
import os

import openpyxl
import pytest
from fixtures import make_xlsx, sheet_xml
from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table, TableStyleInfo
from oracle.llm import Report, _decoded_style, _norm, check_filters, filter_keep, verify

from openrsxl import _native
from openrsxl.extended import llm
from openrsxl.extended.llm import _decode, compress_sheets, compress_workbook, decompress

HERE = os.path.dirname(__file__)
ENCODINGS = ["auto", "grid", "index", "hybrid"]
rs = _native._rs


def save(wb, tmp_path, name="book.xlsx"):
    p = str(tmp_path / name)
    wb.save(p)
    return p


def write(data, tmp_path, name="book.xlsx"):
    p = str(tmp_path / name)
    with open(p, "wb") as f:
        f.write(data)
    return p


def check(path, encoding="auto", full=True, outputs=True, full_cells=False):
    rep, doc = verify(path, encoding=encoding, full=full, check_outputs=outputs, full_cells=full_cells)
    assert rep, str(rep)
    return doc


# -- value tokens ---------------------------------------------------------------
def roundtrip(v, error=False):
    tok = rs.llm_token(v, error)
    got, dt = _decode.parse_token(tok)
    return tok, got, dt


@pytest.mark.parametrize(
    "v",
    [
        0,
        1,
        -1,
        2**63 - 1,
        -(2**63),
        2**80,
        -(2**100),
        0.0,
        -0.0,
        1.0,
        0.1,
        1e16,
        1e15,
        1e-5,
        1e-4,
        1.5e300,
        5e-324,
        1.7976931348623157e308,
        0.30000000000000004,
        True,
        False,
        "",
        " ",
        "a",
        " a",
        "a ",
        "Alice",
        "New York",
        "12",
        "-3",
        "1e5",
        ".5",
        "5.",
        "+1",
        "007",
        "TRUE",
        "FALSE",
        "true",
        "#N/A",
        "#x",
        "@x",
        "=A1",
        '"q"',
        "a|b",
        "a\nb",
        "a\r\nb",
        "tab\there",
        "x*3",
        "*2",
        "B=1",
        "ABC=x",
        "ABCD=x",
        "2024-01-05",
        "2024-01-05 10:30:00",
        "10:30:00",
        "10:30",
        "2024-1-5",
        " ",
        "\x00",
        "\x7f",
        "é",
        "日本語",
        "\U0001f600",
        "１２３",  # full-width digits are text
        "٣",
        "²",
        "2024-０1-05",
        "a\\b",
        "{x}",
        "(x)",
        "x=y",
        "1,000",
        "inf",
        "nan",
        "-",
        datetime.datetime(2024, 1, 5),
        datetime.datetime(2024, 1, 5, 13, 45, 7, 500),
        datetime.datetime(1, 1, 1),
        datetime.datetime(9999, 12, 31, 23, 59, 59, 999999),
        datetime.date(2024, 1, 5),
        datetime.time(0, 0),
        datetime.time(13, 45, 7, 1),
        datetime.timedelta(0),
        datetime.timedelta(days=1, seconds=3, microseconds=4),
        datetime.timedelta(seconds=-1),
        datetime.timedelta(days=-3, microseconds=7),
        datetime.timedelta(hours=1000),
        float("inf"),
        float("-inf"),
    ],
)
def test_token_roundtrip(v):
    tok, got, dt = roundtrip(v)
    assert type(got) is type(v) and got == v, (v, tok, got)
    assert "\n" not in tok and "|" not in tok or tok.startswith('"')


def test_token_nan_and_errors():
    tok, got, _ = roundtrip(float("nan"))
    assert math.isnan(got)
    for e in (
        "#N/A",
        "#DIV/0!",
        "#VALUE!",
        "#REF!",
        "#NAME?",
        "#NUM!",
        "#NULL!",
        "#SPILL!",
        "weird",
        "#a b",
        "#|",
        "#x*2",
    ):
        tok, got, dt = roundtrip(e, True)
        assert (got, dt) == (e, "e"), (e, tok)
    assert rs.llm_token(None) is None


@settings(max_examples=3000, deadline=None, suppress_health_check=list(HealthCheck))
@given(
    st.one_of(
        st.text(),
        st.text(alphabet="0123456789.-+eE:# |\"'@=*TRUEFALS ABC"),
        # digits of other scripts are text (only ASCII digits make numbers)
        st.text(alphabet="0123456789０１٠٣²۰.-:e "),
        st.integers(),
        st.floats(allow_nan=False),
        st.booleans(),
        st.datetimes(),
        st.dates(),
        st.times(),
        st.timedeltas(),
    )
)
def test_token_roundtrip_fuzz(v):
    try:
        tok = rs.llm_token(v)
    except UnicodeEncodeError:
        return
    got, _ = _decode.parse_token(tok)
    assert type(got) is type(v) and got == v, (v, tok, got)
    # a token never breaks a line or a grid field
    if not tok.startswith(('"', '@err:"')):
        assert "|" not in tok and "\n" not in tok


@settings(max_examples=3000, deadline=None)
@given(st.floats(allow_nan=False, allow_infinity=False))
def test_float_repr(f):
    # shortest round-trip digits in repr's notation; a 17-digit tie may be
    # broken differently than CPython (both read back as the same float)
    tok = rs.llm_token(f)
    assert float(tok) == f and str(float(tok)) == repr(f)
    assert len(tok) == len(repr(f)) and ("." in tok or "e" in tok)


def test_surrogates():
    s = "a\ud800b"
    tok = rs.llm_token(s)
    assert _decode.parse_token(tok) == (s, "s")


# -- formula fills ----------------------------------------------------------------
_REF = st.builds(
    lambda d1, c, d2, r: f"{'$' if d1 else ''}{c}{'$' if d2 else ''}{r}",
    st.booleans(),
    st.sampled_from(["A", "B", "Z", "AA", "XFD", "ZZZ", "XFE"]),
    st.booleans(),
    st.integers(1, 1048576),
)
_PIECE = st.one_of(
    _REF,
    st.sampled_from(
        [
            "+",
            "*",
            "(",
            ")",
            ",",
            ":",
            "!",
            "SUM(",
            "Sheet1!",
            "'My Sheet'!",
            '"A1"',
            "'x''y'!",
            "A:A",
            "1:1",
            "LOG10(",
            "Table1[Col]",
            "_x",
            "a1",
            "R1C1",
            "$",
            "0",
            "A0",
            "ABCD1",
            " ",
            "&",
            "@",
            "#",
        ]
    ),
    st.text(max_size=3),
)


@settings(max_examples=3000, deadline=None)
@given(
    st.lists(_PIECE, max_size=8), st.integers(1, 5000), st.integers(1, 300), st.integers(-50, 50), st.integers(-50, 50)
)
def test_fill_matches_rust(pieces, row, col, dr, dc):
    f = "=" + "".join(pieces)
    r2, c2 = max(1, row + dr), max(1, col + dc)
    assert _decode.fill_formula(f, row, col, r2, c2) == rs.llm_fill_formula(f, row, col, r2, c2)
    assert _decode.fill_formula(f, row, col, row, col) == f


def test_fill_semantics():
    f = _decode.fill_formula
    assert f("=B2*C2", 2, 4, 3, 4) == "=B3*C3"
    assert f("=$B$2+B$2+$B2", 2, 4, 5, 6) == "=$B$2+D$2+$B5"
    assert f("=SUM(A1:A9)", 10, 1, 10, 2) == "=SUM(B1:B9)"
    assert f('="A1"&A1', 1, 2, 2, 2) == '="A1"&A2'
    assert f("='My Sheet'!A1", 1, 2, 2, 2) == "='My Sheet'!A2"
    assert f("=A1", 1, 1, 1, 0) is None


def test_fill_flags():
    flags = rs.llm_fill_flags
    assert flags("=A1+B2") == (False, False)
    assert flags("=SUM(A:A)") == (True, False)
    assert flags("=SUM(1:1)") == (False, True)
    assert flags("=Table1[Col]") == (True, False)
    assert flags("=IF(ISNA(A1),#N/A,A1#)") == (False, False)
    assert flags('="A:A"') == (False, False)


# -- grid lines -------------------------------------------------------------------
@pytest.mark.parametrize(
    "line,start,expected",
    [
        ("1:a|b|c", 1, [(1, 1, "a"), (1, 2, "b"), (1, 3, "c")]),
        ("2:||x", 1, [(2, 3, "x")]),
        ("3:K=x|y", 1, [(3, 11, "x"), (3, 12, "y")]),
        ("4:0*3|1", 2, [(4, 2, "0"), (4, 3, "0"), (4, 4, "0"), (4, 5, "1")]),
        ('5:"a|b"*2|"c"', 1, [(5, 1, '"a|b"'), (5, 2, '"a|b"'), (5, 3, '"c"')]),
        ('6:@err:"x|y"|z', 1, [(6, 1, '@err:"x|y"'), (6, 2, "z")]),
        ('7:5"|x', 1, [(7, 1, '5"'), (7, 2, "x")]),
    ],
)
def test_grid_line(line, start, expected):
    assert _decode.parse_grid_line(line, start) == expected


def test_grid_line_unlabelled():
    assert _decode.parse_grid_line("a||b", 1, 8) == [(8, 1, "a"), (8, 3, "b")]
    assert _decode.parse_grid_line("9:a", 1, 8) == [(9, 1, "a")]  # a label wins
    with pytest.raises(_decode.DecodeError):
        _decode.parse_grid_line("a|b", 1)


def test_ranges():
    assert list(_decode.iter_ranges("A1,B2:C3")) == [(1, 1), (2, 2), (2, 3), (3, 2), (3, 3)]
    assert list(_decode.iter_ranges("B6:C14/4")) == [(6, 2), (6, 3), (10, 2), (10, 3), (14, 2), (14, 3)]


# -- row labels -------------------------------------------------------------------
def _grid_lines(doc, sheet=0):
    return doc.sheets[sheet].sections["values"].text.splitlines()


def test_row_labels(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in list(range(1, 25)) + [30, 31]:
        ws.cell(r, 1, f"r{r}")
    p = save(wb, tmp_path)
    doc = check(p, "grid")
    lines = _grid_lines(doc)
    # labelled: the first line, multiples of 10 and after a gap
    labelled = [line.split(":")[0] for line in lines if ":" in line]
    assert labelled == ["1", "10", "20", "30"]
    assert lines[1] == "r2" and lines[-1] == "r31"
    rows = list(range(1, 25)) + [30, 31]
    # row_anchor=1: every line labelled
    doc = compress_workbook(p, include="values", encoding="grid", row_anchor=1)
    assert [line.split(":")[0] for line in _grid_lines(doc)] == [str(r) for r in rows]
    for anchor in (1, 3, 7, 1000):
        doc = compress_workbook(p, include="values", encoding="grid", row_anchor=anchor)
        for src in (doc.text(), doc.to_json()):
            assert decompress(src).sheets[0].values == {(r, 1): f"r{r}" for r in rows}


def test_row_label_kept_when_ambiguous(tmp_path):
    # without its label, a line starting with a time (or "12:..." text) would
    # read as a labelled line
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "x"
    ws["A2"] = datetime.time(13, 45)
    ws["A3"] = "y"
    ws["A4"] = "12:ab"
    ws["A5"] = 7
    ws["B5"] = "5:z"
    p = save(wb, tmp_path)
    doc = check(p, "grid")
    lines = _grid_lines(doc)
    assert lines[1] == "2:13:45:00"
    assert lines[2] == "y"
    assert lines[3].startswith("4:")
    assert lines[4] == "7|5:z"


@pytest.mark.parametrize("encoding", ["grid", "hybrid"])
def test_row_labels_hybrid(tmp_path, encoding):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, 60):
        ws.cell(r, 1, "same")
        ws.cell(r, 2, "same")
        ws.cell(r, 3, "same")
        ws.cell(r, 4, r * 1.5)
        if r % 7:
            ws.cell(r, 5, f"t{r}")
    p = save(wb, tmp_path)
    doc = check(p, encoding)
    if encoding == "hybrid":
        assert doc.sheets[0].sections["values"].encoding == "hybrid"
        lines = _grid_lines(doc)
        first = next(i for i, line in enumerate(lines) if line[:1].isdigit())
        assert all("=" in line for line in lines[:first])  # block lines first
        assert any(not line[:1].isdigit() for line in lines[first:])  # unlabelled rows


def test_ditto_marks(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, 26):
        ws.cell(r, 1, datetime.datetime(2024, 1, 1 + r // 4))  # repeats down
        ws.cell(r, 2, "Receivables - North")
        ws.cell(r, 3, "x" if r % 2 else "^")  # a text "^"
        ws.cell(r, 4, 7)  # one token: never a ditto mark
        ws.cell(r, 5, r * 1.25)
    for c in range(1, 6):
        ws.cell(30, c, "same row twice")
        ws.cell(31, c, "same row twice")
    p = save(wb, tmp_path)
    doc = check(p, "grid")
    lines = _grid_lines(doc)
    assert lines[1].startswith("^|^|")
    assert '"^"' in doc.text()
    for line in lines:
        if _decode._LABEL.match(line):
            assert "^" not in line.replace('"^"', "")  # anchors spell values out
        assert "|7|" in line or line.startswith(("30:", "^*5"))
    assert lines[-1] == "^*5"
    # off: no marks, same cells
    plain = compress_workbook(p, include="values", encoding="grid", ditto=False)
    assert "^" not in plain.sheets[0].sections["values"].text.replace('"^"', "")
    assert decompress(plain.text()).sheets[0].values == decompress(doc.text()).sheets[0].values
    with pytest.raises(_decode.DecodeError):
        decompress('#openrsxl-llm/1\n## sheet 1 "S"\n### values grid start=A\n1:a\n||^\n')


# -- the workbook's style table ---------------------------------------------------
def test_style_table_once(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "A"
    ws["A1"] = 1
    ws["A1"].font = Font(bold=True)
    ws2 = wb.create_sheet("B")
    ws2["A1"] = 2
    ws2["B2"] = 3
    ws2["B2"].fill = PatternFill("solid", fgColor="FF0000")
    ws3 = wb.create_sheet("C")
    ws3["A1"] = "plain"
    p = save(wb, tmp_path)
    doc = check(p)
    text = doc.text()
    assert text.count("\n## styles\n") == 1
    assert text.count("\ns0: ") == 1
    # sheets map cells only; a sheet with only s0 cells has no styles section
    assert all(
        not line.startswith(("font ", "fill ", "border ", "s0:")) for line in doc["A"]["styles"].text.splitlines()
    )
    assert "styles" not in doc["C"].sections
    # a part of the document carries the styles it uses
    only_a = decompress(doc.text(sheets=["A"]))
    sid = only_a["A"].style_ids[(1, 1)]
    assert only_a["A"].styles[sid].font.b
    assert all(sd.fill.fgColor.rgb != "00FF0000" for sd in only_a.styles.values())
    assert "FF0000" not in doc.text(sheets=["A"])
    assert "FF0000" in doc.text(sheets=["B"])
    assert "styles" not in doc.text(sections=["values"], legend=False).split("\n", 1)[1]
    assert doc.to_dict(sheets=["B"])["styles"] == doc.styles.text(doc["B"].used_styles)
    for src in (doc.text(), doc.to_json(), doc.text(sheets=["B"]), doc.to_dict(sheets=["A"])):
        book = decompress(src)
        for s in book.sheets:
            assert 0 in s.styles
            for sid in s.style_ids.values():
                assert sid in s.styles


# -- workbooks -----------------------------------------------------------------
def rich_book():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Q1", "Q2", "Total", "Note", "When"])
    for i in range(2, 60):
        ws.append(["North" if i < 25 else "South", i * 10, i * 1.5, f"=B{i}+C{i}", None, datetime.date(2024, 1, 1)])
    ws["E2"] = "=Sales!A1"
    ws["E3"] = "123"
    ws["E4"] = "  pad"
    ws["E5"] = "a|b"
    ws["E6"] = "#N/A"
    ws["E7"] = "line1\nline2"
    ws["E8"] = 0.1 + 0.2
    ws["E9"] = True
    ws["E10"] = datetime.datetime(2024, 1, 5, 13, 45, 7, 500)
    ws["E11"] = datetime.time(10, 30)
    ws["E12"] = datetime.timedelta(hours=30, minutes=5)
    ws["E13"] = 2**70
    ws["E14"] = "=SUM(A:A)"
    ws["F14"] = "=SUM(B:B)"
    ws["E15"] = "=SUM(1:1)"
    ws["E16"] = "=SUM(2:2)"
    ws["H1"] = ArrayFormula("H1:H3", "=TRANSPOSE(B1:D1)")
    ws["H80"] = "far away"
    ws["D61"] = "=SUM(D2:D59)"
    ws["G2"] = '=IF(A2="North",1,2)'
    for c in ws[1]:
        c.font = Font(bold=True, color="FF0000")
        c.fill = PatternFill("solid", fgColor="FFFF00")
        c.border = Border(bottom=Side(style="thin"))
    ws["B2"].number_format = "0.00%"
    ws["A1"].alignment = Alignment(horizontal="center", wrap_text=True)
    ws["A2"].protection = Protection(locked=False)
    ws["C3"].quotePrefix = True
    for r in range(2, 60, 2):  # zebra stripes
        ws.cell(r, 3).fill = PatternFill("solid", fgColor="DDDDDD")
    ws.merge_cells("J1:K2")
    ws.merge_cells("J5:J9")
    ws["A3"].comment = Comment("check this", "Amir")
    ws["A5"].comment = Comment("multi\nline", None)
    ws["A4"].hyperlink = "https://example.com"
    ws["B4"].hyperlink = "https://example.com"
    ws["C4"].hyperlink = "#Sales!A1"
    dv = DataValidation(type="list", formula1='"a,b,c"', allow_blank=True)
    dv.add("E20:E30")
    ws.add_data_validation(dv)
    ws.conditional_formatting.add(
        "B2:B59", CellIsRule(operator="greaterThan", formula=["100"], font=Font(color="FF0000"))
    )
    ws.conditional_formatting.add(
        "C2:C59", ColorScaleRule(start_type="min", start_color="FF0000", end_type="max", end_color="00FF00")
    )
    ws.conditional_formatting.add(
        "A2:A59", FormulaRule(formula=['A2="North"'], fill=PatternFill("solid", bgColor="CCFFCC"))
    )
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].hidden = True
    ws.row_dimensions[1].height = 30
    for r in range(70, 75):
        ws.row_dimensions[r].height = 15
    ws.freeze_panes = "B2"
    ws.print_title_rows = "1:1"
    ws.print_area = "A1:F59"
    ws.sheet_properties.tabColor = "1072BA"
    ws.page_setup.orientation = "landscape"
    t = wb.create_sheet("Table")
    t.append(["id", "name"])
    for i in range(10):
        t.append([i, f"n{i}"])
    tab = Table(displayName="T1", ref="A1:B11")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    t.add_table(tab)
    sp = wb.create_sheet("Sparse")
    for coord in ("B2", "Z200", "C9", "AA3000"):
        sp[coord] = "x"
    sp["M50"] = 1
    hidden = wb.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "secret"
    wb.create_sheet("Empty")
    style = NamedStyle(name="highlight", font=Font(italic=True))
    wb.add_named_style(style)
    sp["B3"].style = "highlight"
    wb.defined_names["Tax"] = openpyxl.workbook.defined_name.DefinedName("Tax", attr_text="Sales!$B$2")
    wb.properties.creator = "tests"
    wb.properties.title = "A title"
    return wb


@pytest.mark.parametrize("encoding", ENCODINGS)
def test_rich_book(tmp_path, encoding):
    check(save(rich_book(), tmp_path), encoding)


def test_rich_book_text(tmp_path):
    doc = compress_workbook(save(rich_book(), tmp_path))
    text = doc.text()
    assert text.startswith("#openrsxl-llm/1 ")
    # every line is a header, a legend line or section content
    assert all(not line.startswith("#") or line.startswith(("#openrsxl", "## ", "### ")) for line in text.split("\n"))
    f = doc["Sales"]["formulas"].text
    assert "D2:D59=B2+C2" in f
    # whole-column / row references are never filled across the direction
    # Excel would shift them in
    assert "E14=SUM(A:A)" in f and "F14=SUM(B:B)" in f
    assert "E15=SUM(1:1)" in f and "E16=SUM(2:2)" in f
    assert "H1[array H1:H3]=TRANSPOSE(B1:D1)" in f
    # zebra stripes: one strided range
    assert "/2=s" in doc["Sales"]["styles"].text
    assert doc["Sparse"].analysis["encoding"] == "index"
    assert "Empty" in doc.text() and "Hidden" in doc.text()


def test_decoded_objects(tmp_path):
    book = decompress(compress_workbook(save(rich_book(), tmp_path)))
    s = book["Sales"]
    assert s.values[(1, 1)] == "Region"
    assert s.formulas[(3, 4)] == "=B3+C3"
    assert s.values[(12, 5)] == datetime.timedelta(hours=30, minutes=5)
    assert s.cell("A1")["style"].font.b
    assert s.hyperlinks[(4, 1)].target == "https://example.com"
    assert s.comments[(3, 1)].content == "check this"
    assert "J1:K2" in s.merged
    assert book.sheetnames == ["Sales", "Table", "Sparse", "Hidden", "Empty"]
    assert book.sheet_entries[3]["state"] == "hidden"
    assert [d.name for d in book.defined_names] == ["Tax"]


def dense_numeric():
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, 400):
        ws.append([r * 1.25 + c for c in range(12)])
    return wb


def constant_columns():
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, 300):
        ws.append(["USD", "2024", r, r * 2, "EU" if r < 150 else "US", "x"])
    return wb


def sparse_repeats():
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, 2000, 7):
        ws.cell(r, (r * 13) % 50 + 1).value = "yes" if r % 2 else 0
    return wb


@pytest.mark.parametrize(
    "build,expected",
    [(dense_numeric, "grid"), (constant_columns, "hybrid"), (sparse_repeats, "index")],
)
def test_analysis_choice(tmp_path, build, expected):
    p = save(build(), tmp_path)
    a = llm.analyze_workbook(p)
    assert list(a.values())[0]["encoding"] == expected
    for enc in ENCODINGS:
        check(p, enc, full=False, outputs=False)


def test_periodic(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, 400):
        k = r % 4
        ws.append([f"Box {k}", k * 10, k * 20] + [f"=B{r}*{k}"])
    p = save(wb, tmp_path)
    doc = check(p)
    assert len(doc.sheets[0]["formulas"].text.splitlines()) <= 8
    check(p, "index", full=False, outputs=False)
    check(p, "hybrid", full=False, outputs=False)


# -- raw XML edge cases ---------------------------------------------------------
RAW = {
    "types": (
        '<row r="1"><c r="A1"><v>1</v></c><c r="B1"><v>1.0</v></c><c r="C1" t="b"><v>1</v></c>'
        '<c r="D1" t="e"><v>#N/A</v></c><c r="E1" t="s"><v>2</v></c><c r="F1" t="str"><v>txt</v></c>'
        '<c r="G1" t="inlineStr"><is><t>inline</t></is></c><c r="H1" s="1"><v>45000</v></c>'
        '<c r="I1" s="2"><v>1.5</v></c><c r="J1" s="4"><v>45000.25</v></c><c r="K1" t="d"><v>2024-01-05T10:00:00</v></c>'
        '<c r="L1" t="s"><v>3</v></c><c r="M1" s="1"><v>0.5</v></c><c r="N1" t="e"><v>#DIV/0!</v></c>'
        '<c r="O1" t="d"><v>2024-01-05</v></c><c r="P1"><v>1E+30</v></c><c r="Q1"><v>-0</v></c></row>'
    ),
    "formulas": (
        '<row r="1"><c r="A1"><f>B1*2</f><v>4</v></c><c r="B1"><v>2</v></c>'
        '<c r="C1" t="e"><f>1/0</f><v>#DIV/0!</v></c><c r="D1" t="str"><f>"a"&amp;"b"</f><v>ab</v></c>'
        '<c r="E1" t="b"><f>TRUE()</f><v>1</v></c></row>'
        '<row r="2"><c r="A2"><f t="shared" ref="A2:A5" si="0">B2*2</f><v>0</v></c></row>'
        '<row r="3"><c r="A3"><f t="shared" si="0"/><v>0</v></c></row>'
        '<row r="4"><c r="A4"><f t="shared" si="0"/></c></row>'
        '<row r="5"><c r="A5"><f t="shared" si="0"/><v>0</v></c>'
        '<c r="B5"><f t="array" ref="B5:B6">ROW(1:2)</f><v>1</v></c></row>'
        '<row r="7"><c r="C7"><f t="dataTable" ref="C7:D8" dt2D="1" dtr="1" r1="A1" r2="B1"/><v>1</v></c>'
        '<c r="D7"><f>A1&#10;+1</f></c><c r="E7"><f></f><v>3</v></c></row>'
    ),
    "unordered": '<row r="5"><c r="A5"><v>5</v></c></row><row r="2"><c r="A2"><v>2</v></c></row>'
    '<row r="5"><c r="B5"><v>6</v></c></row>',
    "duplicates": '<row r="1"><c r="B1"><v>1</v></c><c r="A1"><v>2</v></c><c r="B1"><v>3</v></c></row>',
    "no_coordinates": '<row><c><v>1</v></c><c><v>2</v></c></row><row><c t="s"><v>0</v></c></row>',
    "styles_only": '<row r="3"><c r="B3" s="3"/><c r="C3" s="3"/><c r="D3" s="6"/></row>',
    "wide": '<row r="1"><c r="A1"><v>1</v></c><c r="XFD1"><v>2</v></c></row>',
    "spaces": '<row r="1"><c r="A1" t="s"><v>2</v></c><c r="B1" t="inlineStr"><is><t xml:space="preserve"> </t></is></c></row>',
}


@pytest.mark.parametrize("name", sorted(RAW))
@pytest.mark.parametrize("encoding", ENCODINGS)
def test_raw(tmp_path, name, encoding):
    # read-only openpyxl skips rows out of order: full mode keeps them, as we do
    check(write(make_xlsx([RAW[name]]), tmp_path), encoding, outputs=encoding == "auto", full_cells=name == "unordered")


def _financial():
    import importlib.util

    spec = importlib.util.spec_from_file_location(
        "make_financial_files", os.path.join(HERE, "..", "tools", "make_financial_files.py")
    )
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


@pytest.mark.parametrize("seed", [1, 2])
@pytest.mark.parametrize("encoding", ENCODINGS)
def test_financial_workbooks(tmp_path, seed, encoding):
    # ledgers, trial balance, P&L, ageing, budget, cash flow, FX, payroll,
    # fixed assets - formulas with cached results, accounting formats
    p = _financial().make(str(tmp_path / f"fin{seed}.xlsx"), seed, scale=0.05)
    doc = check(p, encoding, outputs=encoding == "auto")
    if encoding == "auto":
        # the choice is never worse than the plain grid
        grid = compress_workbook(p, include="values", encoding="grid", legend=False)
        auto = compress_workbook(p, include="values", legend=False)
        for a, g in zip(auto.sheets, grid.sheets):
            if "values" in a.sections:
                assert a.sections["values"].tokens <= g.sections["values"].tokens * 1.01, a.name
        assert doc.sheets


def test_corpus_regressions(tmp_path):
    """Found by the corpus run: cells in row 0 (kept by full mode), hashed
    sheet protection passwords (openpyxl's parser does not hash them again),
    header / footer escapes openpyxl does not read back from its own XML"""
    after = (
        '<sheetProtection password="CF7A" sheet="1" objects="1" scenarios="1"/>'
        "<headerFooter><oddHeader>&amp;LPage &amp;P_x000a_</oddHeader>"
        "<oddFooter>&amp;Rx 8/12_x000D_</oddFooter></headerFooter>"
    )
    data = '<row r="0"><c r="A0"><v>1</v></c></row><row r="1"><c r="A1"><v>2</v></c></row>'
    p = write(make_xlsx(raw_sheets=[sheet_xml(data, after=after)]), tmp_path)
    doc = check(p)
    assert "A0" in doc.text() or "0:" in doc.text()
    book = decompress(doc.text())
    assert book.sheets[0].values[(0, 1)] == 1
    assert book.sheets[0].sheet["protection"].password == "CF7A"


def test_date1904(tmp_path):
    p = write(
        make_xlsx(['<row r="1"><c r="A1" s="1"><v>1000</v></c><c r="B1"><v>3</v></c></row>'], date1904=True), tmp_path
    )
    doc = check(p)
    assert decompress(doc.text()).date1904
    assert "date1904=1" in doc.text()


def test_python_fallback(tmp_path, monkeypatch):
    """The pure Python feeder gives the same document as the native one"""
    p = save(rich_book(), tmp_path)
    for name in sorted(RAW):
        write(make_xlsx([RAW[name]]), tmp_path, name + ".xlsx")
    paths = [p] + [str(tmp_path / (n + ".xlsx")) for n in sorted(RAW)]
    native = [compress_workbook(x).text() for x in paths]
    monkeypatch.setattr(_native, "ENABLED", False)
    for x, t in zip(paths, native):
        assert compress_workbook(x).text() == t, x


def test_native_fallback_on_error(tmp_path, monkeypatch):
    """A NativeFallback while reading switches to the Python feeder"""
    from openrsxl.extended.llm import _encode

    p = save(rich_book(), tmp_path)
    expected = compress_workbook(p).text()

    def boom(enc, ws, n):
        raise _native.NativeFallback("test")

    monkeypatch.setattr(_encode, "_feed_native", boom)
    assert compress_workbook(p).text() == expected


# -- one sheet at a time ------------------------------------------------------------
def sheets_book():
    wb = rich_book()
    from openpyxl.chart import BarChart, Reference

    cs = wb.create_chartsheet("Chart")  # a chartsheet between worksheets
    chart = BarChart()
    chart.add_data(Reference(wb["Sales"], min_col=2, min_row=1, max_row=10), titles_from_data=True)
    cs.add_chart(chart)  # (openpyxl cannot read a chartsheet without a chart)
    wb.move_sheet("Chart", offset=-(len(wb.sheetnames) - 2))
    odd = wb.create_sheet("a<b")  # names that collide as file names
    odd["A1"] = "slash"
    other = wb.create_sheet("a_b")
    other["A1"] = "underscore"
    return wb


@pytest.mark.parametrize("encoding", ENCODINGS)
def test_each_sheet_lossless(tmp_path, encoding):
    """Every worksheet compressed alone decodes to exactly what openpyxl reads"""
    p = save(sheets_book(), tmp_path)
    names = [ws.title for ws in openpyxl.load_workbook(p).worksheets]
    for name in names:
        rep, doc = verify(p, encoding=encoding, sheets=[name])
        assert rep, f"{name}: {rep}"
        assert [s.name for s in doc.sheets] == [name]


def test_compress_sheets_equals_single(tmp_path):
    p = save(sheets_book(), tmp_path)
    for kw in ({}, {"include": "values"}, {"skip_hidden": True, "print_area_only": True}, {"encoding": "index"}):
        each = list(compress_sheets(p, **kw))
        names = [ws.title for ws in openpyxl.load_workbook(p).worksheets]
        assert [n for n, _ in each] == names
        for name, doc in each:
            single = compress_workbook(p, sheets=name, **kw)
            assert doc.text() == single.text(), (name, kw)
            assert doc.to_json() == single.to_json(), (name, kw)
    # a selection, in workbook order whatever the order asked
    assert [n for n, _ in compress_sheets(p, sheets=["Sparse", "Sales"])] == ["Sales", "Sparse"]
    assert [s.name for s in compress_workbook(p, sheets=["Sparse", "Sales"]).sheets] == ["Sales", "Sparse"]


def test_sheet_numbers_are_workbook_positions(tmp_path):
    p = save(sheets_book(), tmp_path)
    positions = {name: i for i, name in enumerate(openpyxl.load_workbook(p).sheetnames, 1)}
    full = compress_workbook(p)
    assert all(s.index == positions[s.name] for s in full.sheets)
    assert "Chart" not in [s.name for s in full.sheets]
    for name, doc in compress_sheets(p):
        (sheet,) = doc.sheets
        assert sheet.index == positions[name]
        assert sheet.header == full[name].header
        assert doc[name].text() == full[name].text()
        assert decompress(doc.text()).sheets[0].index == positions[name]
        json.loads(doc.to_json())  # valid JSON whatever the first index


def test_selected_sheet_decodes_like_full(tmp_path):
    """A sheet's own document decodes to the same cells, styles and metadata
    as the same sheet in the document of the whole workbook"""
    p = save(sheets_book(), tmp_path)
    full = decompress(compress_workbook(p))
    for name, doc in compress_sheets(p):
        a, b = decompress(doc)[name], full[name]
        for attr in ("values", "data_types", "formulas", "merged", "hyperlinks", "comments"):
            assert _norm(getattr(a, attr)) == _norm(getattr(b, attr)), (name, attr)
        assert a.style_ids.keys() == b.style_ids.keys()
        for key in a.style_ids:
            assert _decoded_style(a.cell(key)["style"]) == _decoded_style(b.cell(key)["style"]), (name, key)


def test_compress_sheets_files(tmp_path):
    p = save(sheets_book(), tmp_path)
    base = str(tmp_path / "out" / "book")
    files = {name: doc.files for name, doc in compress_sheets(p, output=["text", "json"], path=base)}
    texts = [os.path.basename(f["text"]) for f in files.values()]
    assert len({t.lower() for t in texts}) == len(texts)  # no name overwrites another
    assert "book.Sales.llm.txt" in texts
    assert {"book.a_b.llm.txt", "book.a_b~2.llm.txt"} <= set(texts)
    for name, f in files.items():
        with open(f["text"], encoding="utf-8") as fh:
            assert fh.read() == compress_workbook(p, sheets=name).text()
        assert decompress(f["json"]).sheets[0].name == name


def test_compress_sheets_lazy(tmp_path):
    """Stopping the iteration skips the remaining sheets (and closes the
    workbook)"""
    p = save(sheets_book(), tmp_path)
    base = str(tmp_path / "lazy")
    it = compress_sheets(p, output="text", path=base)
    name, doc = next(it)
    it.close()
    written = sorted(f for f in os.listdir(tmp_path) if f.startswith("lazy"))
    assert written == [f"lazy.{name}.llm.txt"]


def test_check_sheets_oracle(tmp_path):
    """The corpus check of per-sheet compression on the synthetic books"""
    from oracle.llm import check_sheets

    paths = [save(b(), tmp_path, f"{b.__name__}.xlsx") for b in (sheets_book, rich_book, filter_book)]
    for name in sorted(RAW):
        paths.append(write(make_xlsx([RAW[name]]), tmp_path, name + ".xlsx"))
    for p in paths:
        rep = check_sheets(p, Report())
        assert rep, f"{os.path.basename(p)}: {rep}"


def test_sheet_selection_errors(tmp_path):
    p = save(sheets_book(), tmp_path)
    with pytest.raises(KeyError, match="no worksheet named 'nope'"):
        compress_workbook(p, sheets="nope")
    with pytest.raises(KeyError, match="chartsheet"):
        compress_workbook(p, sheets=["Sales", "Chart"])
    with pytest.raises(KeyError):
        list(compress_sheets(p, sheets=["nope"]))
    with pytest.raises(TypeError):
        compress_workbook(p, sheet="Sales")  # misspelt option


def test_unselected_sheets_not_scanned(tmp_path):
    """The metadata of sheets that are not compressed is not read"""
    from openrsxl.extended import _streaming
    from openrsxl.extended._load import _load_for

    p = save(sheets_book(), tmp_path)
    wb = _load_for(p, ["Sparse"], read_merged_cells=True, read_comments=True)
    try:
        flags = {ws.title: ws._ext.read_merged_cells for ws in wb.worksheets}
    finally:
        wb.close()
    assert flags["Sparse"] and not any(v for k, v in flags.items() if k != "Sparse")
    calls = []
    real = _streaming._native._rs.prescan_sheet

    def counting(src, *a):
        calls.append(1)
        return real(src, *a)

    _streaming._native._rs.prescan_sheet = counting
    try:
        compress_workbook(p, sheets="Sparse")
        one = len(calls)
        calls.clear()
        compress_workbook(p)
        every = len(calls)
    finally:
        _streaming._native._rs.prescan_sheet = real
    assert one == 1 and every > one


# -- sections, outputs -------------------------------------------------------------
def test_sections(tmp_path):
    p = save(rich_book(), tmp_path)
    full = decompress(compress_workbook(p))
    for name in llm.SECTIONS:
        doc = compress_workbook(p, include=name)
        assert doc.included == (name,)
        text = doc.text()
        heads = {line.split(" ")[1] for line in text.split("\n") if line.startswith("### ")}
        assert heads <= {name}, (name, heads)
        if name == "workbook":
            assert "## workbook" in text
        book = decompress(text)
        if name == "values":
            assert [s.values for s in book.sheets] == [s.values for s in full.sheets]
        if name == "formulas":
            assert [_norm(s.formulas) for s in book.sheets] == [_norm(s.formulas) for s in full.sheets]
            assert all(not s.values for s in book.sheets)
    doc = compress_workbook(p, include="cells", exclude="styles")
    assert doc.included == ("values", "formulas")
    doc = compress_workbook(p, include="metadata", sheets=["Sales"])
    assert [s.name for s in doc.sheets] == ["Sales"] and "values" not in doc.included
    with pytest.raises(ValueError):
        compress_workbook(p, include="nope")
    with pytest.raises(KeyError):
        compress_workbook(p, sheets=["nope"])
    with pytest.raises(ValueError):
        compress_workbook(p, encoding="nope")
    # parts of a document
    doc = compress_workbook(p)
    only = doc.text(sections=["formulas"], sheets=["Sales"], legend=False)
    assert "### formulas" in only and "### values" not in only and "## sheet 2" not in only


def test_outputs(tmp_path):
    p = save(rich_book(), tmp_path)
    mem = compress_workbook(p)
    doc = compress_workbook(p, output="all", path=str(tmp_path / "out" / "book"))
    assert doc.text() == mem.text()
    assert set(doc.files) == {"text", "json"}
    with open(doc.files["text"], encoding="utf-8") as f:
        assert f.read() == mem.text()
    with open(doc.files["json"], encoding="utf-8") as f:
        assert json.load(f) == json.loads(mem.to_json())
    # files only: nothing kept in memory, same content
    doc2 = compress_workbook(p, output="text", path=str(tmp_path / "x.llm.txt"))
    assert doc2.files["text"].endswith("x.llm.txt") and not doc2.in_memory
    with pytest.raises(ValueError):
        doc2.text()
    with open(doc2.files["text"], encoding="utf-8") as f:
        assert f.read() == mem.text()
    # default path: next to the workbook
    doc3 = compress_workbook(p, output="json")
    assert doc3.files["json"] == os.path.splitext(p)[0] + ".llm.json"
    # decoding from paths
    assert decompress(doc.files["json"]).sheetnames == decompress(doc.files["text"]).sheetnames
    # file objects
    with open(p, "rb") as f:
        assert compress_workbook(io.BytesIO(f.read())).text().split("\n", 1)[1] == mem.text().split("\n", 1)[1]
    with pytest.raises(ValueError):
        compress_workbook(io.BytesIO(open(p, "rb").read()), output="text")


def test_default_output_is_cheapest(tmp_path):
    """The default is the text form in memory (no files, no JSON); it costs
    fewer tokens than the JSON form of the same document"""
    for wb in (rich_book(), dense_numeric(), filter_book()):
        p = save(wb, tmp_path)
        doc = compress_workbook(p)
        assert doc.files == {} and doc.in_memory
        assert not os.path.exists(p[:-5] + ".llm.txt") and not os.path.exists(p[:-5] + ".llm.json")
        text, js = doc.text(), doc.to_json()
        assert llm.estimate_tokens(text) < llm.estimate_tokens(js)
        assert decompress(text).sheets[0].values == decompress(js).sheets[0].values


def test_legend(tmp_path):
    p = save(rich_book(), tmp_path)
    full = compress_workbook(p).legend
    values = compress_workbook(p, include="values").legend
    assert "### formulas" not in values and "Values:" in values
    assert len(values) < len(full)
    assert compress_workbook(p, legend=False).legend is None
    assert all(line.startswith(("- ", "  ")) for line in llm.LEGEND.split("\n"))


# -- filters: skip_hidden, skip_hidden_sheets, print_area_only ----------------------
def filter_book():
    wb = rich_book()
    ws = wb["Sales"]
    for r in (3, 4, 30, 31, 32):
        ws.row_dimensions[r].hidden = True
    ws.column_dimensions.group("B", "C", hidden=True)  # one <col min=2 max=3>
    ws.column_dimensions["F"].hidden = True
    ws["C2"].comment = Comment("in a hidden column", "a")
    ws["A3"].comment = Comment("in a hidden row", "a")
    ws["A5"].comment = Comment("visible", "a")
    ws["C5"].hyperlink = "https://hidden.example"
    ws["A6"].hyperlink = "https://visible.example"
    ws.print_area = "A1:E20 A40:D45"
    ws.print_title_rows = "1:2"
    ws.print_title_cols = "A:A"
    for state in ("hidden", "veryHidden"):
        h = wb.create_sheet(f"{state} sheet")
        h["A1"] = state
        h.sheet_state = state
    z = wb.create_sheet("zero height")  # rows hidden by default
    for r in range(1, 9):
        z.cell(r, 1, r)
    z.sheet_format.zeroHeight = True
    z.row_dimensions[2].height = 15
    z.row_dimensions[5].height = 15
    return wb


def _range_cells(ref):
    from openpyxl.utils.cell import range_boundaries

    c1, r1, c2, r2 = range_boundaries(ref)
    return [(r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


@pytest.mark.parametrize("encoding", ENCODINGS)
@pytest.mark.parametrize("skip_hidden,area", [(True, False), (False, True), (True, True)])
def test_filters(tmp_path, encoding, skip_hidden, area):
    """A filtered document is exactly the lossless one restricted to the
    cells openpyxl says are visible / printed"""
    p = save(filter_book(), tmp_path)
    full = decompress(compress_workbook(p, encoding=encoding))
    doc = compress_workbook(p, encoding=encoding, skip_hidden=skip_hidden, print_area_only=area)
    for d in (decompress(doc), decompress(doc.text()), decompress(doc.to_json())):
        ref = openpyxl.load_workbook(p)
        assert [s.name for s in d.sheets] == [s.name for s in full.sheets]
        for s in d.sheets:
            keep = filter_keep(ref[s.name], skip_hidden, area)
            f = full[s.name]
            for attr in ("values", "data_types", "formulas", "style_ids", "hyperlinks", "comments"):
                want = {k: v for k, v in getattr(f, attr).items() if keep(k)}
                assert _norm(getattr(s, attr)) == _norm(want), (s.name, attr)
            # merged ranges: those holding a kept cell; other metadata whole
            want = [m for m in f.merged if any(keep(c) for c in _range_cells(m))]
            assert s.merged == want and s.row_dimensions == f.row_dimensions
    sales = decompress(doc)["Sales"]
    expected = (["skip=hidden"] if skip_hidden else []) + (["area=A1:E20,A40:D45+titles:1:2,A:A"] if area else [])
    assert sales.filter == expected
    assert doc["Sales"].analysis["skipped_cells"] > 0
    assert (doc["zero height"].analysis["skipped_cells"] == 6) == skip_hidden
    assert "Filtered sheets:" in doc.legend


def test_filter_rows_cols(tmp_path):
    p = save(filter_book(), tmp_path)
    s = decompress(compress_workbook(p, skip_hidden=True))
    sales = s["Sales"]
    assert not any(r in (3, 4, 30, 31, 32) or c in (2, 3, 6) for r, c in sales.values)
    assert (5, 1) in sales.comments and (2, 3) not in sales.comments and (3, 1) not in sales.comments
    assert (6, 1) in sales.hyperlinks and (5, 3) not in sales.hyperlinks
    assert sorted(s["zero height"].values) == [(2, 1), (5, 1)]
    p_area = decompress(compress_workbook(p, print_area_only=True))["Sales"]
    # B:E: the two areas (+ title rows 1:2); A: title column across rows 1..45
    assert {r for r, c in p_area.values if c > 1} == set(range(1, 21)) | set(range(40, 46))
    assert {r for r, c in p_area.values if c == 1} == set(range(1, 46))
    assert max(c for _, c in p_area.values) == 5


def test_filter_hidden_sheets(tmp_path):
    p = save(filter_book(), tmp_path)
    doc = compress_workbook(p, skip_hidden_sheets=True)
    visible = [ws.title for ws in openpyxl.load_workbook(p) if ws.sheet_state == "visible"]
    assert [s.name for s in doc.sheets] == visible and "Hidden" not in visible
    assert "Filtered sheets:" in doc.legend
    # the workbook section still lists every sheet and its state
    assert "state=veryHidden" in doc.text(sections=["workbook"])
    # combined with sheets=: the hidden sheet is still left out
    picked = compress_workbook(p, sheets=["hidden sheet", "Sales"], skip_hidden_sheets=True)
    assert [s.name for s in picked.sheets] == ["Sales"]


def test_filters_noop(tmp_path):
    """Flags that find nothing to leave out change nothing"""
    p = save(dense_numeric(), tmp_path)
    base = compress_workbook(p).text()
    assert compress_workbook(p, skip_hidden=True, skip_hidden_sheets=True, print_area_only=True).text() == base
    assert "Filtered" not in base


def test_filters_python_fallback(tmp_path, monkeypatch):
    p = save(filter_book(), tmp_path)
    kw = dict(skip_hidden=True, print_area_only=True)
    native = compress_workbook(p, **kw).text()
    monkeypatch.setattr(_native, "ENABLED", False)
    assert compress_workbook(p, **kw).text() == native


def test_filter_encoder_options():
    """The Rust filter: ranges merged, an empty area list keeps nothing"""

    def cells(**kw):
        enc = rs.SheetEncoder(dict(encoding="grid", formulas=False, styles=False, **kw))
        enc.begin_pass(1)
        enc.feed([(r, c, r * 10 + c, "n", None, None) for r in range(1, 6) for c in range(1, 4)], None, None)
        return enc.end_pass()

    assert cells()["cells"] == 15
    assert cells(skip_rows=[(2, 2), (3, 4)])["cells"] == 6
    assert cells(skip_cols=[(1, 2)], skip_rows=[(5, 5)])["cells"] == 4
    a = cells(areas=[(2, 2, 3, 3)])
    assert a["cells"] == 4 and a["skipped_cells"] == 11 and a["dimension"] == "B2:C3"
    assert cells(areas=[])["cells"] == 0
    with pytest.raises(ValueError):
        cells(skip_rows=[(1, 2, 3)])


# -- selection: rows / columns / ranges ---------------------------------------------
def selection_book():
    wb = rich_book()
    ws = wb["Sales"]
    ws["C5"].comment = Comment("in C5", "a")
    ws["H40"].hyperlink = "https://h40.example"
    ws.row_dimensions[7].hidden = True
    return wb


ROWS = lambda *spans: {r for a, b in spans for r in range(a, b + 1)}  # noqa: E731
COLS = lambda *spans: {c for a, b in spans for c in range(a, b + 1)}  # noqa: E731

#: (options, keep((row, col)) written independently of the parser, header items of Sales)
SELECTIONS = [
    (dict(rows="2:4,10"), lambda rc: rc[0] in ROWS((2, 4), (10, 10)), ["rows=2:4,10:10"]),
    (dict(rows=[(5, 9), 30, "40:41"]), lambda rc: rc[0] in ROWS((5, 9), (30, 30), (40, 41)), ["rows=5:9,30:30,40:41"]),
    (dict(columns="B:C,E"), lambda rc: rc[1] in COLS((2, 3), (5, 5)), ["cols=B:C,E:E"]),
    (dict(columns=[1, ("D", "F")]), lambda rc: rc[1] in COLS((1, 1), (4, 6)), ["cols=A:A,D:F"]),
    (
        dict(rows="1:20", columns="A:C"),
        lambda rc: rc[0] <= 20 and rc[1] <= 3,
        ["rows=1:20", "cols=A:C"],
    ),
    (
        dict(ranges="B2:C5,E10:F12"),
        lambda rc: (2 <= rc[0] <= 5 and 2 <= rc[1] <= 3) or (10 <= rc[0] <= 12 and 5 <= rc[1] <= 6),
        ["cells=B2:C5,E10:F12"],
    ),
    (dict(ranges=["$A$1", "D:D", "40:41"]), lambda rc: rc == (1, 1) or rc[1] == 4 or 40 <= rc[0] <= 41, None),
    (
        dict(ranges="A1:F30", rows="5:50", columns="B:H"),
        lambda rc: rc[0] <= 30 and rc[1] <= 6 and rc[0] >= 5 and rc[1] >= 2,
        ["rows=5:50", "cols=B:H", "cells=A1:F30"],
    ),
    (dict(ranges="Z1000:Z1001"), lambda rc: False, ["cells=Z1000:Z1001"]),
]


@pytest.mark.parametrize("encoding", ENCODINGS)
@pytest.mark.parametrize("case", range(len(SELECTIONS)))
def test_selection(tmp_path, encoding, case):
    """A selected document is exactly the lossless one restricted to the
    selected cells - values, formulas, styles, hyperlinks, comments"""
    options, keep, header = SELECTIONS[case]
    p = save(selection_book(), tmp_path)
    full = decompress(compress_workbook(p, encoding=encoding))
    doc = compress_workbook(p, encoding=encoding, **options)
    for d in (decompress(doc), decompress(doc.text()), decompress(doc.to_json())):
        assert [s.name for s in d.sheets] == [s.name for s in full.sheets]
        for s in d.sheets:
            f = full[s.name]
            for attr in ("values", "data_types", "formulas", "style_ids", "hyperlinks", "comments"):
                want = {k: v for k, v in getattr(f, attr).items() if keep(k)}
                assert _norm(getattr(s, attr)) == _norm(want), (case, s.name, attr)
            # merged ranges: those holding a selected cell; other metadata whole
            want = [m for m in f.merged if any(keep(c) for c in _range_cells(m))]
            assert s.merged == want and s.row_dimensions == f.row_dimensions
    if header is not None:
        assert decompress(doc)["Sales"].filter == header
    assert "Filtered sheets:" in doc.legend
    kept = sum(1 for k in full["Sales"].values if keep(k))
    assert len(decompress(doc)["Sales"].values) == kept


def test_selection_combined_with_filters(tmp_path):
    """rows / columns / ranges intersect with skip_hidden and the print area"""
    p = save(filter_book(), tmp_path)
    full = decompress(compress_workbook(p))
    ref = openpyxl.load_workbook(p)
    opts = dict(skip_hidden=True, print_area_only=True, rows="1:42", columns="A:D", ranges="A1:E43")
    doc = decompress(compress_workbook(p, **opts))
    for s in doc.sheets:
        base = filter_keep(ref[s.name], True, True)

        def keep(rc):
            return base(rc) and rc[0] <= 42 and rc[1] <= 4

        for attr in ("values", "formulas", "style_ids", "hyperlinks", "comments"):
            want = {k: v for k, v in getattr(full[s.name], attr).items() if keep(k)}
            assert _norm(getattr(s, attr)) == _norm(want), (s.name, attr)
    assert doc["Sales"].filter[:2] == ["skip=hidden", "area=A1:E20,A40:D45+titles:1:2,A:A"]


def test_selection_per_sheet(tmp_path):
    p = save(selection_book(), tmp_path)
    full = decompress(compress_workbook(p))
    doc = compress_workbook(p, ranges={"Sales": "A1:B3"}, rows={"Sparse": "1:100"})
    d = decompress(doc.text())
    assert set(d["Sales"].values) == {k for k in full["Sales"].values if k[0] <= 3 and k[1] <= 2}
    assert set(d["Sparse"].values) == {k for k in full["Sparse"].values if k[0] <= 100}
    # sheets not named are complete and unmarked
    assert d["Table"].values == full["Table"].values and d["Table"].filter == []
    assert d["Sales"].filter == ["cells=A1:B3"] and d["Sparse"].filter == ["rows=1:100"]
    with pytest.raises(KeyError, match="no worksheet named 'Nope'"):
        compress_workbook(p, ranges={"Nope": "A1"})


def test_selection_with_compress_sheets_and_files(tmp_path):
    p = save(selection_book(), tmp_path)
    opts = dict(rows="2:30", columns="A:C")
    for name, doc in compress_sheets(p, **opts):
        assert doc.text() == compress_workbook(p, sheets=name, **opts).text()
    mem = compress_workbook(p, **opts)
    files = compress_workbook(p, output=["text", "json"], path=str(tmp_path / "sel"), **opts).files
    with open(files["text"], encoding="utf-8") as f:
        assert f.read() == mem.text()
    with open(files["json"], encoding="utf-8") as f:
        assert json.load(f) == json.loads(mem.to_json())


def test_selection_noop_and_analysis(tmp_path):
    p = save(dense_numeric(), tmp_path)
    base = compress_workbook(p, include=["values", "analysis"])
    every = compress_workbook(p, include=["values", "analysis"], rows="1:1048576", columns="A:XFD")
    assert decompress(every).sheets[0].values == decompress(base).sheets[0].values
    a = compress_workbook(p, include="analysis", ranges="A1:B2").sheets[0].analysis
    assert a["cells"] == 4 and a["skipped_cells"] == 399 * 12 - 4 and a["dimension"] == "A1:B2"
    assert llm.analyze_workbook(p, ranges="A1:B2")["Sheet"]["cells"] == 4


@pytest.mark.parametrize(
    "options, error",
    [
        (dict(rows="0:3"), ValueError),
        (dict(rows="5:3"), ValueError),
        (dict(rows=""), ValueError),
        (dict(rows=[]), ValueError),
        (dict(rows="a:b"), ValueError),
        (dict(rows=True), TypeError),
        (dict(rows=2.5), TypeError),
        (dict(columns="A:ZZZZ"), ValueError),
        (dict(columns="B:A"), ValueError),
        (dict(columns=0), ValueError),
        (dict(columns="XFE"), ValueError),
        (dict(ranges="Sheet!A1"), ValueError),
        (dict(ranges="A1:"), ValueError),
        (dict(ranges="nonsense"), ValueError),
        (dict(ranges=[(1, 1, 2, 2)]), TypeError),
        (dict(ranges=[]), ValueError),
    ],
)
def test_selection_errors(tmp_path, options, error):
    p = save(dense_numeric(), tmp_path)
    with pytest.raises(error):
        compress_workbook(p, **options)


def test_filter_header_parsing():
    name = 'a" skip=hidden'
    text = (
        "#openrsxl-llm/1 sections=values\n## sheet 1 %s dim=A1:A1 skip=hidden area=A1:B2\n### values index\nA1=1\n"
        % (json.dumps(name))
    )
    s = decompress(text).sheets[0]
    assert s.name == name and s.dimension == "A1:A1" and s.filter == ["skip=hidden", "area=A1:B2"]
    s = decompress("#openrsxl-llm/1\n## sheet 1 %s\n" % json.dumps(name)).sheets[0]
    assert s.name == name and s.filter == []


@pytest.mark.parametrize("encoding", ENCODINGS)
def test_filters_streamed(tmp_path, encoding):
    """Filtered documents streamed to files equal the in-memory ones"""
    p = save(filter_book(), tmp_path)
    kw = dict(encoding=encoding, skip_hidden=True, print_area_only=True, skip_hidden_sheets=True)
    mem = compress_workbook(p, **kw)
    files = compress_workbook(p, output=["text", "json"], path=str(tmp_path / "out"), **kw).files
    with open(files["text"], encoding="utf-8") as f:
        assert f.read() == mem.text()
    with open(files["json"], encoding="utf-8") as f:
        assert json.load(f) == json.loads(mem.to_json())


@pytest.mark.parametrize("encoding", ENCODINGS)
def test_filters_oracle(tmp_path, encoding):
    """The corpus filter check on every synthetic book and raw XML case"""
    paths = [save(b(), tmp_path, f"{b.__name__}.xlsx") for b in (rich_book, filter_book, dense_numeric, sparse_repeats)]
    for name in sorted(RAW):
        paths.append(write(make_xlsx([RAW[name]]), tmp_path, name + ".xlsx"))
    for p in paths:
        for flags in ((True, True, True), (True, False, False), (False, True, False)):
            rep = check_filters(p, Report(), encoding, *flags)
            assert rep, f"{os.path.basename(p)} {flags}: {rep}"


# -- rebuilding --------------------------------------------------------------------
def _comparable(doc):
    book = decompress(doc.text())
    out = []
    for s in book.sheets:
        out.append(
            (
                s.name,
                {k: v for k, v in s.values.items() if k not in s.formulas},
                s.formulas,
                {**s.style_ids},
                s.merged,
                sorted(s.hyperlinks),
                sorted(s.comments),
                [(sq, rule.type) for sq, rule in s.conditional_formatting],
                len(s.validations),
                len(s.tables),
            )
        )
    return out


def test_to_workbook(tmp_path):
    p = save(rich_book(), tmp_path)
    doc = compress_workbook(p)
    wb = decompress(doc).to_workbook()
    p2 = save(wb, tmp_path, "rebuilt.xlsx")
    doc2 = compress_workbook(p2)
    a, b = _comparable(doc), _comparable(doc2)
    for x, y in zip(a, b):
        # style ids may be renumbered: compare the decoded styles instead
        assert x[0] == y[0]
        assert x[1] == y[1], x[0]
        assert _norm(x[2]) == _norm(y[2]), x[0]
        assert x[4:] == y[4:], x[0]
    s1, s2 = decompress(doc)["Sales"], decompress(doc2)["Sales"]
    for key in s1.style_ids:
        a1, a2 = s1.cell(key)["style"], s2.cell(key)["style"]
        for f in ("font", "fill", "border", "alignment", "protection"):
            assert openpyxl.xml.functions.tostring(getattr(a1, f).to_tree()) == openpyxl.xml.functions.tostring(
                getattr(a2, f).to_tree()
            ), (key, f)
        assert a1.number_format == a2.number_format


# -- files ---------------------------------------------------------------------------
UPSTREAM = sorted(
    glob.glob(os.path.join(HERE, "upstream", "**", "*.xlsx"), recursive=True)
    + glob.glob(os.path.join(HERE, "upstream", "**", "*.xlsm"), recursive=True)
)


def _readable(path):
    try:
        openpyxl.load_workbook(path).close()
        openpyxl.load_workbook(path, read_only=True).close()
        return True
    except Exception:
        return False


@pytest.mark.parametrize("path", UPSTREAM, ids=[os.path.relpath(p, HERE) for p in UPSTREAM])
def test_upstream_files(path):
    import warnings

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        if not _readable(path):
            pytest.skip("openpyxl cannot read it")
        check(path)


@pytest.mark.skipif(not os.environ.get("OPENRSXL_LLM_FILES"), reason="set OPENRSXL_LLM_FILES=a.xlsx;b.xlsx")
@pytest.mark.parametrize("path", [p for p in os.environ.get("OPENRSXL_LLM_FILES", "").split(";") if p])
def test_external_files(path):
    check(path)


# -- 0.4.3: dates, clipped merges, style keys, legend parts, cite_cells --------


def _agent_book():
    from openpyxl.styles import Color

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Disposals"
    ws["B1"] = "Entity No"
    ws["C1"] = "Full legal name"
    ws["B5"] = "Johnny Wilde Co-investment 1 LP"
    ws["C4"] = "Income Tax Allocations"
    ws["D5"] = datetime.datetime(2025, 4, 6)  # a date cell: midnight datetime
    ws["D6"] = datetime.datetime(2025, 4, 6, 13, 45)
    ws["D8"] = "2025-04-06"  # text that reads like a date
    for r in range(10, 40):
        for c in range(1, 9):
            ws.cell(r, c, r * c if (r + c) % 3 else None)
    ws["A1"].font = Font(b=True, name="Arial")
    ws["A2"].font = Font(b=True, name="Calibri", family=2, scheme="minor")
    ws["A3"].font = Font(i=True, color=Color(theme=1))
    ws["A4"].fill = PatternFill("solid", fgColor=Color(theme=0, tint=-0.1499984740745262))
    ws["A5"].fill = PatternFill("solid", fgColor=Color(indexed=10))
    for r in range(1, 60, 3):
        ws.merge_cells(f"K{r}:L{r + 1}")
    return wb


def test_midnight_datetimes_keep_their_time(tmp_path):
    p = save(_agent_book(), tmp_path)
    doc = llm.compress_workbook(p, include="values", encoding="grid")
    text = doc["Disposals"]["values"].text
    assert "2025-04-06 00:00:00" in text and "2025-04-06 13:45:00" in text
    assert '"2025-04-06"' in text
    v = decompress(doc)["Disposals"].values
    assert v[(5, 4)] == datetime.datetime(2025, 4, 6) and type(v[(5, 4)]) is datetime.datetime
    assert v[(8, 4)] == "2025-04-06"
    # a date object (openpyxl reads none from cells) keeps its own token
    assert llm.parse_token("@date:2025-04-06") == (datetime.date(2025, 4, 6), "d")
    # documents written before (bare date = midnight datetime) still decode
    assert llm.parse_token("2025-04-06") == (datetime.datetime(2025, 4, 6), "d")
    assert "00:00:00 at" in doc.legend and "2024-01-05 date;" not in doc.legend


def test_merged_ranges_clipped_to_the_selection(tmp_path):
    p = save(_agent_book(), tmp_path)
    whole = llm.compress_workbook(p, include="merged", legend=False)["Disposals"]["merged"].text
    assert len(whole.strip().split(",")) == 20
    part = llm.compress_workbook(p, include="merged", ranges="A1:K7", legend=False)["Disposals"]["merged"].text
    assert sorted(part.strip().split(",")) == ["K1:L2", "K4:L5", "K7:L8"]
    assert "merged" not in llm.compress_workbook(p, include="merged", ranges="A1:C9")["Disposals"].sections


def test_style_keys(tmp_path):
    p = save(_agent_book(), tmp_path)
    full = llm.compress_workbook(p, include="styles")
    doc = llm.compress_workbook(p, include="styles", style_keys=["font.b", "font.i", "fill"])
    table = doc.styles.text()
    assert "name:" not in table and "scheme" not in table and "protection" not in table and "numFmt" not in table
    assert len(doc.styles.cells) < len(full.styles.cells)  # A1 / A2: one style (bold)
    smap = dict(line.split("=") for line in doc["Disposals"]["styles"].text.splitlines())
    assert smap["A1"] == smap["A2"]
    assert "Only these style keys are written: font.b,font.i,fill" in doc.legend
    decompress(doc)  # still a valid document
    with pytest.raises(ValueError, match="unknown key"):
        llm.compress_workbook(p, include="styles", style_keys=["colour"])
    with pytest.raises(ValueError, match="has no part"):
        llm.compress_workbook(p, include="styles", style_keys=["numFmt.code"])
    # without style_keys: the lossless table, unchanged
    assert full.styles.text() == llm.compress_workbook(p, include="styles").styles.text()


def test_resolve_colors(tmp_path):
    p = save(_agent_book(), tmp_path)
    doc = llm.compress_workbook(p, include="styles", resolve_colors=True)
    t = doc.styles.text()
    assert "theme=" not in t and "indexed=" not in t and "tint=" not in t
    assert "rgb=FFD9D9D9" in t  # white (lt1), darker 15 %
    assert "rgb=FF000000" in t  # dk1 (text) of the default theme
    assert "rgb=FFFF0000" in t  # indexed 10: red
    assert "written as rgb=FFRRGGBB" in doc.legend


def test_legend_of_a_document(tmp_path):
    p = save(_agent_book(), tmp_path)
    with_legend = llm.compress_workbook(p, include=["values", "styles", "merged"])
    without = llm.compress_workbook(p, include=["values", "styles", "merged"], legend=False)
    assert without.legend is None and "## legend" not in without.text()
    assert without.legend_parts == with_legend.legend_parts and "merged" in without.legend_parts
    assert llm.legend(without) == with_legend.legend == llm.legend(with_legend.legend_parts)
    keyed = llm.compress_workbook(p, include="styles", style_keys="font", legend=False)
    assert "Only these style keys" in llm.legend(keyed)


def test_cite_cells(tmp_path):
    p = save(_agent_book(), tmp_path)
    plain = llm.compress_workbook(p, include=["values", "formulas"], encoding="grid")
    cite = llm.compress_workbook(p, include=["values", "formulas"], cite_cells=True)
    grid = llm.compress_workbook(p, include=["values", "formulas"], encoding="grid", cite_cells=True)
    # every grid field names its column, every line its row
    for line in grid["Disposals"]["values"].text.splitlines():
        row, _, rest = line.partition(":")
        assert row.isdigit(), line
        for field in rest.split("|"):
            assert field == "" or (field.split("=", 1)[0].isalpha() and "=" in field), line
    assert "5:B=Johnny Wilde Co-investment 1 LP" in grid["Disposals"]["values"].text
    assert "every field is written COL=v" in grid.legend
    # index when it costs at most 1.3 x the grid; never hybrid
    a = cite["Disposals"].analysis
    assert a["encoding"] in ("index", "grid")
    if a["encoding"] == "grid":
        assert a["index_tokens"] is None or a["index_tokens"] > 1.3 * a["grid_tokens"]
    # lossless as always
    for d in (cite, grid):
        got = decompress(d.text())["Disposals"]
        ref = decompress(plain)["Disposals"]
        assert _norm(got.values) == _norm(ref.values) and got.formulas == ref.formulas
