"""
openrsxl.extended streaming features against the oracle (openpyxl's full
mode), on synthetic workbooks covering each feature and its edge cases.
"""

import io
import itertools
import zipfile

import openpyxl
import pytest
from fixtures import MAIN, make_xlsx, sheet_xml
from oracle.extended import ALL_FLAGS, compare_extended

import openrsxl
import openrsxl.extended as ext

FLAG_NAMES = list(ALL_FLAGS)


def add_files(data, files):
    """Add/replace parts of an xlsx package"""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for item in src.infolist():
            if item.filename not in files:
                z.writestr(item, src.read(item.filename))
        for name, content in files.items():
            z.writestr(name, content)
    return out.getvalue()


def openpyxl_book(build):
    wb = openpyxl.Workbook()
    build(wb)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


# -- workbooks written by openpyxl ------------------------------------------


def _rich(wb):
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Border, Font, PatternFill, Protection, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.formula import ArrayFormula
    from openpyxl.worksheet.table import Table

    thick = Side(style="thick", color="FF0000")
    thin = Side(style="thin")
    ws = wb.active
    ws.title = "Main"
    for r in range(1, 40):
        for c in range(1, 12):
            if (r * c) % 3:
                ws.cell(r, c, r * 100 + c)
    ws["A1"] = "title"
    ws["A1"].border = Border(top=thick, left=thick)
    ws["A1"].font = Font(bold=True)
    ws["C2"].border = Border(bottom=thin, right=thick)
    ws.merge_cells("A1:C2")  # start and end cell with borders
    ws["E1"].protection = Protection(locked=False)
    ws["E1"].border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.merge_cells("E1:Z30")  # large merge (representative styles)
    ws.merge_cells("A5:A9")  # single column
    ws.merge_cells("B12:H12")  # single row
    ws["B12"].fill = PatternFill("solid", fgColor="00FF00")
    ws["A20"] = "=SUM(A21:A25)"
    ws["B20"] = "=A20*2"
    ws["C20"] = ArrayFormula("C20:C22", "=A21:A23*2")
    ws["D20"] = '="text"'
    # hyperlinks: on values, empty cells, merged start / inside, outside data
    ws["A30"].hyperlink = "https://example.com/a"
    ws["B30"] = "has value"
    ws["B30"].hyperlink = "https://example.com/b"
    ws["A1"].hyperlink = "https://example.com/merged"
    ws["AC50"].hyperlink = "#Main!A1"
    ws["AD51"].hyperlink = "=not a formula?"
    # comments: normal, on empty cell, outside data
    ws["D3"].comment = Comment("hello", "me")
    ws["AE60"].comment = Comment("far away", "you", width=300, height=50)
    ws["A30"].comment = Comment("with link", "x")
    ws.freeze_panes = "B2"
    ws.print_title_rows = "1:2"
    ws.print_area = "A1:K39"
    ws.row_dimensions[3].height = 30
    ws.row_dimensions[4].hidden = True
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["D"].hidden = True
    ws.protection.sheet = True
    ws.auto_filter.ref = "A14:K39"
    dv = DataValidation(type="list", formula1='"a,b,c"')
    dv.add("K1:K10")
    ws.add_data_validation(dv)
    ws.conditional_formatting.add(
        "A14:K39", CellIsRule(operator="greaterThan", formula=["1000"], font=Font(color="FF0000"))
    )
    ws.sheet_properties.tabColor = "1072BA"

    ws2 = wb.create_sheet("Table")
    ws2.append(["a", "b", "c"])
    for i in range(10):
        ws2.append([i, i * 2, f"=A{i + 2}+B{i + 2}"])
    ws2.add_table(Table(displayName="T1", ref="A1:C11"))
    ws2["E5"].hyperlink = "https://example.com/t"
    ws2.merge_cells("E1:F2")
    ws2["E1"] = "merged"
    wb.create_sheet("Empty")


RICH = openpyxl_book(_rich)


@pytest.mark.parametrize("data_only", [False, True])
def test_rich_all_flags(data_only):
    assert compare_extended(RICH, data_only=data_only) == []


@pytest.mark.parametrize("flag", FLAG_NAMES)
def test_rich_single_flag(flag):
    assert compare_extended(RICH, flags={flag: True}) == []


def test_rich_flag_pairs():
    for a, b in itertools.combinations(
        ["formula_and_value", "read_comments", "read_hyperlinks", "read_merged_cells", "create_empty_cells"], 2
    ):
        assert compare_extended(RICH, flags={a: True, b: True}) == [], (a, b)


# -- API -----------------------------------------------------------------------


def test_defaults_are_openpyxl():
    wb = ext.load_workbook(io.BytesIO(RICH), read_only=True)
    ws = wb["Main"]
    assert type(ws) is openrsxl.worksheet._read_only.ReadOnlyWorksheet
    c = ws["A20"]
    assert type(c) is openrsxl.cell.read_only.ReadOnlyCell
    assert not hasattr(c, "formula")
    assert not hasattr(ws, "merged_cells")


def test_disabled_features_raise_attribute_error():
    wb = ext.load_workbook(io.BytesIO(RICH), read_only=True, read_comments=True)
    c = wb["Main"]["A20"]
    assert c.comment is None
    with pytest.raises(AttributeError, match="formula"):
        _ = c.formula
    with pytest.raises(AttributeError, match="hyperlink"):
        _ = c.hyperlink
    assert c.value == "=SUM(A21:A25)"


def test_formula_and_value_types():
    wb = ext.load_workbook(io.BytesIO(RICH), read_only=True, formula_and_value=True)
    ws = wb["Main"]
    c = ws["A20"]
    assert c.formula == "=SUM(A21:A25)" and c.value == c.formula and c.data_type == "f"
    assert c.cached_value is None  # openpyxl saved no cached values
    assert type(ws["C20"].formula).__name__ == "ArrayFormula"
    v = ws["B4"]
    assert v.formula is None and v.cached_value == v.value == 402
    empty = ws["ZZ999"]
    assert empty.value is None and empty.formula is None and empty.cached_value is None
    wb2 = ext.load_workbook(io.BytesIO(RICH), read_only=True, data_only=True, formula_and_value=True)
    c = wb2["Main"]["A20"]
    assert c.value is None and c.data_type == "n" and c.formula == "=SUM(A21:A25)"


def test_formula_and_value_requires_read_only():
    with pytest.raises(ValueError):
        ext.load_workbook(io.BytesIO(RICH), formula_and_value=True)


def test_full_mode_flags_are_noops():
    wb = ext.load_workbook(io.BytesIO(RICH), read_comments=True, read_merged_cells=True, create_empty_cells=True)
    assert type(wb["Main"]).__name__ == "Worksheet"


def test_signature_extends_openpyxl():
    import inspect

    base = inspect.signature(openpyxl.load_workbook).parameters
    sig = inspect.signature(ext.load_workbook).parameters
    assert list(sig)[: len(base)] == list(base)
    for name in FLAG_NAMES:
        assert sig[name].kind is inspect.Parameter.KEYWORD_ONLY
        assert sig[name].default is False
    assert ext.open is ext.load_workbook


def test_create_empty_cells():
    # full mode creates a Cell (coordinate, default style) when an empty
    # position is accessed; read-only mode yields the shared EmptyCell
    from oracle.compare import norm  # (ArrayFormula objects have no __eq__, as in openpyxl)

    plain = ext.load_workbook(io.BytesIO(RICH), read_only=True)["Main"]
    wb = ext.load_workbook(io.BytesIO(RICH), read_only=True, create_empty_cells=True)
    ws = wb["Main"]
    full = openrsxl.load_workbook(io.BytesIO(RICH))["Main"]
    n_styles = len(wb._cell_styles)
    empty = 0
    for prow, row in zip(plain.iter_rows(), ws.iter_rows()):
        assert len(prow) == len(row)  # read-only mode's row shape
        for p, c in zip(prow, row):
            if type(p).__name__ != "EmptyCell":
                assert c.coordinate == p.coordinate and norm(c.value) == norm(p.value)
                continue
            assert type(c).__name__ == "ReadOnlyCell" and c.parent is ws
            if (c.row, c.column) in full._cells:
                continue  # eg. a MergedCell: merges are off here (the oracle covers them)
            empty += 1
            f = full.cell(c.row, c.column)  # created on access
            assert (c.coordinate, c.value, c.data_type, c.is_date, c.has_style) == (
                f.coordinate,
                None,
                "n",
                False,
                False,
            )
            for attr in ("number_format", "font", "fill", "border", "alignment", "protection"):
                # (full mode's StyleProxy compares its target: left operand)
                assert getattr(f, attr) == getattr(c, attr), attr
            with pytest.raises(AttributeError):
                _ = c.formula  # disabled features do not exist
    assert empty > 100
    assert len(wb._cell_styles) == n_styles  # the style registry is untouched
    assert norm(list(ws.iter_rows(values_only=True))) == norm(list(plain.iter_rows(values_only=True)))
    # single cells, inside and beyond the rows of the file
    assert type(plain["A3"]).__name__ == "EmptyCell"
    assert ws["A3"].coordinate == "A3" and ws["A3"] is not ws["A3"]  # new objects, nothing kept
    far = ws.cell(1000, 100)
    assert (far.coordinate, far.value, far.row, far.column) == ("CV1000", None, 1000, 100)
    # rows keep read-only mode's shape: none after the last row of the file
    assert len(list(ws.iter_rows(max_row=500))) == len(list(plain.iter_rows(max_row=500))) == 60
    # with the cell features: their attributes exist, all None
    ws = ext.load_workbook(
        io.BytesIO(RICH),
        read_only=True,
        create_empty_cells=True,
        **dict.fromkeys(["formula_and_value", "read_comments", "read_hyperlinks"], True),
    )["Main"]
    c = ws["A3"]
    assert (c.formula, c.cached_value, c.comment, c.hyperlink) == (None, None, None, None)


def test_values_only_with_merges():
    wb = ext.load_workbook(io.BytesIO(RICH), read_only=True, read_merged_cells=True)
    rows = list(wb["Main"].iter_rows(min_row=1, max_row=2, max_col=4, values_only=True))
    assert rows[0][:3] == ("title", None, None)


def test_reiteration_is_stable():
    wb = ext.load_workbook(io.BytesIO(RICH), read_only=True, **ALL_FLAGS)
    ws = wb["Main"]
    from oracle.compare import norm

    first = [[(type(c).__name__, norm(c.value), getattr(c, "_style_id", None)) for c in r] for r in ws.iter_rows()]
    second = [[(type(c).__name__, norm(c.value), getattr(c, "_style_id", None)) for c in r] for r in ws.iter_rows()]
    assert first == second
    n = len(wb._cell_styles)
    list(ws.iter_rows())
    assert len(wb._cell_styles) == n


# -- raw XML fixtures ----------------------------------------------------------

SHEET_RELS = (
    "xl/worksheets/_rels/sheet1.xml.rels",
    '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://example.com/x" TargetMode="External"/>'
    "</Relationships>",
)


def raw(sheet_data, before="", after="", rels=True):
    data = make_xlsx(raw_sheets=[sheet_xml(sheet_data, before=before, after=after)])
    if rels:
        data = add_files(data, dict([SHEET_RELS]))
    return data


CELLS = (
    '<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" s="3"><v>1</v></c><c r="C1"><f>B1*2</f><v>2</v></c></row>'
    '<row r="2" ht="30" customHeight="1"><c r="A2"><f t="shared" ref="A2:A4" si="0">B1+1</f><v>2</v></c><c r="B2" s="6"><v>3.5</v></c></row>'
    '<row r="3"><c r="A3"><f t="shared" si="0"/><v>3</v></c></row>'
    '<row r="4" hidden="1"><c r="A4"><f t="shared" si="0"/><v>4</v></c><c r="D4" t="str"><f t="array" ref="D4:D5">B1:B2</f><v>x</v></c></row>'
    '<row r="6"><c r="A6"><f t="dataTable" ref="A6:B7" dt2D="0" dtr="1" r1="A1"/><v>9</v></c></row>'
)

RAW_CASES = {
    "plain": raw(CELLS),
    "dimension_small": raw(CELLS, before='<dimension ref="A1:B2"/>'),
    "no_dimension_links_outside": raw(
        CELLS,
        after='<hyperlinks><hyperlink ref="H20" r:id="rId1"/><hyperlink ref="A1" location="Sheet1!B2" display="x"/></hyperlinks>',
    ),
    "link_range": raw(
        CELLS,
        before='<dimension ref="A1:D6"/>',
        after='<mergeCells count="1"><mergeCell ref="B2:C3"/></mergeCells><hyperlinks><hyperlink ref="A1:C4" r:id="rId1" tooltip="tip"/></hyperlinks>',
    ),
    "link_merged_inside": raw(
        CELLS,
        before='<dimension ref="A1:D6"/>',
        after='<mergeCells count="1"><mergeCell ref="B2:C3"/></mergeCells><hyperlinks><hyperlink ref="C3" location="Sheet1!A1"/></hyperlinks>',
    ),
    "links_same_cell": raw(
        CELLS,
        after='<hyperlinks><hyperlink ref="E1" location="first"/><hyperlink ref="E1" r:id="rId1"/><hyperlink ref="E1:F1" location="range"/></hyperlinks>',
    ),
    "overlapping_merges": raw(
        CELLS,
        before='<dimension ref="A1:D6"/>',
        after='<mergeCells count="3"><mergeCell ref="A1:B3"/><mergeCell ref="B2:D4"/><mergeCell ref="C4:C6"/></mergeCells>',
    ),
    "merge_outside_data": raw(
        CELLS, before='<dimension ref="A1:D6"/>', after='<mergeCells count="1"><mergeCell ref="X10:Z40"/></mergeCells>'
    ),
    "merge_formula_cells": raw(CELLS, after='<mergeCells count="1"><mergeCell ref="A2:B4"/></mergeCells>'),
    "dims_and_props": raw(
        CELLS,
        before='<sheetPr><tabColor rgb="FF00FF00"/><pageSetUpPr fitToPage="1"/></sheetPr><dimension ref="A1:D6"/>'
        '<sheetViews><sheetView workbookViewId="0"><pane xSplit="1" ySplit="1" topLeftCell="B2" activePane="bottomRight" state="frozen"/></sheetView></sheetViews>'
        '<sheetFormatPr defaultRowHeight="15" baseColWidth="10"/><cols><col min="1" max="3" width="20" customWidth="1" style="3"/></cols>',
        after='<sheetProtection sheet="1" objects="1"/><autoFilter ref="A1:B4"/>'
        '<dataValidations count="1"><dataValidation type="whole" sqref="A1:A4"><formula1>1</formula1><formula2>9</formula2></dataValidation></dataValidations>'
        '<printOptions gridLines="1"/><pageMargins left="0.5" right="0.5" top="1" bottom="1" header="0.3" footer="0.3"/>'
        '<pageSetup orientation="landscape"/><headerFooter><oddHeader>&amp;CTitle</oddHeader></headerFooter>'
        '<rowBreaks count="1" manualBreakCount="1"><brk id="3" max="16383" man="1"/></rowBreaks>',
    ),
    "row_attrs": raw(
        '<row r="1" spans="1:3"><c r="A1"><v>1</v></c></row><row ht="20" customHeight="1"><c><v>2</v></c></row>'
        '<row r="5" s="3" customFormat="1" x14ac:dyDescent="0.25" xmlns:x14ac="http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac"><c r="A5"><v>5</v></c></row>'
    ),
    "cells_without_coordinates": raw(
        '<row><c><v>1</v></c><c s="3"><v>2</v></c></row><row><c/><c s="6"><v>3</v></c></row>',
        after='<mergeCells count="1"><mergeCell ref="A1:B2"/></mergeCells>',
    ),
    "empty_sheet_with_features": raw(
        "",
        after='<mergeCells count="1"><mergeCell ref="B2:C3"/></mergeCells><hyperlinks><hyperlink ref="A1" r:id="rId1"/></hyperlinks>',
    ),
    "comment_like_in_sheetdata": raw(
        '<!-- note --><row r="1"><c r="A1"><v>1</v></c></row>',
        after='<mergeCells count="1"><mergeCell ref="A1:B1"/></mergeCells>',
    ),
    "prefixed_main_ns": make_xlsx(
        raw_sheets=[
            f'<x:worksheet xmlns:x="{MAIN}"><x:sheetData><x:row r="1"><x:c r="A1"><x:v>1</x:v></x:c></x:row></x:sheetData><x:mergeCells><x:mergeCell ref="A1:B2"/></x:mergeCells></x:worksheet>'
        ]
    ),
    "attr_syntax": raw(
        '<row r = \' 2\'\n\tht="1&#50;" customHeight=\'1\' spans="1:2"><c r="A2" s = "3"><v>1</v></c><c\nr="B2"\ts="6"/></row>'
        '<row r="3.0" thickBot="1"><c r="C3" s="1"><v>2</v></c></row><row r="+4" x="a&amp;b&lt;&#x41;"/>',
        after='<mergeCells count="2"><mergeCell ref="A2:B3"/><mergeCell ref="C3:D4"/></mergeCells>',
    ),
    "odd_coordinates": raw(
        '<row r="1"><c r="a1" s="3"><v>1</v></c><c r="b1"><v>2</v></c></row>',
        after='<mergeCells count="1"><mergeCell ref="A1:B2"/></mergeCells>',
    ),
    "cdata_in_sheetdata": raw(
        '<row r="1"><c r="A1" t="inlineStr"><is><t><![CDATA[<c r="Z9" s="5">]]></t></is></c></row>',
        after='<mergeCells count="1"><mergeCell ref="A1:B2"/></mergeCells>',
    ),
    "fake_tags_in_text": raw(
        '<row r="1"><c r="A1" t="inlineStr"><is><t>&lt;c r="B1" s="5"&gt;</t></is></c><c r="B1" s="3"/></row>',
        after='<mergeCells count="1"><mergeCell ref="A1:B1"/></mergeCells>',
    ),
    "self_closing_sheetdata_tail": raw(
        "",
        after='<hyperlinks><hyperlink ref="B2" location="x"/></hyperlinks><dataValidations count="1"><dataValidation sqref="A1"/></dataValidations>',
    ),
    "errors_and_bools": raw(
        '<row r="1"><c r="A1" t="e"><f>1/0</f><v>#DIV/0!</v></c><c r="B1" t="b"><f>TRUE()</f><v>1</v></c><c r="C1" s="1"><f>TODAY()</f><v>45000</v></c><c r="D1" t="str"><f>"a"</f><v>a</v></c><c r="E1" t="inlineStr"><is><t>inline</t></is></c></row>'
    ),
}


@pytest.mark.parametrize("data_only", [False, True])
@pytest.mark.parametrize("case", sorted(RAW_CASES))
def test_raw_cases(case, data_only):
    assert compare_extended(RAW_CASES[case], data_only=data_only) == []


def test_comment_on_merged_cell_warns():
    def build(wb):
        from openpyxl.comments import Comment

        ws = wb.active
        ws["B2"].comment = Comment("lost", "me")
        ws["A1"] = 1

    data = openpyxl_book(build)
    # merge after writing: comment sits on a merged cell
    data = add_files(
        data,
        {
            "xl/worksheets/sheet1.xml": zipfile.ZipFile(io.BytesIO(data))
            .read("xl/worksheets/sheet1.xml")
            .replace(b"</sheetData>", b'</sheetData><mergeCells count="1"><mergeCell ref="A1:C3"/></mergeCells>')
        },
    )
    assert compare_extended(data) == []
    with pytest.warns(UserWarning, match="merged range but has a comment"):
        ext.load_workbook(io.BytesIO(data), read_only=True, read_comments=True, read_merged_cells=True)


# -- random workbooks ------------------------------------------------------------


def _random_book(seed):
    import random

    from openpyxl.comments import Comment
    from openpyxl.styles import Border, Protection, Side

    rnd = random.Random(seed)
    sides = [None, Side(style="thin"), Side(style="thick", color="FF0000")]

    def build(wb):
        ws = wb.active
        nrows, ncols = rnd.randint(1, 30), rnd.randint(1, 12)
        for _ in range(rnd.randint(0, 120)):
            r, c = rnd.randint(1, nrows), rnd.randint(1, ncols)
            kind = rnd.random()
            if kind < 0.4:
                ws.cell(r, c, rnd.randint(-5, 500))
            elif kind < 0.6:
                ws.cell(r, c, rnd.choice(["x", "", " y ", "=A1", "#REF!"]))
            elif kind < 0.8:
                ws.cell(r, c, f"=SUM(A1:B{rnd.randint(1, 9)})")
            else:
                ws.cell(r, c).border = Border(
                    left=rnd.choice(sides), right=rnd.choice(sides), top=rnd.choice(sides), bottom=rnd.choice(sides)
                )
                if rnd.random() < 0.3:
                    ws.cell(r, c).protection = Protection(locked=False)
        coords = lambda: (rnd.randint(1, nrows + 3), rnd.randint(1, ncols + 3))
        for _ in range(rnd.randint(0, 6)):
            r, c = coords()
            ws.cell(r, c).hyperlink = rnd.choice(["https://e.com/" + str(r), f"#Sheet!A{r}"])
        for _ in range(rnd.randint(0, 4)):
            r, c = coords()
            ws.cell(r, c).comment = Comment(f"c{r}", "a")
        merges = []
        for _ in range(rnd.randint(0, 6)):
            r, c = coords()
            merges.append(
                f"{ws.cell(r, c).coordinate}:{ws.cell(r + rnd.randint(0, 12), c + rnd.randint(0, 12)).coordinate}"
            )
        return merges

    holder = {}

    def b(wb):
        holder["merges"] = build(wb)

    data = openpyxl_book(b)
    # merges are inserted in the XML (openpyxl would drop the covered cells
    # when writing, real files keep values / styles / links under merges)
    sheet = zipfile.ZipFile(io.BytesIO(data)).read("xl/worksheets/sheet1.xml")
    if holder["merges"]:
        mc = "".join(f'<mergeCell ref="{m}"/>' for m in holder["merges"]).encode()
        sheet = sheet.replace(b"</sheetData>", b"</sheetData><mergeCells>" + mc + b"</mergeCells>", 1)
        if b"</sheetData>" not in sheet:
            sheet = sheet.replace(b"<sheetData/>", b"<sheetData/><mergeCells>" + mc + b"</mergeCells>", 1)
    return add_files(data, {"xl/worksheets/sheet1.xml": sheet})


@pytest.mark.parametrize("seed", range(60))
def test_random_books(seed):
    data = _random_book(seed)
    assert compare_extended(data, data_only=bool(seed % 2)) == []


def _sig(c):
    from oracle.compare import norm

    if c is None or not hasattr(c, "data_type"):
        return norm(c)
    return (
        type(c).__name__,
        norm(c.value),
        c.data_type,
        norm(getattr(c, "formula", None)),
        norm(getattr(c, "cached_value", None)),
        norm(getattr(c, "hyperlink", None)),
        repr(getattr(c, "comment", None)),
        tuple(c._style) if hasattr(c, "_style") else tuple(getattr(c, "style_array", ())),
    )


@pytest.mark.parametrize("seed", range(12))
def test_sub_rectangles(seed):
    import random

    rnd = random.Random(seed)
    data = RICH if seed == 0 else _random_book(seed + 1000)
    try:
        wb = ext.load_workbook(io.BytesIO(data), read_only=True, **ALL_FLAGS)
    except AttributeError:
        return  # overlapping merges: openpyxl fails as well (oracle tests)
    ws = wb.worksheets[0]
    full = [list(r) for r in ws.iter_rows()]
    full_v = [list(r) for r in ws.iter_rows(values_only=True)]
    if not full:
        return
    for _ in range(15):
        r0 = rnd.randint(1, len(full))
        r1 = rnd.randint(r0, len(full))
        c0 = rnd.randint(1, len(full[0]))
        c1 = rnd.randint(c0, len(full[0]))
        sub = list(ws.iter_rows(min_row=r0, max_row=r1, min_col=c0, max_col=c1))
        assert [[_sig(c) for c in r] for r in sub] == [[_sig(c) for c in r[c0 - 1 : c1]] for r in full[r0 - 1 : r1]]
        sub_v = list(ws.iter_rows(min_row=r0, max_row=r1, min_col=c0, max_col=c1, values_only=True))
        assert [[_sig(c) for c in r] for r in sub_v] == [[_sig(c) for c in r[c0 - 1 : c1]] for r in full_v[r0 - 1 : r1]]
        assert _sig(ws.cell(r0, c0)) == _sig(full[r0 - 1][c0 - 1])


@pytest.mark.parametrize("case", sorted(RAW_CASES) + ["rich"] + [f"random{i}" for i in range(20)])
def test_prescan_matches_python_scan(case):
    from oracle.compare import norm

    from openrsxl.extended import _streaming as S

    if case == "rich":
        data = RICH
    elif case.startswith("random"):
        data = _random_book(int(case[6:]) + 500)
    else:
        data = RAW_CASES[case]
    wb = ext.load_workbook(io.BytesIO(data), read_only=True)
    ws = wb.worksheets[0]
    corners = {(r, c) for r in range(1, 12) for c in range(1, 8)}
    native = S._scan(ws, want_rows=True, corners=corners)
    python = S._python_scan(ws, corners)
    for attr in ("row_dimensions", "column_dimensions", "_corner_styles"):
        assert getattr(native, attr) == getattr(python, attr), attr
    for attr in (
        "merged_cells",
        "hyperlinks",
        "formatting",
        "tables",
        "views",
        "sheet_properties",
        "print_options",
        "page_margins",
        "page_setup",
        "HeaderFooter",
        "auto_filter",
        "data_validations",
        "sheet_format",
        "row_breaks",
        "col_breaks",
        "scenarios",
        "protection",
        "legacy_drawing",
    ):
        assert norm(getattr(native, attr, None)) == norm(getattr(python, attr, None)), attr


def test_cell_errors_raised_while_iterating():
    # full mode fails in load_workbook; streaming (openpyxl read-only and
    # openrsxl.extended) fails while iterating, with the same error
    data = raw(
        '<row r="1"><c r="$A$1" s="3"><v>1</v></c></row>',
        after='<mergeCells count="1"><mergeCell ref="A1:B2"/></mergeCells>',
    )
    with pytest.raises(ValueError):
        openpyxl.load_workbook(io.BytesIO(data))

    def err(lib, **kw):
        wb = lib.load_workbook(io.BytesIO(data), read_only=True, **kw)
        try:
            list(wb.active.iter_rows())
        except Exception as e:
            return type(e).__name__, str(e)

    expected = err(openpyxl)
    assert expected is not None
    assert err(ext, **ALL_FLAGS) == expected


def test_merged_start_cell():
    import copy

    from openrsxl.worksheet.merge import MergedCellRange

    wb = ext.load_workbook(io.BytesIO(RICH), read_only=True, read_merged_cells=True)
    ws = wb["Main"]
    full = openpyxl.load_workbook(io.BytesIO(RICH))["Main"]
    rng = {r.coord: r for r in ws.merged_cells}["A1:C2"]
    assert isinstance(rng, MergedCellRange) and type(rng).__name__ == "MergedCellRange"
    c = rng.start_cell
    exp = {r.coord: r for r in full.merged_cells}["A1:C2"].start_cell
    assert c.coordinate == "A1" and c.value == exp.value == "title"
    # borders of the range: top/left of A1, bottom/right of C2
    assert c.border.top.style == exp.border.top.style == "thick"
    assert c.border.right.style == exp.border.right.style == "thick"
    assert c.border.bottom.style == exp.border.bottom.style == "thin"
    assert rng.start_cell is c  # cached
    assert next(ws.iter_rows(min_row=1, max_row=1, max_col=1))[0].border == c.border
    # created when the file has no top-left cell
    assert {r.coord: r for r in ws.merged_cells}["B12:H12"].start_cell.value is None
    assert str(copy.copy(rng)) == "A1:C2" and copy.copy(rng).start_cell.value == "title"
