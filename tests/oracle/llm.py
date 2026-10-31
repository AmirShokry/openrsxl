"""
Oracle for openrsxl.extended.llm: compress a workbook, decode the document
and compare everything with what *openpyxl* (the real one) reads.

* values (type and value): openpyxl read-only, ``data_only=True``, cells at
  their own coordinates, the file's dimension ignored;
* formulas: openpyxl read-only, ``data_only=False`` (``data_type == 'f'``);
* styles of every cell present in the file: font, fill, border, alignment,
  protection (XML), number format, quotePrefix, pivotButton, named style;
* metadata (``full=True``): openpyxl's full mode objects - merged ranges,
  hyperlinks and comments per cell, column / row dimensions, data
  validation, conditional formatting, tables, sheet settings, workbook
  settings (compared as the XML openpyxl writes for them).

Every output form is checked: the text document, the JSON document and the
files written with output="all" must decode to the same workbook.

CLI::

    python tests/oracle/llm.py file.xlsx [--no-full] [--encoding grid|index|hybrid]
"""

import json
import math
import os
import sys
import tempfile
import warnings

import openpyxl
from openpyxl.xml.functions import tostring

from openrsxl.extended.llm import compress_sheets, compress_workbook, decompress

ALL = "all"


def _xml(obj):
    if obj is None:
        return None
    el = obj.to_tree()
    return None if el is None else tostring(el)


def _same_value(a, b):
    if type(a) is not type(b):
        return False
    if isinstance(a, float) and math.isnan(a) and math.isnan(b):
        return True
    return a == b


class Report:
    def __init__(self, limit=20):
        self.errors = []
        self.limit = limit
        self.counts = {}

    def add(self, where, msg):
        self.counts[where] = self.counts.get(where, 0) + 1
        if len(self.errors) < self.limit:
            self.errors.append(f"{where}: {msg}")

    def __bool__(self):
        return not self.counts

    def __str__(self):
        if not self.counts:
            return "ok"
        return "\n".join(self.errors) + f"\n(total by kind: {self.counts})"


def _rows(path, data_only, full_cells, names):
    """(sheet name, {(row, col): cell}) per row of the file - streamed in
    read-only mode (later cells of a row win, as in full mode)"""
    if full_cells:
        wb = openpyxl.load_workbook(path, data_only=data_only)
        for ws in wb.worksheets:
            if names is None or ws.title in names:
                rows = {}
                for k, c in ws._cells.items():
                    rows.setdefault(k[0], {})[k] = c
                for r in sorted(rows):
                    yield wb, ws.title, rows[r]
        wb.close()
        return
    wb = openpyxl.load_workbook(path, read_only=True, data_only=data_only)
    for ws in wb.worksheets:
        if names is not None and ws.title not in names:
            continue
        ws.reset_dimensions()
        for row in ws.iter_rows():
            cells = {}
            for c in row:
                if c.__class__.__name__ != "EmptyCell":
                    cells[(c.row, c.column)] = c
            if cells:
                yield wb, ws.title, cells
    wb.close()


def _sa(c):
    sa = getattr(c, "_style", None)
    return sa if sa is not None else c.style_array


def _ro_style(wb, c):
    sa = _sa(c)
    names = wb._named_styles.names
    return dict(
        font=_xml(c.font),
        fill=_xml(c.fill),
        border=_xml(c.border),
        alignment=_xml(c.alignment),
        protection=_xml(c.protection),
        number_format=c.number_format,
        quotePrefix=bool(sa.quotePrefix),
        pivotButton=bool(sa.pivotButton),
        style=names[sa.xfId] if sa.xfId < len(names) else None,
    )


def _decoded_style(sd):
    return dict(
        font=_xml(sd.font),
        fill=_xml(sd.fill),
        border=_xml(sd.border),
        alignment=_xml(sd.alignment),
        protection=_xml(sd.protection),
        number_format=sd.number_format,
        quotePrefix=bool(sd.quotePrefix),
        pivotButton=bool(sd.pivotButton),
        style=sd.style,
    )


def _full_cells(path, data_only):
    """cells of full mode (all rows, also out of order ones that read-only
    mode skips); only for files without merged cells / hyperlinks"""
    wb = openpyxl.load_workbook(path, data_only=data_only)
    return wb, {ws.title: dict(ws._cells) for ws in wb.worksheets}


_FULL = {}


def _full_mode_cells(path, data_only):
    """{sheet: {(row, col): cell}} of openpyxl's full mode, without the
    MergedCell placeholders (cached per file)"""
    key = (path, data_only)
    if key not in _FULL:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(path, data_only=data_only)
        _FULL.clear()
        _FULL[key] = {
            ws.title: {k: c for k, c in ws._cells.items() if c.__class__.__name__ != "MergedCell"}
            for ws in wb.worksheets
        }
    return _FULL[key]


def check_cells(path, book, rep, sections, full_cells=False):
    sheets = {s.name: s for s in book.sheets}
    names = set(sheets)
    if "values" in sections:
        seen = {n: set() for n in names}
        for _, name, cells in _rows(path, True, full_cells, names):
            s = sheets[name]
            for k, c in cells.items():
                if c.value is None:
                    continue
                seen[name].add(k)
                if k not in s.values:
                    rep.add("values", f"{name}!{k}: missing (expected {c.value!r})")
                elif not _same_value(s.values[k], c.value):
                    rep.add("values", f"{name}!{k}: {s.values[k]!r} != {c.value!r}")
                elif s.data_types[k] != c.data_type:
                    rep.add("data_type", f"{name}!{k}: {s.data_types[k]!r} != {c.data_type!r}")
        extra = {name: set(sheets[name].values) - seen[name] for name in names}
        if any(extra.values()):
            # cells read-only mode skips (rows out of order, row 0): full mode
            # has them
            full = _full_mode_cells(path, True)
            for name, keys in extra.items():
                for k in keys:
                    c = full.get(name, {}).get(k)
                    v = sheets[name].values[k]
                    if c is None or not _same_value(v, c.value) or sheets[name].data_types[k] != c.data_type:
                        rep.add("values", f"{name}!{k}: unexpected {v!r}")
        del seen
    if "formulas" in sections or "styles" in sections:
        fseen = {n: set() for n in names}
        present = {n: set() for n in names}
        cache = {}
        for wb, name, cells in _rows(path, False, full_cells, names):
            s = sheets[name]
            for k, c in cells.items():
                if "formulas" in sections and c.data_type == "f":
                    f = c.value
                    fseen[name].add(k)
                    got = s.formulas.get(k)
                    if isinstance(f, str):
                        ok = got == f
                    else:
                        ok = type(got).__name__ == type(f).__name__ and dict(got) == dict(f)
                        ok = ok and getattr(got, "text", None) == getattr(f, "text", None)
                    if not ok:
                        rep.add("formulas", f"{name}!{k}: {got!r} != {f!r}")
                if "styles" in sections:
                    if k in s.style_ids:
                        present[name].add(k)
                    sid = s.style_ids.get(k, 0)
                    key = (name, tuple(_sa(c)), sid)
                    if key in cache:
                        continue
                    cache[key] = True
                    sd = s.styles.get(sid)
                    if sd is None:
                        rep.add("styles", f"{name}!{k}: style s{sid} not defined")
                        continue
                    exp = _ro_style(wb, c)
                    got = _decoded_style(sd)
                    for field in exp:
                        if exp[field] != got[field]:
                            rep.add("styles", f"{name}!{k} {field}: {got[field]!r} != {exp[field]!r}")
        for name in names:
            s = sheets[name]
            if "formulas" in sections:
                extra = set(s.formulas) - fseen[name]
                full = _full_mode_cells(path, False) if extra else {}
                for k in extra:
                    c = full.get(name, {}).get(k)
                    if c is None or c.data_type != "f" or _norm({0: c.value}) != _norm({0: s.formulas[k]}):
                        rep.add("formulas", f"{name}!{k}: unexpected {s.formulas[k]!r}")
            if "styles" in sections:
                extra = set(s.style_ids) - present[name]
                full = _full_mode_cells(path, False) if extra else {}
                for k in extra:
                    if k not in full.get(name, {}):
                        rep.add("styles", f"{name}!{k}: styled but not in the file")


def check_metadata(path, book, rep, sections, only=None):
    """`only`: the names of the worksheets the document was made of (None:
    all of them)"""
    if isinstance(only, str):
        only = [only]
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path)
    sheets = {s.name: s for s in book.sheets}
    for ws in wb.worksheets:
        if only is not None and ws.title not in only:
            if ws.title in sheets:
                rep.add(ws.title, "a sheet that was not selected is in the document")
            continue
        s = sheets[ws.title]
        if "merged" in sections:
            exp = [str(r.coord) for r in ws.merged_cells.ranges]
            if exp != s.merged:
                rep.add("merged", f"{ws.title}: {s.merged[:5]} != {exp[:5]}")
        if "hyperlinks" in sections or "comments" in sections:
            links, comments = {}, {}
            for (r, c), cell in ws._cells.items():
                if getattr(cell, "hyperlink", None) is not None:
                    links[(r, c)] = cell.hyperlink
                if getattr(cell, "comment", None) is not None:
                    comments[(r, c)] = cell.comment
            if "hyperlinks" in sections:
                if set(links) != set(s.hyperlinks):
                    rep.add("hyperlinks", f"{ws.title}: cells {sorted(set(links) ^ set(s.hyperlinks))[:5]}")
                for k, link in links.items():
                    got = s.hyperlinks.get(k)
                    for a in ("ref", "target", "location", "tooltip", "display"):
                        if got is not None and getattr(got, a) != getattr(link, a):
                            rep.add("hyperlinks", f"{ws.title}!{k}.{a}: {getattr(got, a)!r} != {getattr(link, a)!r}")
            if "comments" in sections:
                if set(comments) != set(s.comments):
                    rep.add("comments", f"{ws.title}: cells {sorted(set(comments) ^ set(s.comments))[:5]}")
                for k, cm in comments.items():
                    got = s.comments.get(k)
                    if got is None:
                        continue
                    for a in ("content", "author", "width", "height"):
                        if getattr(got, a) != getattr(cm, a):
                            rep.add("comments", f"{ws.title}!{k}.{a}: {getattr(got, a)!r} != {getattr(cm, a)!r}")
        if "dimensions" in sections:
            styles = wb._cell_styles

            first = {}
            for i, sa in enumerate(styles):
                first.setdefault(tuple(sa), i)

            def sid_of(d):
                st = d._style
                return first[tuple(st)] if st is not None and any(st) else None

            for key, cd in ws.column_dimensions.items():
                got = s.column_dimensions.get(key)
                if got is None:
                    rep.add("dimensions", f"{ws.title} col {key}: missing")
                    continue
                for f in ("min", "max", "width", "bestFit", "hidden", "outlineLevel", "collapsed"):
                    if not _same_value(got[f], getattr(cd, f)):
                        rep.add("dimensions", f"{ws.title} col {key}.{f}: {got[f]!r} != {getattr(cd, f)!r}")
                exp_style = sid_of(cd)
                if exp_style is not None:
                    gs = s.styles.get(got["style_id"])
                    if gs is None or tuple(styles[got["style_id"]]) != tuple(cd._style):
                        rep.add("dimensions", f"{ws.title} col {key}.style")
            if set(s.column_dimensions) - set(ws.column_dimensions):
                rep.add("dimensions", f"{ws.title}: unexpected columns")
            for idx, rd in ws.row_dimensions.items():
                got = s.row_dimensions.get(idx)
                if got is None:
                    # full mode creates empty row dimensions on access
                    if any(dict(rd).values()) or rd.ht is not None:
                        rep.add("dimensions", f"{ws.title} row {idx}: missing")
                    continue
                for f in ("ht", "hidden", "outlineLevel", "collapsed", "thickBot", "thickTop"):
                    if not _same_value(got[f], getattr(rd, f)):
                        rep.add("dimensions", f"{ws.title} row {idx}.{f}: {got[f]!r} != {getattr(rd, f)!r}")
                exp_style = sid_of(rd)
                if (exp_style is None) != (got["style_id"] is None):
                    rep.add("dimensions", f"{ws.title} row {idx}.style: {got['style_id']} != {exp_style}")
                elif exp_style is not None and tuple(styles[got["style_id"]]) != tuple(rd._style):
                    rep.add("dimensions", f"{ws.title} row {idx}.style")
        if "validations" in sections:
            exp = [_xml(dv) for dv in ws.data_validations.dataValidation]
            got = [_xml(dv) for dv in s.validations]
            if exp != got:
                rep.add("validations", f"{ws.title}: {got[:2]} != {exp[:2]}")
        if "conditional_formatting" in sections:
            exp = []
            for cf in ws.conditional_formatting:
                for rule in cf.rules:
                    exp.append((str(cf.sqref), _rule_xml(rule), _xml(rule.dxf)))
            got = [(sq, _rule_xml(rule), _xml(rule.dxf)) for sq, rule in s.conditional_formatting]
            if exp != got:
                rep.add("conditional_formatting", f"{ws.title}: {got[:2]} != {exp[:2]}")
        if "tables" in sections:
            exp = [_xml(t) for t in ws.tables.values()]
            got = [_xml(t) for t in s.tables]
            if exp != got:
                rep.add("tables", f"{ws.title}: {got[:1]} != {exp[:1]}")
        if "sheet" in sections:
            from openrsxl.extended.llm._model import SHEET_OBJECTS

            for attr, _cls, factory in SHEET_OBJECTS:
                exp = getattr(ws, attr)
                got = s.sheet.get(attr)
                if got is None:
                    # not listed: openpyxl's default
                    d = factory()
                    if _xml(exp) != _xml(d) and _xml(exp) is not None:
                        rep.add("sheet", f"{ws.title}.{attr}: missing")
                elif _xml(got) != _xml(exp):
                    rep.add("sheet", f"{ws.title}.{attr}: {_xml(got)!r} != {_xml(exp)!r}")
            for attr in ("print_title_rows", "print_title_cols", "print_area"):
                if getattr(s, attr) != getattr(ws, attr):
                    rep.add("sheet", f"{ws.title}.{attr}: {getattr(s, attr)!r} != {getattr(ws, attr)!r}")
    if "workbook" in sections:
        exp = [dict(name=ws.title, state=ws.sheet_state) for ws in wb._sheets]
        got = [dict(name=e["name"], state=e.get("state", "visible")) for e in book.sheet_entries]
        if exp != got:
            rep.add("workbook", f"sheets {got} != {exp}")
        exp_names = sorted(_xml(d) for d in wb.defined_names.values())
        for ws in wb.worksheets:
            exp_names += sorted(_xml(d) for d in ws.defined_names.values())
        got_names = [_xml(d) for d in book.defined_names]
        if sorted(exp_names) != sorted(got_names):
            rep.add("workbook", f"defined names {got_names[:3]} != {exp_names[:3]}")
        if [_xml(v) for v in wb.views] != [_xml(v) for v in book.views]:
            rep.add("workbook", "views")
        import zipfile

        z = zipfile.ZipFile(path)
        if "docProps/core.xml" in z.namelist():
            # times absent from the file are "now" in openpyxl: not data
            core = z.read("docProps/core.xml")
            absent = {n for n in ("created", "modified") if b":" + n.encode() + b">" not in core}

            def props(p):
                if p is None:
                    return None
                el = p.to_tree()
                for child in list(el):
                    if child.tag.rsplit("}", 1)[-1] in absent:
                        el.remove(child)
                return tostring(el)

            if props(wb.properties) != props(book.properties):
                rep.add("workbook", f"properties {props(book.properties)!r} != {props(wb.properties)!r}")
        if book.calculation is not None and _xml(wb.calculation) != _xml(book.calculation):
            rep.add("workbook", "calculation")
        if (wb.epoch.year == 1904) != book.date1904:
            rep.add("workbook", "date1904")
    wb.close()


def _norm(v):
    """formula objects have no __eq__: compare their content"""
    if isinstance(v, dict):
        return {k: _norm(x) for k, x in v.items()}
    if hasattr(v, "__iter__") and not isinstance(v, (str, list, tuple)) and hasattr(v, "t"):
        return (type(v).__name__, dict(v), getattr(v, "text", None))
    return v


def _rule_xml(rule):
    el = rule.to_tree()
    el.attrib.pop("dxfId", None)
    return tostring(el)


def verify(path, *, full=True, encoding="auto", sections=ALL, check_outputs=True, full_cells=False, sheets=None):
    """Compress `path`, decode, compare with openpyxl; returns (Report, doc).
    `full_cells`: cells as full mode reads them (out of order rows);
    `sheets`: only these worksheets"""
    rep = Report()
    doc = compress_workbook(path, include=sections, encoding=encoding, sheets=sheets)
    included = set(doc.included)
    book = decompress(doc.text())
    check_cells(path, book, rep, included, full_cells)
    if full:
        check_metadata(path, book, rep, included, only=sheets)
    if check_outputs:
        # every output form decodes to the same document
        if decompress(doc.to_json()).__dict__.keys() != book.__dict__.keys():
            rep.add("outputs", "json")
        b2 = decompress(doc.to_json())
        for s1, s2 in zip(book.sheets, b2.sheets):
            for attr in ("values", "formulas", "style_ids", "merged"):
                if _norm(getattr(s1, attr)) != _norm(getattr(s2, attr)):
                    rep.add("outputs", f"json {s1.name}.{attr}")
        with tempfile.TemporaryDirectory() as d:
            base = os.path.join(d, "out")
            streamed = compress_workbook(
                path, include=sections, encoding=encoding, output=["text", "json"], path=base, sheets=sheets
            )
            with open(streamed.files["text"], encoding="utf-8") as f:
                t = f.read()
            if t != doc.text():
                rep.add("outputs", "streamed text file differs from the in-memory text")
            with open(streamed.files["json"], encoding="utf-8") as f:
                j = json.load(f)
            if j != json.loads(doc.to_json()):
                rep.add("outputs", "streamed JSON file differs from the in-memory JSON")
    return rep, doc


# -- filters (skip_hidden / skip_hidden_sheets / print_area_only) -------------------
_BIG = 1 << 40


def filter_keep(ws, skip_hidden, area):
    """keep((row, col)) for a filtered document, from openpyxl's full-mode
    worksheet `ws` (independent of the encoder's own computation)"""
    from openpyxl.utils.cell import column_index_from_string

    rows, cols, listed, zero, rects = set(), set(), set(), False, None
    if skip_hidden:
        zero = bool(ws.sheet_format.zeroHeight)
        for i, rd in ws.row_dimensions.items():
            (rows if rd.hidden else listed).add(int(i))
        for key, cd in ws.column_dimensions.items():
            if cd.hidden:
                lo = cd.min or column_index_from_string(key)
                cols.update(range(lo, (cd.max or lo) + 1))
    if area and ws._print_area is not None and ws._print_area.ranges:
        rects = [(r.min_row or 1, r.min_col or 1, r.max_row or _BIG, r.max_col or _BIG) for r in ws._print_area.ranges]
        lo_c, hi_c = min(r[1] for r in rects), max(r[3] for r in rects)
        lo_r, hi_r = min(r[0] for r in rects), max(r[2] for r in rects)
        pr, pc = ws._print_rows, ws._print_cols
        if pr is not None and pr.min_row is not None:
            rects.append((pr.min_row, lo_c, pr.max_row or pr.min_row, hi_c))
        if pc is not None and pc.min_col is not None:
            c1 = column_index_from_string(pc.min_col)
            rects.append((lo_r, c1, hi_r, column_index_from_string(pc.max_col or pc.min_col)))

    def keep(rc):
        r, c = rc
        if c in cols or ((r not in listed) if zero else (r in rows)):
            return False
        return rects is None or any(a <= r <= b and x <= c <= y for a, x, b, y in rects)

    return keep


SHEET_ATTRS = ("values", "data_types", "formulas", "style_ids", "merged", "hyperlinks", "comments")


def check_sheets(path, rep):
    """Each worksheet compressed alone (compress_sheets) equals the same sheet
    of the whole-workbook document, and compress_workbook(sheets=[name])"""
    kw = dict(legend=False, include=["values", "formulas", "styles", "merged", "hyperlinks", "comments"])
    whole = compress_workbook(path, **kw)
    full = decompress(whole.text())
    seen = []
    for name, doc in compress_sheets(path, **kw):
        seen.append(name)
        if doc.text() != compress_workbook(path, sheets=[name], **kw).text():
            rep.add(f"sheets {name}", "compress_sheets != compress_workbook(sheets=[name])")
        (s,) = decompress(doc.text()).sheets
        f = full[name]
        if s.index != f.index:
            rep.add(f"sheets {name}", f"number {s.index} != {f.index}")
        for attr in SHEET_ATTRS:
            if _norm(getattr(s, attr)) != _norm(getattr(f, attr)):
                rep.add(f"sheets {name}", attr)
        if doc[name].text() != whole[name].text():
            rep.add(f"sheets {name}", "sheet text differs from the whole document's")
    if seen != [s.name for s in full.sheets]:
        rep.add("sheets", f"{seen} != {[s.name for s in full.sheets]}")
    return rep


FILTER_SECTIONS = ["values", "formulas", "styles", "hyperlinks", "comments"]
FILTER_ATTRS = ("values", "data_types", "formulas", "style_ids", "hyperlinks", "comments")


def check_filters(path, rep, encoding="auto", skip_hidden=True, area=True, hidden_sheets=True):
    """A filtered document equals the lossless one restricted to the cells
    filter_keep keeps, and leaves out exactly the hidden sheets"""
    kw = dict(include=FILTER_SECTIONS, encoding=encoding, legend=False)
    full = decompress(compress_workbook(path, **kw).text())
    doc = compress_workbook(path, skip_hidden=skip_hidden, print_area_only=area, skip_hidden_sheets=hidden_sheets, **kw)
    book = decompress(doc.text())
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        ref = openpyxl.load_workbook(path)
    try:
        want_names = [s.name for s in full.sheets if not hidden_sheets or ref[s.name].sheet_state == "visible"]
        if [s.name for s in book.sheets] != want_names:
            rep.add("filters", f"sheets {[s.name for s in book.sheets]} != {want_names}")
            return rep
        for s in book.sheets:
            keep = filter_keep(ref[s.name], skip_hidden, area)
            f = full[s.name]
            for attr in FILTER_ATTRS:
                want = {k: v for k, v in getattr(f, attr).items() if keep(k)}
                got = getattr(s, attr)
                if _norm(got) != _norm(want):
                    extra = sorted(set(got) - set(want))[:3]
                    missing = sorted(set(want) - set(got))[:3]
                    rep.add(f"filters {s.name}", f"{attr}: extra {extra} missing {missing}")
    finally:
        ref.close()
    return rep


def main(argv):
    """--no-full: no metadata (openpyxl's full mode needs the whole file in
    memory); --per-sheet: one sheet at a time (big files); --encoding E"""
    full = "--no-full" not in argv
    per_sheet = "--per-sheet" in argv
    enc = "auto"
    if "--encoding" in argv:
        enc = argv[argv.index("--encoding") + 1]
    files = [a for a in argv if not a.startswith("--") and a not in ("grid", "index", "hybrid", "auto")]
    bad = 0
    for p in files:
        if per_sheet:
            wb = openpyxl.load_workbook(p, read_only=True)
            names = [ws.title for ws in wb.worksheets]
            wb.close()
            for n in names:
                sec = [*("values", "formulas", "styles", "analysis")]
                rep, doc = verify(p, full=False, encoding=enc, sections=sec, check_outputs=False, sheets=[n])
                print(f"{p} [{n}]: {rep} ({[s.analysis['encoding'] for s in doc.sheets]})", flush=True)
                bad += not rep
                del doc
            continue
        rep, doc = verify(p, full=full, encoding=enc)
        print(f"{p}: {rep} ({doc.tokens} tokens, {[s.analysis['encoding'] for s in doc.sheets]})")
        bad += not rep
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))


def compare_llm(path):
    """Corpus check (tools/corpus.py "llm"): differences as strings. Runs the
    automatic encoding with every output form, and one forced encoding
    chosen per file (reproducibly) so the corpus covers all of them."""
    import zlib

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        try:
            openpyxl.load_workbook(path).close()
            openpyxl.load_workbook(path, read_only=True).close()
        except Exception:
            # openpyxl cannot read the file: nothing to compare with (ours
            # may fail too, it must only not hang or crash)
            try:
                compress_workbook(path)
            except Exception:
                pass
            return []
        out = []
        forced = ("grid", "index", "hybrid")[zlib.crc32(os.path.basename(path).encode()) % 3]
        for enc, outputs in (("auto", True), (forced, False)):
            try:
                rep, _ = verify(path, encoding=enc, check_outputs=outputs)
            except Exception as e:
                import traceback

                out.append(f"[{enc}] {type(e).__name__}: {e}\n{traceback.format_exc()[-1200:]}")
                continue
            if not rep:
                out.append(f"[{enc}] {rep}")
        try:
            rep = check_sheets(path, Report())
        except Exception as e:
            import traceback

            out.append(f"[sheets] {type(e).__name__}: {e}\n{traceback.format_exc()[-1200:]}")
        else:
            if not rep:
                out.append(f"[sheets] {rep}")
        try:
            rep = check_filters(path, Report())
        except Exception as e:
            import traceback

            out.append(f"[filters] {type(e).__name__}: {e}\n{traceback.format_exc()[-1200:]}")
        else:
            if not rep:
                out.append(f"[filters] {rep}")
        return out
