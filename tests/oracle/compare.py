"""
Differential comparison of openpyxl (oracle) and openrsxl.

Used both by the pytest suite and as a CLI::

    python tests/oracle/compare.py file.xlsx [--read-only] [--data-only] [--rich-text]

Everything observable is compared: every cell (value, Python type, data_type,
style array, number format, font, fill, border, alignment, protection,
hyperlink, comment), worksheet properties (dimensions, merged cells,
row/column dimensions, conditional formatting, data validation, views,
print settings, tables, ...), workbook properties, defined names and the
complete output of ``save`` (byte for byte, per package part).
"""

import datetime
import gc
import hashlib
import io
import re
import sys
import warnings
import weakref
import zipfile

MODULE_RE = re.compile(r"\bopenrsxl\b")


def norm_text(s):
    """Map openrsxl module paths to openpyxl ones for comparison."""
    return MODULE_RE.sub("openpyxl", s)


def type_name(v):
    t = type(v)
    return norm_text(f"{t.__module__}.{t.__qualname__}")


#: id(long string) -> (string, digest): a string shared by many cells is
#: hashed once (the string is kept so that its id cannot be reused)
_LONG = {}


def _long_text(v):
    hit = _LONG.get(id(v))
    if hit is None or hit[0] is not v:
        if len(_LONG) > 100_000:
            _LONG.clear()
        data = v.encode("utf-8", "surrogatepass") if isinstance(v, str) else v
        hit = _LONG[id(v)] = (v, f"<{len(v)} chars, sha1 {hashlib.sha1(data).hexdigest()}>")
    return hit[1]


def norm(v, depth=0):
    """Normalise an arbitrary value into comparable plain data."""
    if depth > 30:
        return "<deep>"
    if isinstance(v, (str, bytes)) and len(v) > 256:
        # (states stay small: big shared strings are compared by digest)
        return (type_name(v), _long_text(v))
    if v is None or isinstance(v, (bool, int, str, bytes)):
        return (type_name(v), v)
    if isinstance(v, float):
        if v != v:
            return ("float", "nan")
        return ("float", repr(v))
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time, datetime.timedelta)):
        return (type_name(v), repr(v))
    if isinstance(v, (list, tuple)):
        return (type_name(v), [norm(x, depth + 1) for x in v])
    if isinstance(v, dict):
        return (type_name(v), [(norm(k, depth + 1), norm(x, depth + 1)) for k, x in v.items()])
    # Serialisable objects: use __attrs__/__elements__/__nested__
    fields = []
    for group in ("__attrs__", "__nested__", "__elements__"):
        names = getattr(type(v), group, None)
        if names:
            fields.extend(names)
    if fields:
        out = []
        for f in dict.fromkeys(fields):
            try:
                out.append((f, norm(getattr(v, f), depth + 1)))
            except Exception as e:  # pragma: no cover
                out.append((f, ("error", type(e).__name__)))
        return (type_name(v), out)
    attrs = getattr(v, "__dict__", None)
    if attrs is None and hasattr(type(v), "__slots__"):
        attrs = {k: getattr(v, k) for k in type(v).__slots__ if hasattr(v, k)}
    if attrs:
        return (type_name(v), [(k, norm(x, depth + 1)) for k, x in attrs.items() if k not in ("parent", "_parent")])
    return (type_name(v), ADDR_RE.sub("", norm_text(repr(v))))


ADDR_RE = re.compile(r" at 0x[0-9A-Fa-f]+")

#: workbook -> {(cell type, style array): style state}; weak keys: the
#: entries of a workbook die with it (ids are reused, workbooks are not)
_STYLE_CACHE = weakref.WeakKeyDictionary()


def style_state(c):
    """Style-derived properties, cached per (workbook, style array)."""
    sty = getattr(c, "_style", None)
    if sty is None and hasattr(c, "style_array"):
        sty = c.style_array
    cache = _STYLE_CACHE.get(c.parent.parent)
    if cache is None:
        cache = _STYLE_CACHE[c.parent.parent] = {}
    key = (type(c).__name__, tuple(sty) if sty is not None else None)
    st = cache.get(key)
    if st is None:
        st = {
            "number_format": c.number_format,
            "font": norm(c.font),
            "fill": norm(c.fill),
            "border": norm(c.border),
            "alignment": norm(c.alignment),
            "protection": norm(c.protection),
        }
        cache[key] = st
    return st


def cell_state(c):
    st = {
        "coord": c.coordinate,
        "row": c.row,
        "column": c.column,
        "type": type_name(c),
        "value": norm(c.value),
        "data_type": c.data_type,
    }
    style = getattr(c, "_style", None)
    st["style"] = list(style) if style is not None else None
    st["style_id_attr"] = getattr(c, "_style_id", None)
    if hasattr(c, "has_style"):
        st.update(style_state(c))
    st["hyperlink"] = norm(getattr(c, "hyperlink", None))
    st["comment"] = norm_text(repr(getattr(c, "comment", None)))
    if getattr(c, "comment", None) is not None:
        st["comment_text"] = (c.comment.text, c.comment.author)
    return st


WS_ATTRS = [
    "title",
    "sheet_state",
    "dimensions",
    "min_row",
    "max_row",
    "min_column",
    "max_column",
    "freeze_panes",
    "print_area",
    "print_title_rows",
    "print_title_cols",
    "sheet_properties",
    "sheet_format",
    "views",
    "page_margins",
    "page_setup",
    "print_options",
    "HeaderFooter",
    "protection",
    "auto_filter",
    "data_validations",
    "row_breaks",
    "col_breaks",
    "scenarios",
    "legacy_drawing",
    "_current_row",
]


#: sheets with more cells are compared through digests (sheet_state)
BIG_SHEET = 200_000


def _chunk_digests(items, state, size=2000):
    out = []
    h = hashlib.sha1()
    n = 0
    for item in items:
        h.update(state(item).encode("utf-8", "surrogatepass"))
        n += 1
        if n == size:
            out.append(h.hexdigest())
            h = hashlib.sha1()
            n = 0
    if n:
        out.append(h.hexdigest())
    return out


def sheet_state(ws):
    out = {}
    for a in WS_ATTRS:
        try:
            out[a] = norm(getattr(ws, a))
        except Exception as e:
            out[a] = ("error", type(e).__name__, str(e))
    out["merged"] = [str(r) for r in getattr(ws, "merged_cells", [])]
    if len(ws._cells) <= BIG_SHEET:
        out["cells_order"] = list(ws._cells.keys())
        out["cells"] = [cell_state(c) for c in ws._cells.values()]
    else:
        # bounded memory: the order and state of the cells as digests, a few
        # thousand cells per entry (the first difference is still located)
        out["cells_order"] = _chunk_digests(ws._cells.keys(), repr)
        out["cells"] = _chunk_digests(ws._cells.values(), lambda c: repr(cell_state(c)))
    out["row_dimensions"] = [
        (k, norm_text(repr(v)), list(v._style) if v._style else None) for k, v in ws.row_dimensions.items()
    ]
    out["column_dimensions"] = [
        (k, norm_text(repr(v)), norm(v.width), list(v._style) if v._style else None)
        for k, v in ws.column_dimensions.items()
    ]
    cf = []
    for rng in ws.conditional_formatting:
        cf.append((str(rng.sqref), [norm(r) for r in rng.rules]))
    out["conditional_formatting"] = cf
    out["tables"] = [(k, norm(t)) for k, t in ws.tables.items()]
    out["hyperlinks"] = [norm(h) for h in getattr(ws, "_hyperlinks", [])]
    out["images"] = len(getattr(ws, "_images", []))
    out["charts"] = len(getattr(ws, "_charts", []))
    out["defined_names"] = sorted((k, norm(v)) for k, v in ws.defined_names.items())
    return out


def props_state(p):
    """Document properties; timestamps defaulting to "now" are masked."""
    st = norm(p)
    now = datetime.datetime.now(datetime.timezone.utc).replace(tzinfo=None)
    fields = []
    for k, v in st[1]:
        if k in ("created", "modified"):
            raw = getattr(p, k)
            if (
                isinstance(raw, datetime.datetime)
                and abs((raw.replace(tzinfo=None) - now).total_seconds()) < 7200
                or isinstance(raw, datetime.datetime)
                and abs((raw.replace(tzinfo=None) - datetime.datetime.now()).total_seconds()) < 7200
            ):
                v = ("now", type(raw).__name__)
        fields.append((k, v))
    return (st[0], fields)


def workbook_state(wb):
    out = {
        "sheetnames": wb.sheetnames,
        "active": wb.active.title if wb.active is not None else None,
        "epoch": repr(wb.epoch),
        "properties": props_state(wb.properties),
        "calculation": norm(wb.calculation),
        "defined_names": sorted((k, norm(v)) for k, v in wb.defined_names.items()),
        "named_styles": list(wb.named_styles),
        "cell_styles": [list(s) for s in wb._cell_styles],
        "fonts": [norm(x) for x in wb._fonts],
        "fills": [norm(x) for x in wb._fills],
        "borders": [norm(x) for x in wb._borders],
        "number_formats": list(wb._number_formats),
        "date_formats": sorted(wb._date_formats),
    }
    return out


def first_diff(a, b, path="$"):
    if type(a) == type(b) and a == b:
        # equal nested data is recognised in C, repr tells 1, 1.0 and True
        # apart; big containers are checked child by child (bounded memory)
        if not isinstance(a, (list, tuple, dict)) or len(a) <= 64:
            if repr(a) == repr(b):
                return None
    if type(a) != type(b):
        return f"{path}: type {type(a).__name__} != {type(b).__name__}: {a!r:.200} vs {b!r:.200}"
    if isinstance(a, dict):
        for k in a.keys() | b.keys():
            if k not in a or k not in b:
                return f"{path}.{k}: missing"
            d = first_diff(a[k], b[k], f"{path}.{k}")
            if d:
                return d
        return None
    if isinstance(a, (list, tuple)):
        if len(a) != len(b):
            for i, (x, y) in enumerate(zip(a, b)):
                d = first_diff(x, y, f"{path}[{i}]")
                if d:
                    return d
            return f"{path}: len {len(a)} != {len(b)}"
        for i, (x, y) in enumerate(zip(a, b)):
            d = first_diff(x, y, f"{path}[{i}]")
            if d:
                return d
        return None
    if a != b:
        return f"{path}: {a!r:.300} != {b!r:.300}"
    return None


def guarded(fn, *args):
    """Run fn capturing exceptions and warnings as comparable data."""
    with warnings.catch_warnings(record=True) as w:
        warnings.simplefilter("always")
        try:
            res = ("ok", fn(*args))
        except Exception as e:
            res = ("error", type(e).__name__, norm_text(str(e)))
    return res, _messages(w)


def _messages(w):
    # ResourceWarnings are emitted whenever the garbage collector finalizes
    # an object (eg. the temporary file openpyxl leaves open after a failed
    # save): their timing says nothing about the library being compared
    return [(x.category.__name__, norm_text(str(x.message))) for x in w if not issubclass(x.category, ResourceWarning)]


def load(lib, src, **kw):
    gc.collect()
    with warnings.catch_warnings(record=True) as w:
        warnings.simplefilter("always")
        if isinstance(src, bytes):
            src = io.BytesIO(src)
        try:
            wb = lib.load_workbook(src, **kw)
            err = None
        except Exception as e:
            wb, err = None, (type(e).__name__, norm_text(str(e)))
    return wb, err, _messages(w)


FIXED = datetime.datetime(2020, 1, 1, 12, 0, 0)


#: saved parts bigger than this are compared by size and SHA-1, read in
#: chunks (openpyxl writes every string inline: a 1 MB shared string used by
#: 10 000 cells becomes a 10 GB worksheet part)
BIG_PART = 64 * 2**20


def saved_parts(wb):
    wb.properties.created = FIXED
    wb.properties.modified = FIXED
    b = io.BytesIO()
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb.save(b)
    z = zipfile.ZipFile(b)
    parts = {}
    for info in z.infolist():
        name = info.filename
        if info.file_size > BIG_PART:
            h = hashlib.sha1()
            with z.open(info) as f:
                for chunk in iter(lambda: f.read(1 << 20), b""):
                    h.update(chunk)
            parts[name] = ("<big part>", info.file_size, h.hexdigest())
            continue
        data = z.read(name)
        if name == "docProps/core.xml":
            data = re.sub(rb"<dcterms:modified[^<]*</dcterms:modified>", b"", data)
        parts[name] = data
    return z.namelist(), parts


def _values_state(row):
    return [norm(v) for v in row]


def _cells_state(row):
    return [cell_state(c) if hasattr(c, "coordinate") else repr(c) for c in row]


def _row_digests(rows, state):
    """
    One digest per streamed row, run-length encoded ([digest, count] pairs):
    memory stays small for sheets whose declared dimension makes read-only
    mode pad millions of empty cells.
    """
    out = []
    for row in rows:
        s = repr(state(row))
        dig = s if len(s) <= 64 else hashlib.sha1(s.encode("utf-8", "surrogatepass")).hexdigest()
        if out and out[-1][0] == dig:
            out[-1][1] += 1
        else:
            out.append([dig, 1])
    return out


def _row_details(sa, sb, ra, rb, label):
    """The full state of the first row whose digest differs."""
    if ra[0] != "ok" or rb[0] != "ok":
        return ""
    row = 0
    for (da, na), (db, nb) in zip(ra[1], rb[1]):
        if da != db:
            break
        if na != nb:
            row += min(na, nb)
            break
        row += na
    state = _values_state if label == "values" else _cells_state
    out = []
    for ws in (sa, sb):
        try:
            for i, r in enumerate(ws.iter_rows(values_only=label == "values")):
                if i == row:
                    out.append(repr(state(r))[:1500])
                    break
        except Exception as e:  # pragma: no cover
            out.append(repr(e))
    return f"\n    first differing row {row + 1}:\n      " + "\n      ".join(out)


def compare_workbooks(src, check_save=True, **kw):
    """Return a list of differences (empty if identical)."""
    import openpyxl

    import openrsxl

    diffs = []
    wa, ea, wa_msgs = load(openpyxl, src, **kw)
    wb, eb, wb_msgs = load(openrsxl, src, **kw)
    if ea or eb:
        if ea != eb:
            diffs.append(f"load error: {ea} != {eb}")
        return diffs
    if wa_msgs != wb_msgs:
        diffs.append(f"warnings: {wa_msgs} != {wb_msgs}")
    d = first_diff(workbook_state(wa), workbook_state(wb))
    if d:
        diffs.append("workbook " + d)
    for sa, sb in zip(wa.worksheets, wb.worksheets):
        if kw.get("read_only"):
            for label, fn in (
                ("values", lambda ws: _row_digests(ws.iter_rows(values_only=True), _values_state)),
                ("cells", lambda ws: _row_digests(ws.iter_rows(), _cells_state)),
                (
                    "bounded",
                    lambda ws: [
                        [norm(v) for v in r]
                        for r in ws.iter_rows(min_row=2, max_row=4, min_col=2, max_col=3, values_only=True)
                    ],
                ),
                ("dims", lambda ws: (ws.min_row, ws.max_row, ws.min_column, ws.max_column)),
            ):
                ra, rb = guarded(fn, sa), guarded(fn, sb)
                d = first_diff(ra, rb, f"[{sa.title}].{label}")
                if d and label in ("values", "cells"):
                    d += _row_details(sa, sb, ra, rb, label)
                if d:
                    diffs.append(d)
            continue
        d = first_diff(sheet_state(sa), sheet_state(sb), f"[{sa.title}]")
        if d:
            diffs.append(d)
    if check_save and not kw.get("read_only"):
        # 1. save right after loading (openrsxl: compact cells, nothing
        #    materialised) 2. save after every cell has been accessed
        fa, _, _ = load(openpyxl, src, **kw)
        fb, _, _ = load(openrsxl, src, **kw)
        diffs.extend("direct " + d for d in _compare_saves(fa, fb))
        diffs.extend(_compare_saves(wa, wb))
    return diffs


def _compare_saves(wa, wb):
    diffs = []
    ra, rb = guarded(saved_parts, wa)[0], guarded(saved_parts, wb)[0]
    if ra[0] != "ok" or rb[0] != "ok":
        if ra != rb:
            diffs.append(f"save: {ra[:3]} != {rb[:3]}")
        return diffs
    if True:
        na, pa = ra[1]
        nb, pb = rb[1]
        if na != nb:
            diffs.append(f"saved part names differ: {na} != {nb}")
        for name in na:
            if pa.get(name) != pb.get(name):
                x, y = pa.get(name) or b"", pb.get(name) or b""
                if not isinstance(x, bytes) or not isinstance(y, bytes):
                    diffs.append(f"saved part {name} differs: {x!r:.120} != {y!r:.120}")
                    continue
                i = next((i for i in range(min(len(x), len(y))) if x[i] != y[i]), min(len(x), len(y)))
                diffs.append(
                    f"saved part {name} differs at byte {i}: {x[max(0,i-80):i+80]!r} != {y[max(0,i-80):i+80]!r}"
                )
    return diffs


if __name__ == "__main__":
    args = sys.argv[1:]
    kw = {}
    for flag, key in (("--read-only", "read_only"), ("--data-only", "data_only"), ("--rich-text", "rich_text")):
        if flag in args:
            args.remove(flag)
            kw[key] = True
    for path in args:
        diffs = compare_workbooks(path, **kw)
        print(path, kw, "IDENTICAL" if not diffs else "DIFFERENT")
        for d in diffs[:20]:
            print("  ", d)
