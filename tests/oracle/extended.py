"""
Oracle for openrsxl.extended streaming features.

The expected result of every feature is what openpyxl's *full* mode returns
for the same file:

* cells, values, types, styles, hyperlinks, comments, MergedCell objects:
  openpyxl ``load_workbook(data_only=<same>)``;
* ``cell.formula``: openpyxl ``data_only=False`` value of formula cells;
* ``cell.cached_value``: openpyxl ``data_only=True`` value;
* worksheet attributes (merged cells, dimensions, properties, ...): the same
  attributes of openpyxl's full mode worksheet.

CLI::

    python tests/oracle/extended.py file.xlsx [--data-only]
"""

import gc
import io
import os
import sys
import warnings
import weakref

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

from oracle.compare import _messages, norm, norm_text, sheet_state, type_name  # noqa: E402


def _unproxy(v):
    # full mode wraps styles in StyleProxy, read-only cells return them as is
    return getattr(v, "_StyleProxy__target", v)


#: workbook -> {(cell kind, style array): style state} (weak keys: ids of
#: collected workbooks are reused)
_STYLES = weakref.WeakKeyDictionary()


def style_state(c):
    sty = getattr(c, "_style", None)
    if sty is None:
        sty = getattr(c, "style_array", None)
    cache = _STYLES.get(c.parent.parent)
    if cache is None:
        cache = _STYLES[c.parent.parent] = {}
    key = (kind(c), tuple(sty) if sty is not None else None)
    st = cache.get(key)
    if st is None:
        st = cache[key] = _style_state(c)
    return st


def _style_state(c):
    return {
        "number_format": c.number_format,
        "font": norm(_unproxy(c.font)),
        "fill": norm(_unproxy(c.fill)),
        "border": norm(_unproxy(c.border)),
        "alignment": norm(_unproxy(c.alignment)),
        "protection": norm(_unproxy(c.protection)),
    }


ALL_FLAGS = dict(
    formula_and_value=True,
    read_comments=True,
    read_hyperlinks=True,
    read_merged_cells=True,
    read_dimensions=True,
    read_sheet_properties=True,
    read_data_validations=True,
    read_conditional_formatting=True,
    read_tables=True,
    create_empty_cells=True,
)

PROPERTY_ATTRS = [
    "freeze_panes",
    "print_area",
    "print_title_rows",
    "print_title_cols",
    "sheet_properties",
    "sheet_format",
    "views",
    "page_margins",
    "page_setup",
    "print_options",
    "HeaderFooter",
    "protection",
    "auto_filter",
    "row_breaks",
    "col_breaks",
    "scenarios",
    "legacy_drawing",
]


def _src(src):
    return io.BytesIO(src) if isinstance(src, bytes) else src


def _load(lib, src, **kw):
    gc.collect()
    with warnings.catch_warnings(record=True) as w:
        warnings.simplefilter("always")
        wb = lib.load_workbook(_src(src), **kw)
    return wb, _messages(w)


def kind(c):
    n = type(c).__name__
    if n in ("Cell", "ReadOnlyCell"):
        return "Cell"
    return n


def expected_cell(c, fcell, vcell, flags):
    st = {
        "kind": kind(c),
        "value": norm(c.value),
        "data_type": c.data_type,
    }
    st.update(style_state(c))
    if flags.get("formula_and_value"):
        if kind(c) == "MergedCell":
            st["formula"] = norm(None)
            st["cached_value"] = norm(None)
        else:
            st["formula"] = norm(fcell.value if fcell.data_type == "f" else None)
            st["cached_value"] = norm(vcell.value)
    if flags.get("read_hyperlinks") and kind(c) != "MergedCell":
        st["hyperlink"] = norm(c.hyperlink)
    if flags.get("read_comments") and kind(c) != "MergedCell":
        st["comment"] = _comment(c.comment)
    return st


def _comment(cm):
    if cm is None:
        return None
    return (cm.text, cm.author, cm.width, cm.height, cm.parent.coordinate if cm.parent else None)


def actual_cell(c, flags):
    st = {
        "kind": kind(c),
        "value": norm(c.value),
        "data_type": c.data_type,
    }
    st.update(style_state(c))
    if flags.get("formula_and_value"):
        st["formula"] = norm(c.formula)
        st["cached_value"] = norm(c.cached_value)
    if flags.get("read_hyperlinks") and kind(c) != "MergedCell":
        st["hyperlink"] = norm(c.hyperlink)
    if flags.get("read_comments") and kind(c) != "MergedCell":
        st["comment"] = _comment(c.comment)
    return st


def compare_extended(src, data_only=False, flags=None, rsxl=None, oracle=None, limit=20):
    """Return a list of differences (empty = identical)."""
    if rsxl is None:
        import openrsxl.extended as rsxl
    if oracle is None:
        import openpyxl as oracle
    flags = dict(ALL_FLAGS if flags is None else flags)
    # the expectations follow the resolved features (read_all: every feature
    # not given explicitly); the loader gets the flags as given
    given = flags
    if "read_all" in flags:
        flags = {k: bool(flags["read_all"]) for k in ALL_FLAGS}
        flags.update((k, v) for k, v in given.items() if k != "read_all" and v is not None)
    diffs = []

    def diff(msg):
        diffs.append(msg)
        return len(diffs) >= limit

    # features that are off do not change the cells: no merge / hyperlink
    # binding in the oracle either
    from openpyxl.cell.cell import Cell
    from openpyxl.worksheet._reader import WorksheetReader

    patched = {}
    if not flags.get("read_merged_cells"):
        patched["bind_merged_cells"] = WorksheetReader.bind_merged_cells
    if not flags.get("read_hyperlinks"):
        patched["bind_hyperlinks"] = WorksheetReader.bind_hyperlinks
    for k in patched:
        setattr(WorksheetReader, k, lambda self: None)
    try:
        try:
            full_f, warn_f = _load(oracle, src, data_only=False)
            full_v, warn_v = _load(oracle, src, data_only=True)
            exp_err = None
        except Exception as err:
            exp_err = (type(err).__name__, norm_text(str(err)))
    finally:
        for k, v in patched.items():
            setattr(WorksheetReader, k, v)
    try:
        ext, warn_e = _load(rsxl, src, read_only=True, data_only=data_only, **given)
        got_err = None
    except Exception as err:
        got_err = (type(err).__name__, norm_text(str(err)))
    if exp_err or got_err:
        if exp_err != got_err and not (
            got_err is None and _deferred(oracle, ext, src, data_only, diff, flags, exp_err)
        ):
            diff(f"load error {got_err} != {exp_err}")
        return diffs
    primary = full_v if data_only else full_f
    exp_warn = warn_v if data_only else warn_f
    # openpyxl's full mode warns while loading; read-only mode also warns
    # while iterating (cell level warnings are compared below)
    if flags.get("read_comments"):
        cw = [w for w in exp_warn if "merged range but has a comment" in w[1]]
        gw = [w for w in warn_e if "merged range but has a comment" in w[1]]
        if cw != gw:
            diff(f"comment warnings {gw} != {cw}")

    for name in primary.sheetnames:
        pws, ews = primary[name], ext[name]
        if type_name(pws).endswith("Chartsheet"):
            continue
        fws, vws = full_f[name], full_v[name]
        pcells = dict(pws._cells)
        fcells = dict(fws._cells)
        vcells = dict(vws._cells)
        if not flags.get("read_comments"):
            # full mode always binds comments, creating their cells (no
            # style yet, no value); streaming without read_comments does not
            for k, c in list(pcells.items()):
                if getattr(c, "_comment", None) is not None and c._style is None and c._value is None:
                    del pcells[k]
        seen = set()
        for rows_vals in (False, True):
            for r_i, row in enumerate(ews.iter_rows(values_only=rows_vals), start=1):
                for c_i, c in enumerate(row, start=1):
                    exp = pcells.get((r_i, c_i))
                    if rows_vals:
                        ev = None if exp is None else exp.value
                        if norm(ev) != norm(c):
                            if diff(f"{name}!{(r_i, c_i)} values_only: {c!r} != {ev!r}"):
                                return diffs
                        continue
                    seen.add((r_i, c_i))
                    if exp is None:
                        if not flags.get("create_empty_cells"):
                            if type(c).__name__ != "EmptyCell":
                                if diff(f"{name}!{(r_i, c_i)}: unexpected {type(c).__name__} {c!r}"):
                                    return diffs
                            continue
                        # the cell full mode creates when the position is
                        # accessed (not kept: the oracle workbook is unchanged)
                        exp = fexp = vexp = Cell(pws, row=r_i, column=c_i)
                        pos = (getattr(c, "row", None), getattr(c, "column", None), getattr(c, "coordinate", None))
                        if pos != (r_i, c_i, exp.coordinate):
                            if diff(f"{name}!{exp.coordinate}: empty cell at {pos}"):
                                return diffs
                    else:
                        fexp, vexp = fcells.get((r_i, c_i)), vcells.get((r_i, c_i))
                    e = expected_cell(exp, fexp, vexp, flags)
                    try:
                        a = actual_cell(c, flags)
                    except Exception as err:
                        a = ("error", repr(err))
                    if e != a:
                        if diff(f"{name}!{exp.coordinate}:\n  exp {e}\n  got {a}"):
                            return diffs
        # read-only mode only produces file cells inside the declared
        # dimension; cells created by the features are always produced
        plain = oracle.load_workbook(_src(src), read_only=True)[name]
        pr, pc = plain.max_row, plain.max_column

        def required(k, c):
            if k[0] < 1 or k[1] < 1:
                # (row 0 of r="A0" ...): iter_rows starts at row / column 1
                return False
            if (pr is None or k[0] <= pr) and (pc is None or k[1] <= pc):
                return True
            return (
                kind(c) == "MergedCell"
                or any(k == (r.min_row, r.min_col) for r in getattr(pws.merged_cells, "ranges", ()))
                or (flags.get("read_hyperlinks") and c.hyperlink is not None)
                or (flags.get("read_comments") and c.comment is not None)
            )

        missing = [k for k, c in pcells.items() if k not in seen and required(k, c)]
        if missing:
            diff(f"{name}: cells not produced by iter_rows: {missing[:10]} ({len(missing)})")

        if flags.get("exact_dimensions"):
            keys = list(pcells)
            if keys:
                exp_b = (
                    min(k[0] for k in keys),
                    min(k[1] for k in keys),
                    max(k[0] for k in keys),
                    max(k[1] for k in keys),
                )
            else:
                exp_b = (1, 1, 1, 1)
            got_b = (ews.min_row, ews.min_column, ews.max_row, ews.max_column)
            if got_b != exp_b:
                diff(f"{name}: bounds {got_b} != {exp_b}")
            if not keys and list(ews.iter_rows()):
                diff(f"{name}: iter_rows() of a sheet without cells")
        if flags.get("random_access"):
            if _check_random_access(name, ews, pws, pcells, fcells, vcells, flags, diff):
                return diffs

        # worksheet attributes
        pstate = sheet_state(pws)
        if flags.get("read_merged_cells"):
            got = [str(r) for r in ews.merged_cells]
            if got != pstate["merged"]:
                diff(f"{name}: merged {got} != {pstate['merged']}")
            # start_cell: the top-left cell of full mode
            eranges = {r.coord: r for r in ews.merged_cells}
            franges = {x.coord: x for x in fws.merged_cells}
            vranges = {x.coord: x for x in vws.merged_cells}
            for r in pws.merged_cells:
                er = eranges.get(r.coord)
                if er is None:
                    continue
                if type(er).__name__ != type(r).__name__ or not isinstance(er, rsxl.worksheet.merge.MergedCellRange):
                    diff(f"{name}: range type {type(er)}")
                # formula / cached value: the same range's start_cell of the
                # data_only=False / True workbooks (may be orphaned cells)
                fr = franges[r.coord]
                vr = vranges[r.coord]
                e = expected_cell(r.start_cell, fr.start_cell, vr.start_cell, flags)
                try:
                    a = actual_cell(er.start_cell, flags)
                except Exception as err:
                    a = ("error", repr(err))
                if e != a:
                    diff(f"{name}: {r.coord}.start_cell\n  exp {e}\n  got {a}")
                if str(__import__("copy").copy(er)) != str(r):
                    diff(f"{name}: copy of {r.coord}")
        if flags.get("read_dimensions"):
            for k in ("row_dimensions", "column_dimensions"):
                got = _dims(ews, k)
                if got != pstate[k]:
                    diff(f"{name}: {k}\n  exp {pstate[k][:5]}\n  got {got[:5]}")
        if flags.get("read_sheet_properties"):
            for a in PROPERTY_ATTRS:
                try:
                    got = norm(getattr(ews, a))
                except Exception as err:
                    got = ("error", type(err).__name__, str(err))
                if got != norm(getattr(pws, a)):
                    diff(f"{name}: {a}\n  exp {norm(getattr(pws, a))}\n  got {got}")
        if flags.get("read_data_validations"):
            if norm(ews.data_validations) != norm(pws.data_validations):
                diff(f"{name}: data_validations")
        if flags.get("read_conditional_formatting"):
            got = [(str(r.sqref), [norm(x) for x in r.rules]) for r in ews.conditional_formatting]
            if got != pstate["conditional_formatting"]:
                diff(f"{name}: conditional_formatting\n  exp {pstate['conditional_formatting']}\n  got {got}")
        if flags.get("read_tables"):
            got = [(k, norm(t)) for k, t in ews.tables.items()]
            if got != pstate["tables"]:
                diff(f"{name}: tables")
    return diffs


def _rows_sig(ws, flags):
    out = []
    for values_only in (False, True):
        for row in ws.iter_rows(values_only=values_only):
            if values_only:
                out.append(tuple(norm(v) for v in row))
            else:
                out.append(tuple(_sig(c, flags) for c in row))
    return out


def _sig(c, flags):
    if type(c).__name__ == "EmptyCell":
        return "EmptyCell"
    return (getattr(c, "row", None), getattr(c, "column", None), str(actual_cell(c, flags)))


def _check_random_access(name, ews, pws, pcells, fcells, vcells, flags, diff):
    """
    random_access: ws.cell() / ws[...] / ws._cells against full mode's cells
    (positions without a cell: the cell full mode creates, or EmptyCell
    without create_empty_cells); iter_rows unchanged by the store.
    Returns True when the limit of differences is reached.
    """
    from openpyxl.cell.cell import Cell

    before = _rows_sig(ews, flags)

    def expected(k):
        exp = pcells.get(k)
        if exp is None:
            if not flags.get("create_empty_cells"):
                return "EmptyCell"
            exp = Cell(pws, row=k[0], column=k[1])
            return str(expected_cell(exp, exp, exp, flags))
        return str(expected_cell(exp, fcells.get(k), vcells.get(k), flags))

    def got(c, k):
        if type(c).__name__ == "EmptyCell":
            return "EmptyCell"
        if (c.row, c.column) != k:
            return f"at {(c.row, c.column)}"
        return str(actual_cell(c, flags))

    # ws._cells: the cells full mode keeps (at valid positions)
    exp_keys = sorted(k for k in pcells if k[0] >= 1 and k[1] >= 1)
    got_keys = sorted(ews._cells)
    if got_keys != exp_keys:
        if diff(f"{name}: _cells keys {got_keys[:10]} != {exp_keys[:10]} ({len(got_keys)} / {len(exp_keys)})"):
            return True
    if len(ews._cells) != len(got_keys) or bool(ews._cells) != bool(got_keys):
        if diff(f"{name}: _cells len / bool"):
            return True
    for k in exp_keys:
        if k in ews._cells and got(ews._cells.get(k), k) != expected(k):
            if diff(f"{name}: _cells[{k}]\n  exp {expected(k)}\n  got {got(ews._cells[k], k)}"):
                return True
    # every position of the sheet and beyond
    max_r = min(max([k[0] for k in exp_keys] + [ews.max_row or 1]) + 2, 80)
    max_c = min(max([k[1] for k in exp_keys] + [ews.max_column or 1]) + 2, 40)
    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            a, e = got(ews.cell(r, c), (r, c)), expected((r, c))
            if a != e:
                if diff(f"{name}: cell({r}, {c})\n  exp {e}\n  got {a}"):
                    return True
    if ews.cell(1, 1) is not ews.cell(1, 1) and (1, 1) in pcells:
        diff(f"{name}: cell(1, 1) is not the same object")
    # ranges, rows, columns: every position (full mode)
    from openpyxl.utils import get_column_letter

    last = get_column_letter(max_c)
    rect = ews[f"A1:{last}{max_r}"]
    if [[got(c, (r, k)) for k, c in enumerate(row, 1)] for r, row in enumerate(rect, 1)] != [
        [expected((r, k)) for k in range(1, max_c + 1)] for r in range(1, max_r + 1)
    ]:
        diff(f"{name}: ws['A1:{last}{max_r}']")
    if ews["A1"] is not ews.cell(1, 1) and (1, 1) in pcells:
        diff(f"{name}: ws['A1'] is not ws.cell(1, 1)")
    # (unsized sheet: the bounds of the cells, as full mode)
    nrows = ews.max_row or max([k[0] for k in exp_keys], default=1)
    ncols = ews.max_column or max([k[1] for k in exp_keys], default=1)
    col = ews["B"]
    if [got(c, (r, 2)) for r, c in enumerate(col, 1)] != [expected((r, 2)) for r in range(1, nrows + 1)]:
        diff(f"{name}: ws['B']")
    row2 = ews[2]
    if [got(c, (2, k)) for k, c in enumerate(row2, 1)] != [expected((2, k)) for k in range(1, ncols + 1)]:
        diff(f"{name}: ws[2]")
    after = _rows_sig(ews, flags)
    if after != before:
        n = next((i for i, (x, y) in enumerate(zip(after, before)) if x != y), min(len(after), len(before)))
        diff(f"{name}: iter_rows changed by the store at row {n}: {after[n:n + 1]} != {before[n:n + 1]}")
    return False


def _deferred(oracle, ext, src, data_only, diff, flags=ALL_FLAGS, full_err=None):
    """
    Full mode failed while loading, the streaming features did not: accepted
    when openpyxl's read-only mode loads the file too (the error comes from
    cell data, which streaming only reads while iterating). The cells must
    then be those of read-only mode, including the errors raised by them
    (with formula_and_value, formulas are read as with data_only=False too:
    the errors of that reading are expected as well). Iterating may also
    raise full mode's own error (`full_err`) when a feature needs the faulty
    data (eg. the style of a merged range's corner cell).
    """
    try:
        plain = oracle.load_workbook(_src(src), read_only=True, data_only=data_only)
        other = None
        if data_only and flags.get("formula_and_value"):
            other = oracle.load_workbook(_src(src), read_only=True, data_only=False)
    except Exception:
        return False
    for name in plain.sheetnames:
        pws, ews = plain[name], ext[name]
        if type_name(pws).endswith("Chartsheet"):
            continue
        exp, got = _cells_or_error(pws), _cells_or_error(ews)
        if other is not None and not isinstance(exp, tuple):
            alt = _cells_or_error(other[name])
            if isinstance(alt, tuple):
                exp = alt
        if isinstance(got, tuple) and full_err is not None and got[1:] == tuple(full_err):
            continue
        if isinstance(exp, tuple) or isinstance(got, tuple):
            if exp != got:
                diff(f"{name}: iteration {got} != {exp}")
            continue
        covered = set()
        for r in getattr(ews, "merged_cells", ()):
            covered.update(c for c in r.cells if c != (r.min_row, r.min_col))
        for k, e in exp.items():
            a = got.get(k)
            if k in covered or a is None:
                continue
            if a != e and not (e[0] == norm(None) and a[0] != norm(None)):
                # (a value where read-only mode has none: link targets)
                diff(f"{name}!{k}: {a} != {e}")
    return True


def _cells_or_error(ws):
    out = {}
    try:
        for row in ws.iter_rows():
            for c in row:
                if not hasattr(c, "coordinate"):
                    continue
                try:
                    font = norm(_unproxy(c.font))
                except Exception as err:
                    font = ("error", type(err).__name__, str(err))
                out[(c.row, c.column)] = (norm(c.value), c.data_type, font)
    except Exception as err:
        return ("error", type(err).__name__, norm_text(str(err)))
    return out


def _dims(ws, k):
    if k == "row_dimensions":
        return [(k_, norm_text(repr(v)), list(v._style) if v._style else None) for k_, v in ws.row_dimensions.items()]
    return [
        (k_, norm_text(repr(v)), norm(v.width), list(v._style) if v._style else None)
        for k_, v in ws.column_dimensions.items()
    ]


if __name__ == "__main__":
    fn = sys.argv[1]
    d = compare_extended(fn, data_only="--data-only" in sys.argv)
    print("\n".join(d) if d else "IDENTICAL")
