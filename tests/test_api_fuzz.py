"""
Weird usage patterns of the openpyxl API, applied to workbooks *loaded* by
openpyxl and by openrsxl (whose worksheets keep compact cells and save
without materialising them): every operation result, the final state of
every worksheet and workbook and the saved package must be identical.

Complements test_cellstore.py (dict semantics of ``ws._cells``) with a richer
source workbook (shared / array formulae, comments, hyperlinks, merged cells,
data validation, conditional formatting, tables, defined names, several
sheets), every load mode and many more operations.
"""

import datetime
import decimal
import io
import os
import zipfile

import openpyxl
import pytest
from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st
from oracle.compare import ADDR_RE, first_diff, norm, norm_text, saved_parts, sheet_state, workbook_state

import openrsxl

EXAMPLES = int(os.environ.get("OPENRSXL_FUZZ_EXAMPLES", "300"))
SETTINGS = settings(max_examples=EXAMPLES, deadline=None, suppress_health_check=list(HealthCheck))


def make_source(epoch1904=False):
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.formula import ArrayFormula
    from openpyxl.worksheet.table import Table

    wb = openpyxl.Workbook()
    if epoch1904:
        from openpyxl.utils.datetime import MAC_EPOCH

        wb.epoch = MAC_EPOCH
    ws = wb.active
    ws.title = "Data"
    ws.append(["name", "qty", "price", "when", "flag", "total"])
    for r in range(2, 12):
        ws.append([f"item {r}", r, r * 1.25, datetime.datetime(2021, 1, r, 8, r), r % 2 == 0, f"=B{r}*C{r}"])
    ws["G2"] = ArrayFormula("G2:G4", "=B2:B4*2")
    ws["H1"] = "=SUM(F2:F11)"
    ws["H2"] = datetime.time(13, 30)
    ws["H3"] = datetime.timedelta(hours=30)
    ws["H4"] = datetime.date(1999, 12, 31)
    ws["I1"] = "#N/A"
    ws["I2"] = 1e300
    ws["I3"] = -0.0
    ws["B2"].font = Font(bold=True, color="FF0000")
    ws["C3"].number_format = "0.00%"
    ws["D2"].alignment = Alignment(horizontal="center", wrap_text=True)
    ws["E2"].fill = PatternFill("solid", fgColor="00FF00")
    ws["A5"].border = Border(left=Side(style="thin"), bottom=Side(style="double"))
    ws["A2"].comment = Comment("a comment", "me")
    ws["A3"].hyperlink = "https://example.com/x?a=1&b=2"
    ws["A4"].hyperlink = "#Other!B2"
    ws.merge_cells("J1:K3")
    ws.merge_cells("A14:C15")
    ws["A14"] = "merged"
    dv = DataValidation(type="list", formula1='"a,b,c"', allow_blank=True)
    dv.add("E2:E11")
    ws.add_data_validation(dv)
    ws.conditional_formatting.add("B2:B11", CellIsRule(operator="greaterThan", formula=["5"], font=Font(italic=True)))
    ws.add_table(Table(displayName="Items", ref="A1:F11"))
    ws.row_dimensions[3].height = 30
    ws.column_dimensions["A"].width = 25
    ws.freeze_panes = "B2"
    ws.print_title_rows = "1:1"
    wb.defined_names["total"] = DefinedName("total", attr_text="Data!$H$1")
    other = wb.create_sheet("Other")
    other["B2"] = "=Data!B2+1"
    other["A1"] = "x" * 300
    for r in range(1, 30, 3):
        other.cell(row=r, column=r % 7 + 1, value=r)
    b = io.BytesIO()
    wb.save(b)
    data = b.getvalue()
    # a shared formula (openpyxl never writes them)
    z = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for name in z.namelist():
            part = z.read(name)
            if name == "xl/worksheets/sheet2.xml":
                part = part.replace(
                    b"</sheetData>",
                    b'<row r="40"><c r="A40"><f t="shared" ref="A40:C41" si="0">B40+$A$1</f><v>1</v></c>'
                    b'<c r="B40"><f t="shared" si="0"/><v>2</v></c></row>'
                    b'<row r="41"><c r="C41"><f t="shared" si="0"/><v>3</v></c></row></sheetData>',
                )
            zo.writestr(name, part)
    return out.getvalue()


SOURCES = {False: make_source(), True: make_source(epoch1904=True)}

coords = st.tuples(st.integers(1, 45), st.integers(1, 14))
far = st.tuples(st.sampled_from([1, 2, 1048575, 1048576]), st.sampled_from([1, 2, 16383, 16384]))
values = st.one_of(
    st.none(),
    st.integers(-10, 10),
    st.integers(-(2**70), 2**70),
    st.floats(allow_nan=True, allow_infinity=True),
    st.text(max_size=6),
    st.sampled_from(["=A1+1", "=SUM(A:A)", "#DIV/0!", "=", " ", "\x01bad", "x" * 40000, "é€"]),
    st.booleans(),
    st.just(datetime.datetime(2020, 2, 29, 23, 59, 59, 999999)),
    st.just(datetime.date(1900, 2, 28)),
    st.just(datetime.time(0, 0, 0, 1)),
    st.just(datetime.timedelta(days=-1, seconds=5)),
    st.just(decimal.Decimal("1.10")),
    st.just(b"bytes"),
    st.just(object()),
)
sheets = st.sampled_from(["Data", "Other"])
ranges = st.sampled_from(["A1:C4", "B2:B2", "J1:K3", "A10:F15", "A1:N45", "G2:G4", "C38:D42"])

BASE_OPS = [
    st.tuples(st.just("set"), sheets, coords, values),
    st.tuples(st.just("get"), sheets, coords),
    st.tuples(st.just("del"), sheets, coords),
    # (no whole-row / whole-column slices: after a write to XFD1048576 they
    # create 16384 / 1048576 cells per row / column in both libraries, which
    # copy_sheet multiplies)
    st.tuples(st.just("slice"), sheets, st.sampled_from(["A1:C3", "A1", "C5:A2", "B2:B9"])),
    st.tuples(st.just("rows"), sheets, st.integers(1, 6), st.integers(0, 4), st.booleans()),
    st.tuples(st.just("cols"), sheets, st.integers(1, 6), st.integers(0, 4), st.booleans()),
    st.tuples(st.just("move"), sheets, ranges, st.integers(-3, 5), st.integers(-3, 5), st.booleans()),
    st.tuples(st.just("merge"), sheets, ranges),
    st.tuples(st.just("unmerge"), sheets, ranges),
    st.tuples(st.just("style"), sheets, coords, st.integers(0, 6)),
    st.tuples(st.just("row_style"), sheets, st.integers(1, 45)),
    st.tuples(st.just("comment"), sheets, coords),
    st.tuples(st.just("link"), sheets, coords),
    st.tuples(st.just("append"), sheets, st.lists(values, max_size=5)),
    st.tuples(st.just("copy_sheet"), sheets),
    st.tuples(st.just("rename"), sheets),
    st.tuples(st.just("move_sheet"), sheets, st.integers(-2, 2)),
    st.tuples(st.just("dims"), sheets),
    st.tuples(st.just("values"), sheets),
    st.tuples(st.just("data_type"), sheets, coords, st.sampled_from(["n", "s", "f", "b", "e", "d", "inlineStr", "x"])),
    st.tuples(
        st.just("number_format"),
        sheets,
        ranges,
        st.sampled_from(["General", "yyyy-mm-dd", "0.0E+00", "@", '[h]:mm;"x"']),
    ),
    st.tuples(st.just("translate_formula"), sheets, coords),
    st.tuples(st.just("clear_cells"), sheets),
]
# openpyxl's insert / delete of rows and columns create a Cell for every
# position from the insertion point to max_row x max_column (_move_cells):
# combined with cells at row 1048576 / column XFD this is billions of cells,
# in both libraries. Each kind of operation has its own test.
ROW_COL_OPS = [
    st.tuples(st.just("insert_rows"), sheets, st.integers(1, 45), st.integers(1, 3)),
    st.tuples(st.just("delete_rows"), sheets, st.integers(1, 45), st.integers(1, 5)),
    st.tuples(st.just("insert_cols"), sheets, st.integers(1, 14), st.integers(1, 3)),
    st.tuples(st.just("delete_cols"), sheets, st.integers(1, 14), st.integers(1, 3)),
]
FAR_OPS = [st.tuples(st.just("set_far"), sheets, far, values)]
ops = st.one_of(*BASE_OPS, *ROW_COL_OPS)
far_ops = st.one_of(*BASE_OPS, *FAR_OPS)


def describe(x):
    if hasattr(x, "coordinate") and hasattr(x, "parent"):
        return (
            "cell",
            type(x).__name__,
            x.coordinate,
            norm(x.value),
            x.data_type,
            list(x._style) if getattr(x, "_style", None) is not None else None,
        )
    if isinstance(x, (list, tuple)):
        return [describe(y) for y in x]
    return norm(x)


STYLES = [
    lambda lib, c: setattr(c, "font", lib.styles.Font(italic=True, size=13)),
    lambda lib, c: setattr(c, "fill", lib.styles.PatternFill("solid", fgColor="123456")),
    lambda lib, c: setattr(c, "border", lib.styles.Border(top=lib.styles.Side(style="thick"))),
    lambda lib, c: setattr(c, "alignment", lib.styles.Alignment(vertical="top", indent=2)),
    lambda lib, c: setattr(c, "protection", lib.styles.Protection(locked=False)),
    lambda lib, c: setattr(c, "style", "Comma"),
    lambda lib, c: setattr(c, "number_format", "0.000"),
]


def apply(lib, wb, op):
    kind = op[0]
    ws = wb[op[1]] if op[1] in wb.sheetnames else wb.worksheets[0]
    if kind == "set":
        ws.cell(row=op[2][0], column=op[2][1]).value = op[3]
        return None
    if kind == "set_far":
        return ws.cell(row=op[2][0], column=op[2][1], value=op[3])
    if kind == "get":
        return ws.cell(row=op[2][0], column=op[2][1])
    if kind == "del":
        del ws[f"{lib.utils.get_column_letter(op[2][1])}{op[2][0]}"]
        return None
    if kind == "slice":
        key = op[2]
        if key.isdigit():
            key = int(key)
        elif ":" in key and key.replace(":", "").isdigit():
            a, b = key.split(":")
            key = slice(int(a), int(b))
        return ws[key]
    if kind == "rows":
        return list(ws.iter_rows(min_row=op[2], max_row=op[2] + op[3], max_col=6, values_only=op[4]))
    if kind == "cols":
        return list(ws.iter_cols(min_col=op[2], max_col=op[2] + op[3], max_row=8, values_only=op[4]))
    if kind == "insert_rows":
        ws.insert_rows(op[2], op[3])
        return None
    if kind == "delete_rows":
        ws.delete_rows(op[2], op[3])
        return None
    if kind == "insert_cols":
        ws.insert_cols(op[2], op[3])
        return None
    if kind == "delete_cols":
        ws.delete_cols(op[2], op[3])
        return None
    if kind == "move":
        ws.move_range(op[2], rows=op[3], cols=op[4], translate=op[5])
        return None
    if kind == "merge":
        ws.merge_cells(op[2])
        return None
    if kind == "unmerge":
        ws.unmerge_cells(op[2])
        return None
    if kind == "style":
        c = ws.cell(row=op[2][0], column=op[2][1])
        STYLES[op[3]](lib, c)
        return (c.has_style, c.style_id)
    if kind == "row_style":
        ws.row_dimensions[op[2]].font = lib.styles.Font(bold=True)
        return None
    if kind == "comment":
        c = ws.cell(row=op[2][0], column=op[2][1])
        c.comment = lib.comments.Comment("fuzz", "author")
        return None
    if kind == "link":
        c = ws.cell(row=op[2][0], column=op[2][1])
        c.hyperlink = "https://example.com/fuzz"
        return c.value
    if kind == "append":
        ws.append(op[2])
        return ws._current_row
    if kind == "copy_sheet":
        cp = wb.copy_worksheet(ws)
        return [cp.title, len(cp._cells)]
    if kind == "rename":
        ws.title = ws.title + "x"
        return wb.sheetnames
    if kind == "move_sheet":
        wb.move_sheet(ws, op[2])
        return wb.sheetnames
    if kind == "dims":
        return (ws.min_row, ws.max_row, ws.min_column, ws.max_column, ws.dimensions, ws.calculate_dimension())
    if kind == "values":
        # bounded: a cell at XFD1048576 makes ws.values 17 billion cells long
        return list(ws.iter_rows(max_row=min(ws.max_row, 60), max_col=min(ws.max_column, 20), values_only=True))
    if kind == "data_type":
        c = ws.cell(row=op[2][0], column=op[2][1])
        c.data_type = op[3]
        return None
    if kind == "number_format":
        for row in ws[op[2]]:
            for c in row:
                c.number_format = op[3]
        return None
    if kind == "translate_formula":
        c = ws.cell(row=op[2][0], column=op[2][1])
        if isinstance(c.value, str) and c.value.startswith("="):
            tr = lib.formula.translate.Translator(c.value, origin=c.coordinate)
            return tr.translate_formula("C10")
        return None
    if kind == "clear_cells":
        for k in list(ws._cells)[::3]:
            del ws._cells[k]
        return len(ws._cells)
    raise AssertionError(kind)


def run(lib, seq, mode):
    data_only, rich_text, epoch1904 = mode
    wb = lib.load_workbook(io.BytesIO(SOURCES[epoch1904]), data_only=data_only, rich_text=rich_text)
    out = []
    for op in seq:
        try:
            out.append(("ok", describe(apply(lib, wb, op))))
        except Exception as e:
            out.append(("error", type(e).__name__, ADDR_RE.sub("", norm_text(str(e)))))
    final = []
    for fn, obj in [(workbook_state, wb)] + [(sheet_state, ws) for ws in wb.worksheets]:
        try:
            final.append(fn(obj))
        except Exception as e:  # eg. cells moved to column 0
            final.append(("error", type(e).__name__, ADDR_RE.sub("", norm_text(str(e)))))
    try:
        saved = saved_parts(wb)
    except Exception as e:
        saved = ("error", type(e).__name__, norm_text(str(e)))
    return out, final, saved


modes = st.tuples(st.booleans(), st.booleans(), st.booleans())


def check(seq, mode):
    a = run(openpyxl, seq, mode)
    b = run(openrsxl, seq, mode)
    for i, (x, y) in enumerate(zip(a[0], b[0])):
        assert x == y, (i, seq[i], first_diff(x, y))
    assert not first_diff(a[1], b[1]), first_diff(a[1], b[1])
    assert not first_diff(a[2], b[2]), first_diff(a[2], b[2])


@SETTINGS
@given(st.lists(ops, min_size=1, max_size=15), modes)
def test_api_sequences(seq, mode):
    check(seq, mode)


@SETTINGS
@given(st.lists(far_ops, min_size=1, max_size=15), modes)
def test_api_sequences_extreme_coordinates(seq, mode):
    check(seq, mode)


@pytest.mark.parametrize(
    "mode", [(False, False, False), (True, False, False), (False, True, False), (False, False, True)]
)
def test_untouched_roundtrip(mode):
    """load + save right away (compact cells, nothing materialised)"""
    a = run(openpyxl, [], mode)
    b = run(openrsxl, [], mode)
    assert a == b
