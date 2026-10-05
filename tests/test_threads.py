"""
openpyxl objects can be used from any thread (under the GIL); so must
openrsxl's (the native classes must not be bound to their creating thread).
"""

import io
from concurrent.futures import ThreadPoolExecutor

import openpyxl
from oracle.compare import saved_parts

import openrsxl


def source():
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, 200):
        ws.append([r, f"s{r}", r / 7, f"=A{r}*2"])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


SOURCE = source()


def work(lib):
    wb = lib.load_workbook(io.BytesIO(SOURCE))  # loaded in the calling thread
    ws = wb.active

    def in_thread():
        vals = list(ws.iter_rows(values_only=True))
        ws["E5"] = "added"
        ws.delete_rows(3)
        cells = [c.coordinate for row in ws.iter_rows(max_row=3) for c in row]
        return vals, cells, saved_parts(wb)

    with ThreadPoolExecutor(2) as ex:
        return ex.submit(in_thread).result()


def test_workbook_used_from_other_thread():
    a, b = work(openpyxl), work(openrsxl)
    assert a[0] == b[0]
    assert a[1] == b[1]
    assert a[2] == b[2]


def test_read_only_generator_consumed_in_other_thread():
    def run(lib):
        wb = lib.load_workbook(io.BytesIO(SOURCE), read_only=True)
        gen = wb.active.iter_rows(values_only=True)
        first = next(gen)  # started in this thread
        with ThreadPoolExecutor(1) as ex:
            rest = ex.submit(lambda: list(gen)).result()
        return [first] + rest

    assert run(openpyxl) == run(openrsxl)


def test_parallel_loads():
    def load_values(_):
        wb = openrsxl.load_workbook(io.BytesIO(SOURCE))
        return list(wb.active.values)

    expected = list(openpyxl.load_workbook(io.BytesIO(SOURCE)).active.values)
    with ThreadPoolExecutor(4) as ex:
        for res in ex.map(load_values, range(8)):
            assert res == expected
