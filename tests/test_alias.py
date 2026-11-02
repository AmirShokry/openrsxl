"""
openrsxl.extended.install_as_openpyxl(): ``import openpyxl`` gives openrsxl.

Each test runs in a fresh interpreter (the test process imports the real
openpyxl: the oracle).
"""

import importlib.util
import subprocess
import sys
import textwrap

import pytest


def run(code, strict=True):
    # strict: every warning is an error (eg. "__package__ != __spec__.parent"
    # of relative imports in modules whose spec an alias would have replaced)
    flags = ["-W", "error"] if strict else []
    return subprocess.run(
        [sys.executable, *flags, "-c", textwrap.dedent(code)], capture_output=True, text=True, timeout=600
    )


def test_openpyxl_is_openrsxl():
    proc = run("""
        import io, sys
        import openrsxl.extended as ext
        ext.install_as_openpyxl()
        ext.install_as_openpyxl()  # a no-op
        import openpyxl
        import openpyxl.chart
        import openpyxl.styles.stylesheet  # saving runs its function-level relative import
        from openpyxl.worksheet.formula import ArrayFormula
        import openrsxl
        assert openpyxl is openrsxl and sys.modules["openpyxl"] is openrsxl
        assert ArrayFormula is openrsxl.worksheet.formula.ArrayFormula
        assert openpyxl.__version__ == "3.1.5"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = 1
        ws["B1"] = ArrayFormula("B1:B2", "=A1:A2*2")
        chart = openpyxl.chart.BarChart()  # saving runs function-level relative imports
        chart.add_data(openpyxl.chart.Reference(ws, min_col=1, min_row=1, max_row=1))
        ws.add_chart(chart, "D2")
        buf = io.BytesIO()
        wb.save(buf)
        data = buf.getvalue()
        # isinstance checks written for openpyxl hold for openrsxl's objects
        for kw in ({}, {"read_only": True}):
            v = openpyxl.load_workbook(io.BytesIO(data), **kw).active["B1"].value
            assert isinstance(v, ArrayFormula) and v.text == "=A1:A2*2", kw
        c = ext.load_workbook(io.BytesIO(data), read_only=True, formula_and_value=True).active["B1"]
        assert isinstance(c.formula, ArrayFormula)
        # the aliased modules keep their own spec
        import openrsxl.chart._chart
        assert openrsxl.worksheet.formula.__spec__.name == "openrsxl.worksheet.formula"
        assert openrsxl.chart._chart.__spec__.name == "openrsxl.chart._chart"
        assert ext.styles.__spec__.name == "openrsxl.styles"
        print("ok")
        """)
    assert proc.returncode == 0, proc.stderr
    assert proc.stdout.strip() == "ok"


def test_extended_aliases_keep_module_specs():
    # openrsxl.extended.<x> is openrsxl.<x>: importing it through the alias
    # must not replace the module's spec (0.1.0 did: DeprecationWarning
    # "__package__ != __spec__.parent" on its relative imports)
    proc = run("""
        import io
        import openrsxl
        import openrsxl.extended.styles.stylesheet
        assert openrsxl.extended.styles.stylesheet is openrsxl.styles.stylesheet
        assert openrsxl.styles.stylesheet.__spec__.name == "openrsxl.styles.stylesheet"
        wb = openrsxl.Workbook()
        wb.active["A1"] = 1
        wb.save(io.BytesIO())  # write_stylesheet: a function-level relative import
        print("ok")
        """)
    assert proc.returncode == 0, proc.stderr
    assert proc.stdout.strip() == "ok"


def test_refused_after_openpyxl_was_imported():
    proc = run("""
        import openpyxl
        import openrsxl.extended as ext
        try:
            ext.install_as_openpyxl()
        except RuntimeError as err:
            print("RuntimeError:", err)
        """)
    assert proc.returncode == 0, proc.stderr
    assert "openpyxl is already imported" in proc.stdout


@pytest.mark.skipif(importlib.util.find_spec("pandas") is None, reason="pandas not installed")
def test_pandas_engine_openpyxl_uses_openrsxl():
    proc = run(
        """
        import io, sys
        import openrsxl.extended as ext
        ext.install_as_openpyxl()
        import pandas as pd
        import openrsxl
        wb = openrsxl.Workbook()
        wb.active.append(["a", "b"])
        wb.active.append([1, 2.5])
        buf = io.BytesIO()
        wb.save(buf)
        df = pd.read_excel(io.BytesIO(buf.getvalue()), engine="openpyxl")
        assert df.to_dict("list") == {"a": [1], "b": [2.5]}, df
        out = io.BytesIO()
        df.to_excel(out, index=False, engine="openpyxl")
        assert openrsxl.load_workbook(io.BytesIO(out.getvalue())).active["B2"].value == 2.5
        assert sys.modules["openpyxl"] is openrsxl
        print("ok")
        """,
        strict=False,  # (pandas' own deprecation warnings are not ours)
    )
    assert proc.returncode == 0, proc.stderr
    assert proc.stdout.strip() == "ok"


def test_reprs_name_openpyxl_modules():
    # after install_as_openpyxl() the classes report openpyxl's module names
    # (object repr, Serialisable's repr, pickle) - modules imported before and
    # after the call; without it openrsxl's
    proc = run("""
        import io, pickle
        import openrsxl
        from openrsxl.worksheet.formula import ArrayFormula  # imported before
        before = repr(ArrayFormula("A1", "=1"))
        import openrsxl.extended as ext
        ext.install_as_openpyxl()
        import openpyxl
        from openpyxl.chart.bar_chart import BarChart  # imported after
        f = ArrayFormula("A1", "=1")
        assert before.startswith("<openrsxl.worksheet.formula.ArrayFormula object at "), before
        assert repr(f).startswith("<openpyxl.worksheet.formula.ArrayFormula object at "), repr(f)
        assert ArrayFormula.__module__ == "openpyxl.worksheet.formula"
        assert BarChart.__module__ == "openpyxl.chart.bar_chart"
        assert repr(openpyxl.styles.Font()).startswith("<openpyxl.styles.fonts.Font object>")
        assert pickle.loads(pickle.dumps(openpyxl.styles.Font(b=True))).b is True
        assert type(pickle.loads(pickle.dumps(f))) is ArrayFormula
        wb = openpyxl.Workbook()
        wb.active["A1"] = f
        buf = io.BytesIO()
        wb.save(buf)
        assert repr(openpyxl.load_workbook(buf).active["A1"].value).startswith("<openpyxl.worksheet.formula.")
        assert ext.load_workbook.__module__ == "openrsxl.extended._load"  # extended: its own
        print("ok")
        """)
    assert proc.returncode == 0, proc.stderr
    assert proc.stdout.strip() == "ok"
