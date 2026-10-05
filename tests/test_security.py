"""
Hostile input: every case must behave exactly like openpyxl (same values or
the same exception) in every load mode and with the openrsxl.extended
streaming features, without crashing the interpreter (stack overflows,
Rust panics), leaking local files (XXE) or expanding entities without bound.
"""

import io
import zipfile

import pytest
from fixtures import MAIN, cells, make_xlsx, sheet_xml, strings_xml
from oracle.compare import compare_workbooks
from oracle.extended import compare_extended

import openrsxl
import openrsxl.extended

MODES = [
    {},
    {"read_only": True},
    {"data_only": True},
    {"rich_text": True},
    {"read_only": True, "data_only": True},
]

SECRET = "openrsxl-xxe-secret-7f3a"


def c(value=None, t=None, r="A1", inner=None, extra=""):
    attrs = f' r="{r}"' if r is not None else ""
    if t is not None:
        attrs += f' t="{t}"'
    attrs += extra
    if inner is None:
        inner = "" if value is None else f"<v>{value}</v>"
    return f"<c{attrs}>{inner}</c>"


def check_all_modes(data, extended=True):
    for mode in MODES:
        diffs = compare_workbooks(data, **mode)
        assert not diffs, f"{mode}: " + "\n".join(diffs)
    for data_only in (False, True):
        if extended:
            diffs = compare_extended(data, data_only=data_only)
        else:
            diffs = extended_like_read_only(data, data_only)
        assert not diffs, f"extended data_only={data_only}: " + "\n".join(diffs)


def extended_like_read_only(data, data_only):
    """
    For cells whose coordinate disagrees with their <row> (r="A0" in row 1)
    openpyxl's full and read-only modes place them differently; streaming
    follows read-only mode: compare the cells with it.
    """
    import openpyxl
    from oracle.extended import ALL_FLAGS, _deferred

    diffs = []
    try:
        ext = openrsxl.extended.load_workbook(io.BytesIO(data), read_only=True, data_only=data_only, **ALL_FLAGS)
    except Exception as err:
        try:
            openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=data_only)
        except Exception as exp:
            if (type(err).__name__, str(err)) != (type(exp).__name__, str(exp)):
                diffs.append(f"load error {err!r} != {exp!r}")
            return diffs
        return [f"load error {err!r}, read-only mode loads"]
    if not _deferred(openpyxl, ext, data, data_only, diffs.append):
        diffs.append("read-only mode fails to load")
    for ws in ext.worksheets:
        try:
            list(ws.iter_rows())
        except Exception:
            continue  # start_cell reads the same rows (compared above)
        for r in ws.merged_cells:
            if r.start_cell is None:
                diffs.append(f"{r.coord}.start_cell is None")
    return diffs


def all_values(data):
    """Every cell value openrsxl reads in any mode (errors are fine)."""
    out = []
    for kw in MODES + [dict(read_only=True, formula_and_value=True, read_comments=True, read_hyperlinks=True)]:
        lib = openrsxl.extended if "formula_and_value" in kw else openrsxl
        try:
            wb = lib.load_workbook(io.BytesIO(data), **kw)
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for cell in row:
                        out.append(repr(getattr(cell, "value", None)))
                        out.append(repr(getattr(cell, "comment", None)))
            out.extend(repr(n) for n in wb.defined_names)
        except Exception as err:
            out.append(repr(err))
    return "\n".join(out)


# --- XML external entities -------------------------------------------------


def _xxe_doctype(root, path):
    uri = path.as_uri()
    return f'<?xml version="1.0"?><!DOCTYPE {root} [<!ENTITY xxe SYSTEM "{uri}">' f'<!ENTITY % pe SYSTEM "{uri}">]>'


@pytest.fixture
def secret_file(tmp_path):
    p = tmp_path / "secret.txt"
    p.write_text(SECRET)
    return p


def test_xxe_in_worksheet(secret_file):
    doc = _xxe_doctype("worksheet", secret_file) + sheet_xml(
        cells(c("&xxe;", t="str"), c(None, t="inlineStr", r="B1", inner="<is><t>&xxe;</t></is>")), decl=False
    )
    data = make_xlsx(raw_sheets=[doc])
    check_all_modes(data)
    assert SECRET not in all_values(data)


def test_xxe_in_shared_strings(secret_file):
    sst = (
        _xxe_doctype("sst", secret_file)
        + strings_xml("<si><t>&xxe;</t></si>")[len('<?xml version="1.0" encoding="UTF-8"?>') :]
    )
    data = make_xlsx([cells(c("0", t="s"))], strings=sst)
    check_all_modes(data)
    assert SECRET not in all_values(data)


@pytest.mark.parametrize(
    "part", ["xl/workbook.xml", "xl/styles.xml", "[Content_Types].xml", "xl/_rels/workbook.xml.rels"]
)
def test_xxe_in_package_parts(secret_file, part):
    data = make_xlsx([cells(c("1"))])
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for name in zin.namelist():
            xml = zin.read(name)
            if name == part:
                text = xml.decode("utf-8")
                body = text[text.index("?>") + 2 :] if text.startswith("<?xml") else text
                root = body[1 : body.index(" ")]
                xml = (_xxe_doctype(root, secret_file) + body).encode("utf-8")
            z.writestr(name, xml)
    data = out.getvalue()
    check_all_modes(data)
    assert SECRET not in all_values(data)


# --- entity expansion ------------------------------------------------------

LAUGHS = (
    '<!DOCTYPE {root} [<!ENTITY lol "lol">'
    + "".join(f'<!ENTITY lol{i} "{("&lol%s;" % (i - 1 if i > 1 else "")) * 10}">' for i in range(1, 10))
    + "]>"
)


def test_billion_laughs_worksheet():
    doc = '<?xml version="1.0"?>' + LAUGHS.format(root="worksheet") + sheet_xml(cells(c("&lol9;", t="str")), decl=False)
    check_all_modes(make_xlsx(raw_sheets=[doc]))


def test_billion_laughs_shared_strings():
    sst = '<?xml version="1.0"?>' + LAUGHS.format(root="sst") + f'<sst xmlns="{MAIN}"><si><t>&lol9;</t></si></sst>'
    check_all_modes(make_xlsx([cells(c("0", t="s"))], strings=sst))


# --- nesting and repetition (no recursion in the Rust engine) --------------

DEPTH = 100_000


@pytest.mark.parametrize("where", ["cell", "inline-string", "row", "sheetData", "worksheet", "extLst"])
def test_deep_nesting(where):
    deep = "<x>" * DEPTH + "</x>" * DEPTH
    if where == "cell":
        doc = sheet_xml(cells(c(None, inner="<v>1</v>" + deep)))
    elif where == "inline-string":
        doc = sheet_xml(cells(c(None, t="inlineStr", inner="<is><t>a</t>" + deep + "</is>")))
    elif where == "row":
        doc = sheet_xml('<row r="1"><c r="A1"><v>1</v></c>' + deep + "</row>")
    elif where == "sheetData":
        doc = sheet_xml(cells(c("1")) + deep)
    elif where == "worksheet":
        doc = sheet_xml(cells(c("1")), before=deep, after=deep)
    else:
        doc = sheet_xml(cells(c("1")), after="<extLst>" + deep + "</extLst>")
    check_all_modes(make_xlsx(raw_sheets=[doc]))


def test_deep_nesting_shared_strings():
    deep = "<x>" * DEPTH + "</x>" * DEPTH
    data = make_xlsx([cells(c("0", t="s"))], strings=strings_xml("<si><t>a</t>" + deep + "</si>"))
    check_all_modes(data)


@pytest.mark.parametrize("where", ["before", "after", "both"])
def test_many_comments_and_pis_outside_root(where):
    junk = "<!-- c -->\n<?pi x?>\n" * 100_000
    body = sheet_xml(cells(c("1")), decl=False)
    doc = (
        '<?xml version="1.0"?>\n'
        + (junk if where != "after" else "")
        + body
        + ("\n" + junk if where != "before" else "")
    )
    check_all_modes(make_xlsx(raw_sheets=[doc]))


def test_many_comments_inside_sheet_data():
    doc = sheet_xml(cells(c("1")) + "<!-- c -->" * 100_000 + cells(c("2", r="A2"), r=2))
    check_all_modes(make_xlsx(raw_sheets=[doc]))


# --- numbers and coordinates beyond any limit -----------------------------


@pytest.mark.parametrize(
    "value",
    [
        "1" * 4300,
        "1" * 4301,
        "9" * 100_000,
        "-" + "1" * 5000,
        "1" * 5000 + ".5",
        "1e" + "9" * 5000,
    ],
    ids=[
        "4300-digits",
        "4301-digits",
        "100k-digits",
        "negative-5000-digits",
        "5000-digits-fraction",
        "5000-digit-exponent",
    ],
)
def test_huge_digit_strings(value):
    check_all_modes(make_xlsx([cells(c(value), c(value, t="n", r="B1", extra=' s="1"'))]))


@pytest.mark.parametrize(
    "coord",
    [
        "A99999999999999999999",
        "A9223372036854775808",
        "XFD1048577",
        "XFE1",
        "AAAA1",
        "ZZZ1",
        "A" + "9" * 5000,
        "ZZZZZZZZZZZZZZZZZZZZZZZZ1",
        "A-1",
        "A0",
    ],
    ids=lambda v: v if len(v) < 30 else f"A-{len(v) - 1}-digits",
)
def test_coordinate_overflow(coord):
    doc = sheet_xml(
        cells(c("1", r=coord)),
        before='<dimension ref="A1:B2"/>',
        after='<mergeCells count="1"><mergeCell ref="A1:B2"/></mergeCells>',
    )
    # every coordinate but ZZZ1 disagrees with its <row r="1">
    check_all_modes(make_xlsx(raw_sheets=[doc]), extended=coord == "ZZZ1")


@pytest.mark.parametrize(
    "row",
    ["99999999999999999999", "9223372036854775808", "-9223372036854775809", "1048577", "0", "-1", "1" * 5000],
    ids=lambda v: v if len(v) < 30 else "5000-digits",
)
def test_row_number_overflow(row):
    doc = sheet_xml(
        f'<row r="{row}"><c><v>1</v></c></row>',
        before='<dimension ref="A1:B2"/>',
        after='<mergeCells count="1"><mergeCell ref="A1:B2"/></mergeCells>',
    )
    check_all_modes(make_xlsx(raw_sheets=[doc]))


@pytest.mark.parametrize(
    "style",
    ["-1", "99999999999999999999", "1" * 5000, "1.5", " 1"],
    ids=lambda v: repr(v) if len(v) < 30 else "5000-digits",
)
def test_style_index_overflow(style):
    check_all_modes(make_xlsx([cells(c("1", extra=f' s="{style}"'))]))


@pytest.mark.parametrize(
    "formula",
    [
        "=" + "(" * 50_000 + "1" + ")" * 50_000,
        "=" + "A1+" * 50_000 + "1",
        "=ZZZ1+XFD1048576+$ZZZ$1",
    ],
    ids=["50k-parentheses", "50k-references", "extreme-references"],
)
def test_formula_size_and_extreme_references(formula):
    # shared formulae are translated to their dependents (possibly beyond the
    # last column / row)
    inner = f'<f t="shared" ref="A1:C3" si="0">{formula[1:]}</f><v>1</v>'
    doc = sheet_xml(
        cells(c(None, inner=inner), c(None, r="C1", inner='<f t="shared" si="0"/><v>2</v>'))
        + cells(c(None, r="A3", inner='<f t="shared" si="0"/>'), r=3)
    )
    check_all_modes(make_xlsx(raw_sheets=[doc]))


# --- malformed references in the extended pre-scan (former Rust panic) ----


@pytest.mark.parametrize("ref", ["&#x;", "&#0;", "&#xZZ;", "&#99999999999;", "&#xD800;", "&#;", "&nbsp;", "&#x110000;"])
@pytest.mark.parametrize("where", ["row", "cell"])
def test_bad_character_reference_in_prescanned_tags(ref, where):
    row_attr = f' ht="{ref}"' if where == "row" else ""
    cell_attr = f' s="{ref}"' if where == "cell" else ""
    doc = sheet_xml(
        f'<row r="1"{row_attr}><c r="A1"{cell_attr}><v>1</v></c></row>',
        before='<dimension ref="A1:B2"/>',
        after='<mergeCells count="1"><mergeCell ref="A1:B2"/></mergeCells>',
    )
    data = make_xlsx(raw_sheets=[doc])
    for flags in ({"read_dimensions": True}, {"read_merged_cells": True}):
        diffs = compare_extended(data, flags=flags)
        assert not diffs, "\n".join(diffs)


# --- package level --------------------------------------------------------


@pytest.mark.parametrize(
    "target",
    [
        "../../../../etc/passwd",
        "/../../secret.xml",
        "C:/Windows/win.ini",
        "worksheets/../../../x.xml",
        "\\\\server\\share\\x.xml",
    ],
)
def test_relationship_targets_outside_package(target):
    data = make_xlsx([cells(c("1"))])
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for name in zin.namelist():
            xml = zin.read(name)
            if name == "xl/_rels/workbook.xml.rels":
                xml = xml.replace(b'Target="worksheets/sheet1.xml"', f'Target="{target}"'.encode())
            z.writestr(name, xml)
    check_all_modes(out.getvalue())


def test_duplicate_zip_members():
    data = make_xlsx([cells(c("1"))])
    out = io.BytesIO(data)
    with pytest.warns(UserWarning):  # zipfile: "Duplicate name"
        with zipfile.ZipFile(out, "a") as z:
            z.writestr("xl/worksheets/sheet1.xml", sheet_xml(cells(c("2"))))
    check_all_modes(out.getvalue())
